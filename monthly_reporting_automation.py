#!/usr/bin/env python3
"""Monthly reporting automation runner.

This script updates Sarmayacar's master Monthly Reporting workbook from
portfolio-company KPI packs. ABHI is implemented as the first reusable profile.
"""

from __future__ import annotations

import argparse
import io
import json
import re
from dataclasses import dataclass
from functools import lru_cache
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Mapping
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from formula_memory import apply_formula_memory, audit_formula_memory
from normalization_memory import (
    memory_aliases_for_profile,
    memory_sheet_aliases_for_profile,
    normalise_memory_key,
)


DEFAULT_MONTHS = ((2026, 1), (2026, 2), (2026, 3))
DEMO_NOTICE = "DEMO — ILLUSTRATIVE DATA"

ONELOAD_FX_RATES = {
    (2025, 4): 280.9,
    (2025, 5): 282.1,
    (2025, 6): 280.2,
    (2025, 7): 282.1,
    (2025, 8): 280.2,
    (2025, 9): 282.2447,
    (2025, 10): 280.2,
    (2025, 11): 282.2447,
    (2025, 12): 282.2447,
    (2026, 1): 279.93,
    (2026, 2): 279.6,
    (2026, 3): 279.02,
}

TAPMAD_FX_RATES = {
    (2025, 4): 280.713,
    (2025, 5): 281.66,
    (2025, 6): 283.0,
    (2025, 7): 284.2133,
    (2025, 8): 282.2447,
    (2025, 9): 283.0,
    (2025, 10): 282.2447,
    (2025, 11): 283.0,
    (2025, 12): 282.2447,
}

FX_MANAGED_PROFILES = ("simpaisa", "oneload", "tapmad", "roomy", "procheck")
NBP_RATE_SHEET_URL_TEMPLATE = (
    "https://www.nbp.com.pk/RateSheetFiles/NBP-RateSheet-{day:02d}-{month:02d}-{year}.pdf"
)
NBP_LIVE_FX_PROFILES = {"simpaisa", "oneload", "tapmad"}


@lru_cache(maxsize=256)
def _download_nbp_rate_sheet_pdf(url: str) -> bytes | None:
    request = Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
            "Accept": "application/pdf,*/*",
            "Referer": "https://www.nbp.com.pk/RateSheet/index.aspx",
            "Origin": "https://www.nbp.com.pk",
        },
    )
    try:
        with urlopen(request, timeout=20) as response:
            return response.read()
    except (HTTPError, URLError, TimeoutError, OSError):
        return None


def _nbp_rate_sheet_candidates(month_key: tuple[int, int]) -> list[tuple[date, str]]:
    year, month = month_key
    last_day = monthrange(year, month)[1]
    return [
        (date(year, month, day), NBP_RATE_SHEET_URL_TEMPLATE.format(year=year, month=month, day=day))
        for day in range(last_day, 0, -1)
    ]


def _parse_nbp_usd_rate(pdf_bytes: bytes) -> float | None:
    try:
        import pdfplumber
    except ImportError:
        return None

    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                for raw_line in text.splitlines():
                    line = " ".join(raw_line.split())
                    upper = line.upper()
                    if "US DOLLAR" not in upper or "USD" not in upper:
                        continue
                    numbers = [
                        float(number.replace(",", ""))
                        for number in re.findall(r"-?\d[\d,]*(?:\.\d+)?", line)
                    ]
                    if len(numbers) >= 2:
                        return numbers[1]
    except Exception:
        return None
    return None


@lru_cache(maxsize=128)
def _nbp_month_end_rate(month_key: tuple[int, int]) -> tuple[float, date, str] | None:
    for rate_date, url in _nbp_rate_sheet_candidates(month_key):
        pdf_bytes = _download_nbp_rate_sheet_pdf(url)
        if not pdf_bytes:
            continue
        rate = _parse_nbp_usd_rate(pdf_bytes)
        if rate is not None:
            return rate, rate_date, url
    return None


def _fx_rate_metadata(
    profile_name: str,
    month_key: tuple[int, int],
    kpi_workbook=None,
) -> tuple[object | None, str, str]:
    if profile_name in NBP_LIVE_FX_PROFILES:
        nbp_rate = _nbp_month_end_rate(month_key)
        if nbp_rate is not None:
            rate, rate_date, url = nbp_rate
            return (
                rate,
                "NBP month-end PDF",
                f"Pulled from {rate_date:%d %b %Y} NBP rate sheet ({Path(url).name}).",
            )
    if profile_name == "oneload":
        value = ONELOAD_FX_RATES.get(month_key)
        if value is not None:
            return value, "Fallback rate table", "Used only if the live NBP PDF lookup is unavailable."
    if profile_name == "tapmad":
        value = TAPMAD_FX_RATES.get(month_key)
        if value is not None:
            return value, "Fallback rate table", "Used only if the live NBP PDF lookup is unavailable."
    if profile_name == "simpaisa":
        legacy_fx = {
            (2024, 1): 280.3206,
            (2024, 2): 279.18,
            (2024, 3): 278.705,
            (2025, 4): 280.713,
            (2025, 5): 281.66,
            (2025, 6): 283.0,
            (2025, 7): 284.2133,
            (2025, 8): 282.2447,
            (2025, 9): 282.2447,
        }
        value = legacy_fx.get(month_key)
        if value is not None:
            return value, "Fallback rate table", "Used only if the live NBP PDF lookup is unavailable."
        return 282, "Fallback rate table", "Used only if the live NBP PDF lookup is unavailable."
    if profile_name == "roomy" and kpi_workbook is not None:
        source_sheet = find_workbook_month_sheet(kpi_workbook, month_key)
        if source_sheet is None:
            return None, "KPI workbook label not found", "Editable if the KPI sheet uses a new FX line."
        value = find_roomy_fx_label(source_sheet)
        return value, "KPI workbook label", "Pulled from the uploaded KPI file."
    return None, "Not configured", "No FX source is configured for this company."


def default_fx_rate(
    profile_name: str,
    month_key: tuple[int, int],
    kpi_workbook=None,
) -> object | None:
    value, _, _ = _fx_rate_metadata(profile_name, month_key, kpi_workbook)
    return value


def describe_fx_rate(
    profile_name: str,
    month_key: tuple[int, int],
    kpi_workbook=None,
) -> tuple[object | None, str, str]:
    return _fx_rate_metadata(profile_name, month_key, kpi_workbook)


def resolve_fx_rate(
    profile_name: str,
    month_key: tuple[int, int],
    kpi_workbook=None,
    fx_overrides: Mapping[str, Mapping[tuple[int, int], object]] | None = None,
) -> object | None:
    if fx_overrides:
        override = fx_overrides.get(profile_name, {}).get(month_key)
        if override not in (None, ""):
            return override
    return default_fx_rate(profile_name, month_key, kpi_workbook)


@dataclass(frozen=True)
class SectionMap:
    start_row: int
    end_row: int
    source_sheet: str


@dataclass(frozen=True)
class CompanyProfile:
    name: str
    display_name: str
    report_sheet: str
    sections: tuple[SectionMap, ...]


@dataclass(frozen=True)
class SourceSpec:
    sheet_name: str
    label_columns: tuple[int, ...]
    month_value_columns: Mapping[int, int] | None = None
    value_column: int | None = None
    sheet_month: int | None = None


@dataclass(frozen=True)
class LabelCopyProfile:
    name: str
    display_name: str
    report_sheet: str
    source_specs: tuple[SourceSpec, ...]
    aliases: Mapping[str, str] | None = None
    row_rules: Mapping[str, tuple[tuple[str, int, float], ...]] | None = None
    fixed_month_sheet_rules: Mapping[str, tuple[int, int, float]] | None = None
    zero_as_dash_labels: tuple[str, ...] = ()
    date_rows: tuple[int, ...] = ()
    formula_rows: tuple[int, ...] = ()
    ignored_rows: tuple[int, ...] = ()
    round_values: bool = False


@dataclass(frozen=True)
class ValidationResult:
    profile: str
    company: str
    status: str
    confidence: int
    issues: tuple[str, ...] = ()
    warnings: tuple[str, ...] = ()


ABHI_PROFILE = CompanyProfile(
    name="abhi",
    display_name="ABHI",
    report_sheet="Abhi Finance",
    sections=(
        SectionMap(592, 756, "Abhi Private Ltd"),
        SectionMap(757, 844, "Abhi Payriff"),
        SectionMap(845, 1038, "Abhi Microfinance Bank"),
        SectionMap(1039, 1127, "Abhi FinTech"),
        SectionMap(1128, 1208, "Abhi Consulting"),
        SectionMap(1209, 1297, "Abhi Middle East"),
        SectionMap(1298, 1371, "Abhi Technologies"),
        SectionMap(1372, 1448, "Abhi LLC"),
        SectionMap(1449, 1526, "Abhi Limited"),
        SectionMap(1527, 2210, "Consolidated"),
    ),
)

ABHI_BASE_SECTION_START_ROWS = {
    section.source_sheet: section.start_row for section in ABHI_PROFILE.sections
}

ABHI_SECTION_HEADING_ALIASES = {
    "Abhi Private Ltd": ("abhi pvt", "abhi private", "abhi pak"),
    "Abhi Payriff": ("abhi payriff",),
    "Abhi Microfinance Bank": ("abhi microfinance bank",),
    "Abhi FinTech": ("abhi fintech",),
    "Abhi Consulting": ("abhi consulting",),
    "Abhi Middle East": ("abhi middle east",),
    "Abhi Technologies": ("abhi technologies", "abhi technologies ksa"),
    "Abhi LLC": ("abhi llc", "abhi llc oman"),
    "Abhi Limited": ("abhi limited",),
    "Consolidated": ("abhi group consolidated", "consolidated"),
}

ABHI_EXPLICIT_ONLY_SHEETS = {
    "Abhi Payriff",
    "Abhi Consulting",
    "Abhi Middle East",
    "Abhi Technologies",
}

ABHI_ZERO_AS_DASH_ROWS = {
    597, 598, 600, 603, 605, 627, 637, 646, 656, 677, 678, 682, 698,
    699, 700, 701, 765, 767, 768, 772, 773, 775, 777, 781, 842, 867,
    879, 882, 883, 884, 892, 900, 907, 918, 920, 949, 954, 1010, 1044,
    1046, 1052, 1089, 1090, 1101, 1112, 1113, 1118, 1124, 1140, 1142,
    1145, 1188, 1227, 1286, 1312, 1317, 1319, 1320, 1339, 1358, 1383,
    1385, 1390, 1391, 1399, 1400, 1403, 1414, 1436, 1443, 1460, 1462,
    1467, 1468, 1472, 1475, 1498,
}

ABHI_ZERO_AS_BLANK_ROWS = {
    671, 672, 674, 675, 676, 679, 696, 697, 749, 947, 1043, 1045,
    1047, 1048, 1049, 1050, 1051, 1054, 1055, 1063, 1093, 1107, 1108,
    1109, 1110, 1111, 1116, 1137, 1138, 1139, 1143, 1144, 1146, 1152,
    1157, 1161, 1164, 1171, 1183, 1203, 1217, 1219, 1220, 1224, 1225,
    1232, 1233, 1238, 1271, 1309, 1310, 1311, 1315, 1316, 1324, 1341,
    1351, 1352, 1357, 1364, 1369, 1380, 1381, 1382, 1384, 1386, 1387,
    1388, 1389, 1392, 1393, 1394, 1395, 1396, 1397, 1398, 1401, 1402,
    1407, 1418, 1420, 1422, 1429, 1430, 1431, 1435, 1441, 1442, 1447,
    1457, 1458, 1459, 1461, 1463, 1464, 1465, 1466, 1470, 1471, 1473,
    1474, 1478, 1479, 1496, 1500, 1520, 1525,
}

ABHI_FORCE_BLANK_ROWS = {
    1243, 1260, 1269, 1271, 1369, 1456, 1484, 1497,
}

ABHI_ROUND_WHOLE_ROWS = {828, 836, 966, 991, 1002, 1270}

ABHI_CARRY_FORWARD_ON_ZERO_ROWS = {1115, 1117}

ABHI_LITERAL_ROW_VALUES: dict[int, object] = {
    765: "-",
    775: "-",
}

ABHI_LITERAL_MONTH_VALUES: dict[tuple[int, tuple[int, int]], object] = {
    (670, (2026, 1)): "Product shifted to Abhi Bank",
    (982, (2026, 3)): ".",
    (1414, (2026, 1)): "-",
    (1414, (2026, 2)): "-",
}

ABHI_EXPLICIT_ROW_MAPS: dict[str, dict[int, int]] = {
    "Abhi Private Ltd": {
        129: 709,
    },
    "Abhi Payriff": {
        16: 762, 17: 763, 18: 764, 20: 766, 21: 767, 22: 768, 23: 769,
        24: 770, 25: 771, 26: 772, 27: 773, 28: 774, 30: 776, 31: 777,
        32: 778, 33: 779, 34: 780, 35: 781, 36: 782, 41: 787, 42: 788,
        43: 789, 44: 790, 47: 793, 48: 794, 49: 795, 50: 796, 51: 797,
        53: 799, 56: 802, 57: 803, 58: 804, 61: 807, 62: 808, 63: 809,
        64: 810, 66: 812, 74: 819, 77: 822, 78: 823, 80: 825, 81: 826,
        83: 828, 84: 829, 87: 832, 88: 833, 89: 834, 91: 836, 92: 837,
        94: 839, 95: 840, 96: 841, 97: 842,
    },
    "Abhi FinTech": {
        80: 1107, 81: 1108, 82: 1109, 83: 1110, 84: 1111, 85: 1112,
        86: 1113, 88: 1115, 89: 1116, 90: 1117, 91: 1118,
        96: 1123, 97: 1124, 98: 1125, 99: 1126,
    },
    "Abhi Microfinance Bank": {
        17: 850, 178: 1011,
    },
    "Abhi Consulting": {
        16: 1132, 17: 1133, 18: 1134, 19: 1135, 20: 1136, 21: 1137,
        22: 1138, 23: 1139, 24: 1140, 25: 1141, 26: 1142, 27: 1143,
        28: 1144, 29: 1145, 30: 1146, 31: 1147, 32: 1148, 33: 1149,
        34: 1150, 35: 1151, 36: 1152, 37: 1153, 40: 1157, 43: 1159,
        44: 1160, 45: 1161, 46: 1162, 47: 1163, 48: 1164, 49: 1165,
        51: 1167, 54: 1170, 55: 1171, 56: 1172, 57: 1173, 60: 1176,
        61: 1177, 64: 1180, 65: 1181, 66: 1182, 67: 1183, 68: 1184,
        70: 1186, 72: 1188, 78: 1194, 79: 1195, 80: 1196, 81: 1197,
        82: 1198, 84: 1202, 85: 1203, 87: 1205, 88: 1206,
    },
    "Abhi Middle East": {
        16: 1213, 17: 1214, 18: 1215, 19: 1216, 20: 1217, 21: 1218,
        22: 1219, 23: 1220, 24: 1221, 25: 1222, 26: 1223, 27: 1224,
        28: 1225, 29: 1226, 30: 1227, 31: 1228, 32: 1229, 33: 1230,
        34: 1231, 35: 1232, 36: 1233, 37: 1234, 41: 1238, 44: 1241,
        45: 1242, 46: 1243, 47: 1244, 48: 1245, 49: 1246, 50: 1247,
        52: 1249, 55: 1252, 57: 1254, 58: 1255, 61: 1258, 62: 1259,
        66: 1263, 67: 1264, 68: 1265, 69: 1266, 71: 1268, 73: 1270,
        79: 1277, 80: 1278, 81: 1279, 82: 1281, 83: 1282, 84: 1283,
        86: 1285, 87: 1286, 92: 1291, 93: 1292, 94: 1293,
    },
    "Abhi Technologies": {
        17: 1303, 19: 1305, 20: 1306, 21: 1307, 23: 1309, 24: 1310,
        25: 1311, 26: 1312, 27: 1313, 28: 1314, 29: 1315, 30: 1316,
        31: 1317, 32: 1318, 33: 1319, 34: 1320, 35: 1321, 36: 1322,
        37: 1323, 38: 1324, 39: 1325, 46: 1339, 47: 1340, 48: 1341,
        49: 1342, 50: 1343, 51: 1344, 52: 1345, 54: 1347, 57: 1350,
        58: 1351, 59: 1352, 60: 1353, 61: 1354, 64: 1357, 65: 1358,
        68: 1361, 69: 1362, 70: 1363, 71: 1364, 72: 1365, 74: 1367,
        76: 1369, 80: 1379,
    },
    "Abhi Limited": {
        43: 1490, 44: 1491, 45: 1492, 48: 1495, 49: 1496, 50: 1497,
        51: 1498, 52: 1499, 53: 1500, 54: 1501, 56: 1503, 59: 1506,
        60: 1507, 61: 1508, 62: 1509, 63: 1510, 66: 1513, 67: 1514,
        70: 1517, 71: 1518, 72: 1519, 73: 1520, 74: 1521, 76: 1523,
        78: 1525,
    },
    "Consolidated": {
        103: 1560, 104: 1561, 105: 1562, 107: 1563, 112: 1569,
    },
}

ABHI_CONSOLIDATED_SOURCE_LABEL_ALIASES = {
    "total invoices factored": (
        "total invoices factored",
        "total b2b transactions",
    ),
    "total bank loans": (
        "total bank loans",
        "total gold + other loans",
    ),
    "total outstanding portfolio": (
        "total outstanding portfolio",
    ),
    "bank staff": (
        "abhi microfinance bank - pakistan",
    ),
}

BUILT_IN_SOURCE_SHEET_ALIASES = {
    "abhi": {
        "Abhi Private Ltd": (
            "Abhi Private - Pakistan (USD)",
        ),
        "Abhi Payriff": (
            "AbhiPayriff - Pakistan (USD)",
        ),
        "Abhi FinTech": (
            "Abhi Fintech - UAE",
        ),
    },
}

SIMPAISA_PROFILE = LabelCopyProfile(
    name="simpaisa",
    display_name="SimPaisa",
    report_sheet="SimPaisa ",
    source_specs=(SourceSpec("Sheet1", (2,)),),
    aliases={
        "Net Revenue": "Net Revenue",
        "Publishex Revenue %": "Simpaisa Revenue %",
        "Total income for the period": "Total income for the period",
    },
    formula_rows=(),
)

TAPMAD_PROFILE = LabelCopyProfile(
    name="tapmad",
    display_name="Tapmad",
    report_sheet="Tapmad",
    source_specs=(SourceSpec("Sheet1", (3,)),),
    date_rows=(45,),
    formula_rows=(6, 7, 8, 9, 10, 11, 13, 14, 16, 17, 18, 19, 21, 23, 24, 25, 26, 27, 28, 29, 31, 33, 35, 46, 47, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58),
)

REVOLVING_GAMES_PROFILE = LabelCopyProfile(
    name="revolving_games",
    display_name="Revolving Games",
    report_sheet="Revolving Games",
    source_specs=(SourceSpec("P_L_for_Q1_2025", (1,)),),
    aliases={
        "Operations and BD": "Operations",
        "Net Income (loss)": "Net Burn",
    },
    zero_as_dash_labels=("Software and Servers",),
    ignored_rows=tuple(range(26, 36)),
)

ROOMY_PROFILE = LabelCopyProfile(
    name="roomy",
    display_name="Roomy",
    report_sheet="Roomy",
    source_specs=(
        SourceSpec("Jan 2026", (3, 4), value_column=7, sheet_month=1),
        SourceSpec("Feb 2026", (3, 4), value_column=7, sheet_month=2),
        SourceSpec("Mar 2026", (3, 4), value_column=7, sheet_month=3),
    ),
    aliases={
        "Numbers in 000s": "Numebrs in 000s",
        "Less: SG&A Expense (USD)": "LESS: SG&A EXPENSE",
        "Net Profit Before Tax / (Loss)": "EBITDA",
        "Less: Tax": "LESS: Income Tax",
        "Net Profit / (Loss)": "NET PROFIT / (LOSS)",
    },
    fixed_month_sheet_rules={
        "Cash bal ($)": (5, 12, 1),
        "Cash bal (PKR)": (4, 12, 1),
        "Numbers in 000s": (3, 7, 1),
        "Total Rooms": (36, 5, 1),
        "# Of RMNTS Available": (37, 5, 1),
        "# Of RMNTS Occupied": (38, 5, 1),
        "# Of Rooms Available": (39, 5, 1),
        "Occupancy %": (41, 5, 1),
        "Operational Rooms": (44, 5, 1),
        "RevPar": (50, 7, 1),
        "ADR": (48, 7, 1),
    },
    formula_rows=(8, 13, 14, 15, 16, 17, 18, 19, 20, 21, 25, 26, 27, 28),
)

BYKEA_PROFILE = LabelCopyProfile(
    name="bykea",
    display_name="Bykea",
    report_sheet="Bykea",
    source_specs=(SourceSpec("Plan", (3, 4)),),
    aliases={
        "GTV": "GBV / GMV / GTV",
        "Bookings / day": "Gross Daily Bookings",
        "MAU (Monthly Active Users) Net": "Consolidated MTUs",
        "MAD (Monthly Active Drivers)": "Consolidated MADs",
        "Net Tranactions/day": "Net Transactions/day",
    },
    row_rules={
        "Net Revenue": (("Plan", 1067, 1),),
        "Driver Incentives": (("Plan", 1129, 1),),
        "Marketing": (("Plan", 1154, 1),),
        "Tech": (("Plan", 1158, 1),),
        "Overheads": (("Plan", 1152, 1), ("Plan", 1157, 1), ("Plan", 1158, 1), ("Plan", 1119, 1)),
    },
    formula_rows=(22, 26, 30, 31, 35, 36, 37, 38, 39, 41, 42, 43, 44),
)

BYKEA_VISIBLE_SUMMARY_LABELS = {
    "net revenue": ("pc1",),
    "driver incentives": ("driver incent.", "driver incentives"),
    "marketing": ("marketing",),
    "tech": ("cloud variable costs", "total cloud variable costs"),
    "overheads": ("people", "facilities", "other g&a", "others"),
}

BYKEA_COMPOSITE_SUMMARY_LABELS = {
    "driver incentives": (("commission refund",), ("driver incent.", "driver incentives")),
    "overheads": (("people",), ("facilities",), ("other g&a",), ("others",)),
}

PROCHECK_PROFILE = LabelCopyProfile(
    name="procheck",
    display_name="Procheck",
    report_sheet="Procheck",
    source_specs=(SourceSpec("Procheck", (1,)),),
    aliases={
        "Revenue (New Lines/Codes)": "Revenue (New Lines)",
        "Cost of Sales": "COGS",
        "Software Development Cost": "Infrastructure Cost",
        "Other Expenses": "Subscription Cost",
        "No. of Lines (OEE)": "No. of Lines (Analytics)",
    },
    row_rules={
        "Revenue (Recurring Lines/Codes)": (("Procheck", 28, 1), ("Procheck", 32, 1), ("Procheck", 36, 1)),
        "Codes Generated (TnT)": (("Procheck", 61, 1),),
    },
    round_values=True,
    formula_rows=(18, 20, 25, 26, 101, 102, 103, 104, 105, 106, 107, 108, 109, 110, 111),
    ignored_rows=tuple(range(57, 100)),
)

OLADOC_PROFILE = LabelCopyProfile(
    name="oladoc",
    display_name="Oladoc",
    report_sheet="Oladoc",
    source_specs=(SourceSpec("2022-2025 ", (3,)),),
    aliases={
        "Billed Bookings (0nline)": "Billed Bookings",
        "Total Cost of Services": "Total COS",
        "Salaries": "Payroll",
        "Marketing": "Marketing - Digital Advertising & Discounts",
        "Telco - VAS User Acquisition": "Marketing - SEO & Content",
        "Tech - Subscriptions": "Tech - Subscriptions",
        "G&A": "Admin + Overhead + Infrastructure",
        "Others": "Collection Loss (Bad Debt)",
    },
    row_rules={
        "Total Cost of Services": (("2022-2025 ", 28, 1),),
        "Salaries": (("2022-2025 ", 34, -1),),
        "Marketing": (("2022-2025 ", 35, -1), ("2022-2025 ", 37, -1)),
        "Telco - VAS User Acquisition": (("2022-2025 ", 36, -1),),
        "Tech - Subscriptions": (("2022-2025 ", 38, -1),),
        "G&A": (("2022-2025 ", 39, -1),),
        "Others": (("2022-2025 ", 40, -1),),
    },
    formula_rows=(22, 30, 31),
    round_values=True,
)

JIYE_PROFILE = LabelCopyProfile(
    name="jiye_technologies",
    display_name="Jiye Technologies",
    report_sheet="Jiye Technologies",
    source_specs=(SourceSpec("2. Historical Performance ", (1,)),),
    aliases={
        "Net Revenue": "Total Revenue",
    },
    row_rules={
        "Salaries": (("2. Historical Performance ", 122, -1),),
        "Tech": (("2. Historical Performance ", 124, 0),),
        "Office Space": (("2. Historical Performance ", 124, 0),),
        "Entertainment": (("2. Historical Performance ", 124, 0),),
        "Accounting & Legal": (("2. Historical Performance ", 124, 0),),
        "Travel": (("2. Historical Performance ", 124, 0),),
        "Others": (("2. Historical Performance ", 121, -1),),
    },
    date_rows=(41,),
    formula_rows=(15, 16, 18, 28, 29, 43, 44, 45, 46, 47, 48, 50, 51, 52, 53, 54, 55, 56),
    ignored_rows=(84, 89, 105, 107, 121, 122, 141, 143, 152, 178),
)

ONELOAD_PROFILE = LabelCopyProfile(
    name="oneload",
    display_name="OneLoad",
    report_sheet="OneLoad",
    source_specs=(SourceSpec("Monthly Tracker OPSPL", (2,)),),
    ignored_rows=(50, 75),
)

DOT_AND_LINE_PROFILE = LabelCopyProfile(
    name="dot_and_line",
    display_name="Dot & Line",
    report_sheet="Dot & Line",
    source_specs=(
        SourceSpec("Monthly Report ", (1,)),
        SourceSpec("Monthly Expenses", (1, 4, 7, 10, 13)),
    ),
    aliases={
        "Revenue from Product Sales (IGNITE)": "Revenue from Teacher Training Vertical (IGNITE)",
        "Active Students w/o short course students": "Active Students with Returning Student",
    },
    formula_rows=(15, 16, 17, 18, 19, 20, 22, 23, 24, 26, 27, 31, 32, 33, 34, 36, 37, 38, 39, 40, 64, 70, 75, 87, 89, 196, 197, 198, 199, 200, 201, 202, 203, 204, 205),
    ignored_rows=(66, 135, 160, 168, 179, 189),
)

PROFILES = {
    profile.name: profile
    for profile in (
        ABHI_PROFILE,
        SIMPAISA_PROFILE,
        BYKEA_PROFILE,
        DOT_AND_LINE_PROFILE,
        PROCHECK_PROFILE,
        ROOMY_PROFILE,
        OLADOC_PROFILE,
        REVOLVING_GAMES_PROFILE,
        TAPMAD_PROFILE,
        JIYE_PROFILE,
        ONELOAD_PROFILE,
    )
}

PENDING_COMPANIES = (
)

INACTIVE_COMPANIES = (
    "Dawaai",
    "Jugnu",
    "Patari",
    "Buyzilla",
    "Trukker",
)

FUTURE_COMPANIES = PENDING_COMPANIES + INACTIVE_COMPANIES


def parse_month(value: str) -> tuple[int, int]:
    try:
        parsed = datetime.strptime(value, "%Y-%m")
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"{value!r} is invalid. Use YYYY-MM, for example 2026-03."
        ) from exc
    return parsed.year, parsed.month


def is_target_month(value: object, year: int, month: int) -> bool:
    return isinstance(value, (datetime, date)) and value.year == year and value.month == month


def normalise_label(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    label = value.strip()
    return label or None


def _precision_safe_round(value: float, ndigits: int | None = None) -> int | float:
    """Round, but never collapse a nonzero value to zero.

    Ratio/percentage rows (Yield, Net MDR %, OneLoad ratios) sometimes sit in
    integer-formatted cells; rounding them to a whole number destroys the value.
    When rounding would produce 0 from a nonzero input, keep 6dp instead.
    """
    rounded = round(value, ndigits) if ndigits is not None else round(value)
    if rounded == 0 and value != 0:
        return round(value, 6)
    return rounded


def clean_number(value: object) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return _precision_safe_round(value)
    return None


def abhi_label_key(value: object) -> str | None:
    label = normalise_label(value)
    if label is None:
        return None

    label = label.replace("–", "-").replace("—", "-").replace("\n", " ")
    label = re.sub(r"(?<=[A-Za-z])0(?=[A-Za-z])", "-", label)
    return " ".join(label.lower().split())


def abhi_report_label(sheet: Worksheet, row: int) -> str | None:
    for col in (3, 2):
        label = abhi_label_key(sheet.cell(row=row, column=col).value)
        if label:
            return label
    return None


def _abhi_heading_matches(label: str | None, source_sheet: str) -> bool:
    if not label or "financial statements" not in label:
        return False
    aliases = ABHI_SECTION_HEADING_ALIASES.get(source_sheet, ())
    return any(alias in label for alias in aliases)


def abhi_section_start_row(report_sheet: Worksheet, source_sheet: str) -> int | None:
    base_start = ABHI_BASE_SECTION_START_ROWS.get(source_sheet)
    candidates: list[int] = []
    for row in range(1, report_sheet.max_row + 1):
        if _abhi_heading_matches(abhi_report_label(report_sheet, row), source_sheet):
            candidates.append(row)
    if not candidates:
        return base_start
    if base_start is None:
        return candidates[-1]
    return min(candidates, key=lambda row: abs(row - base_start))


def resolve_abhi_report_sections(
    report_sheet: Worksheet,
    sections: tuple[SectionMap, ...],
) -> tuple[SectionMap, ...]:
    resolved: list[SectionMap] = []
    starts: dict[str, int] = {}
    for section in sections:
        starts[section.source_sheet] = abhi_section_start_row(report_sheet, section.source_sheet) or section.start_row

    ordered = sorted(sections, key=lambda section: starts[section.source_sheet])
    for index, section in enumerate(ordered):
        start_row = starts[section.source_sheet]
        projected_end = start_row + (section.end_row - section.start_row)
        next_start = starts[ordered[index + 1].source_sheet] if index + 1 < len(ordered) else None
        end_row = min(projected_end, report_sheet.max_row)
        if next_start is not None:
            end_row = min(end_row, next_start - 1)
        if end_row < start_row:
            end_row = section.end_row
        resolved.append(SectionMap(start_row, end_row, section.source_sheet))

    return tuple(resolved)


def abhi_rows_by_label(
    report_sheet: Worksheet,
    label: str,
    start_row: int,
    end_row: int,
) -> list[int]:
    wanted = abhi_label_key(label)
    rows: list[int] = []
    for row in range(start_row, min(end_row, report_sheet.max_row) + 1):
        if abhi_report_label(report_sheet, row) == wanted:
            rows.append(row)
    return rows


def abhi_first_row_by_label(
    report_sheet: Worksheet,
    label: str,
    start_row: int,
    end_row: int,
) -> int | None:
    rows = abhi_rows_by_label(report_sheet, label, start_row, end_row)
    return rows[0] if rows else None


def clean_abhi_number(
    value: object,
    report_label: str | None,
    number_format: str | None,
) -> int | float | None:
    if isinstance(value, bool):
        return None
    if not isinstance(value, (int, float)):
        return None

    if report_label == "__round_whole__":
        return _precision_safe_round(value)

    label = report_label or ""
    if "check" in label:
        # Check rows hold rounding residue; the team always shows 0.
        return round(value)
    if "capital adequacy ratio" in label:
        return _precision_safe_round(value, 4)

    decimal_label_keywords = (
        "average revenue / transaction",
        "average transaction size",
        "number of transactions/user",
    )
    if any(keyword in label for keyword in decimal_label_keywords):
        return _precision_safe_round(value, 2)

    number_format = number_format or "General"
    if "%" in number_format:
        decimal_places = 0
        if "." in number_format:
            decimal_places = len(number_format.split("%", 1)[0].rsplit(".", 1)[-1])
        return _precision_safe_round(value, decimal_places + 2)

    if "." in number_format and "0" in number_format.rsplit(".", 1)[-1]:
        decimal_places = number_format.rsplit(".", 1)[-1].count("0")
        return _precision_safe_round(value, decimal_places)

    if number_format != "General":
        return _precision_safe_round(value)

    decimal_keywords = (
        "average",
        "avg",
        "rating",
        "transactions/user",
        "days outstanding",
    )
    if abs(value) < 1 or any(keyword in label for keyword in decimal_keywords):
        return _precision_safe_round(value, 2)
    return _precision_safe_round(value)


def apply_abhi_display_convention(value: object, row: int) -> object:
    if row in ABHI_FORCE_BLANK_ROWS:
        return None
    if row in ABHI_ZERO_AS_DASH_ROWS and value is None:
        return "-"
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value == 0:
        if row in ABHI_ZERO_AS_DASH_ROWS:
            return "-"
        if row in ABHI_ZERO_AS_BLANK_ROWS:
            return None
    return value


def find_month_columns(
    sheet: Worksheet,
    months: Iterable[tuple[int, int]],
    header_search_rows: int = 25,
) -> dict[tuple[int, int], int]:
    wanted = set(months)
    candidates: dict[tuple[int, int], tuple[int, int]] = {}

    def explicit_years(value: object) -> set[int]:
        if isinstance(value, (datetime, date)):
            return {value.year}
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            year = int(value)
            return {year} if 1900 <= year <= 2100 else set()
        if isinstance(value, str):
            return {int(match) for match in re.findall(r"\b(19\d{2}|20\d{2}|21\d{2})\b", value)}
        return set()

    def has_conflicting_stacked_year(row: int, col: int, year: int) -> bool:
        nearby_years: set[int] = set()
        for nearby_row in range(max(1, row - 2), row):
            nearby_years.update(explicit_years(sheet.cell(row=nearby_row, column=col).value))
        return bool(nearby_years) and year not in nearby_years

    def date_header_display_year(value: object, number_format: str | None) -> int | None:
        if not isinstance(value, (datetime, date)):
            return None
        fmt = (number_format or "").lower().replace("\\", "")
        # Some Monthly Reporting headers display "Jun-25" using an Excel date
        # stored as 2025-06-25 with format mmm-d. In that convention, the day
        # is the visible two-digit year, not the day of the month.
        if "y" not in fmt and "m" in fmt and "d" in fmt and 20 <= value.day <= 40:
            return 2000 + value.day
        return None

    def score_month_header(value: object, year: int, month: int, number_format: str | None = None) -> int:
        if not isinstance(value, (datetime, date)) or value.month != month:
            return 0

        display_year = date_header_display_year(value, number_format)
        if display_year is not None:
            return 120 if display_year == year else 0

        score = 0
        if value.year == year:
            score += 100
        # The Monthly Reporting workbook sometimes stores labels like Oct-21 as
        # 2025-10-21. Use that as a fallback, but never above a real year match.
        if value.year != year and value.day == year % 100:
            score += 30
        return score

    for row in range(1, min(header_search_rows, sheet.max_row) + 1):
        for col in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row, column=col)
            value = cell.value
            for year, month in wanted:
                key = (year, month)
                score = score_month_header(value, year, month, cell.number_format)
                if score and isinstance(value, (datetime, date)) and has_conflicting_stacked_year(row, col, year):
                    score = 0
                if score == 0:
                    continue
                if key not in candidates or score > candidates[key][0]:
                    candidates[key] = (score, col)

    return {month_key: col for month_key, (_score, col) in candidates.items()}


def _month_aliases(month: int) -> set[str]:
    short = datetime(2000, month, 1).strftime("%b").lower()
    long = datetime(2000, month, 1).strftime("%B").lower()
    return {short, long, short.rstrip(".")}


def _string_mentions_month(value: object, month_key: tuple[int, int]) -> bool:
    if not isinstance(value, str):
        return False
    year, month = month_key
    text = value.strip().lower().replace("'", "")
    if not text:
        return False
    year_tokens = {str(year), str(year)[-2:]}
    if not any(token in text for token in year_tokens):
        return False
    return any(alias in text for alias in _month_aliases(month))


def workbook_sheet_name(workbook, sheet_name: str) -> str | None:
    if sheet_name in workbook.sheetnames:
        return sheet_name

    wanted = " ".join(sheet_name.strip().lower().split())
    for candidate in workbook.sheetnames:
        if " ".join(candidate.strip().lower().split()) == wanted:
            return candidate

    wanted_compact = re.sub(r"[^a-z0-9]", "", sheet_name.lower())
    for candidate in workbook.sheetnames:
        candidate_compact = re.sub(r"[^a-z0-9]", "", candidate.lower())
        if candidate_compact == wanted_compact:
            return candidate

    wanted_fy_family = re.sub(r"fy[- ]?\d{2,4}", "fy", wanted)
    for candidate in workbook.sheetnames:
        candidate_normalised = " ".join(candidate.strip().lower().split())
        candidate_fy_family = re.sub(r"fy[- ]?\d{2,4}", "fy", candidate_normalised)
        if wanted_fy_family == candidate_fy_family and "fy" in wanted_fy_family:
            return candidate

    wanted_year_range = re.search(r"(\d{4})\s*[-–]\s*(\d{4})", sheet_name)
    if wanted_year_range:
        wanted_start_year = wanted_year_range.group(1)
        for candidate in workbook.sheetnames:
            candidate_year_range = re.search(r"(\d{4})\s*[-–]\s*(\d{4})", candidate)
            if candidate_year_range and candidate_year_range.group(1) == wanted_start_year:
                return candidate
    return None


def workbook_sheet(workbook, sheet_name: str) -> Worksheet | None:
    actual_name = workbook_sheet_name(workbook, sheet_name)
    return workbook[actual_name] if actual_name is not None else None


def profile_workbook_sheet_name(workbook, profile_name: str, expected_sheet_name: str) -> str | None:
    actual_name = workbook_sheet_name(workbook, expected_sheet_name)
    if actual_name is not None:
        return actual_name

    for alias in BUILT_IN_SOURCE_SHEET_ALIASES.get(profile_name, {}).get(expected_sheet_name, ()):
        alias_name = workbook_sheet_name(workbook, alias)
        if alias_name is not None:
            return alias_name

    sheet_aliases = memory_sheet_aliases_for_profile(profile_name)
    expected_key = normalise_memory_key(expected_sheet_name)
    memory_sheet_name = sheet_aliases.get(expected_key or "")
    if not memory_sheet_name:
        return None
    return workbook_sheet_name(workbook, memory_sheet_name)


def profile_workbook_sheet(workbook, profile_name: str, expected_sheet_name: str) -> Worksheet | None:
    actual_name = profile_workbook_sheet_name(workbook, profile_name, expected_sheet_name)
    return workbook[actual_name] if actual_name is not None else None


def _value_mentions_year(value: object, year: int) -> bool:
    if isinstance(value, (datetime, date)):
        return value.year == year
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value) == year
    if isinstance(value, str):
        text = value.strip().replace("'", "")
        return str(year) in text or str(year)[-2:] in text
    return False


def find_flexible_month_columns(
    sheet: Worksheet,
    months: Iterable[tuple[int, int]],
    header_search_rows: int = 25,
    max_header_columns: int = 300,
) -> dict[tuple[int, int], int]:
    """Find month columns across date cells and common text header layouts."""
    wanted = set(months)
    found = find_month_columns(sheet, wanted, header_search_rows=header_search_rows)

    for row in range(1, min(header_search_rows, sheet.max_row) + 1):
        for col in range(1, min(sheet.max_column, max_header_columns) + 1):
            value = sheet.cell(row=row, column=col).value
            for year, month in wanted:
                key = (year, month)
                if key in found:
                    continue
                if _string_mentions_month(value, key):
                    found[key] = col
                    continue
                if isinstance(value, str) and value.strip().lower() in _month_aliases(month):
                    nearby_values = [
                        sheet.cell(row=nearby_row, column=col).value
                        for nearby_row in range(max(1, row - 2), min(sheet.max_row, row + 2) + 1)
                        if nearby_row != row
                    ]
                    if any(_value_mentions_year(nearby_value, year) for nearby_value in nearby_values):
                        found[key] = col

    return found


def _sheet_name_mentions_month(sheet_name: str, month_key: tuple[int, int]) -> bool:
    year, month = month_key
    text = sheet_name.strip().lower().replace("'", "")
    return any(alias in text for alias in _month_aliases(month)) and (
        str(year) in text or str(year)[-2:] in text
    )


def find_workbook_month_sheet(workbook, month_key: tuple[int, int]) -> Worksheet | None:
    for sheet_name in workbook.sheetnames:
        if _sheet_name_mentions_month(sheet_name, month_key):
            return workbook[sheet_name]
    return None


def _profile_sheet_candidates(profile_name: str, profile: CompanyProfile | LabelCopyProfile) -> tuple[str, ...]:
    candidates: list[str] = []
    if isinstance(profile, CompanyProfile):
        candidates.extend(section.source_sheet for section in profile.sections)
    else:
        candidates.extend(spec.sheet_name for spec in profile.source_specs)

    legacy_candidates = {
        "abhi": (
            "Abhi Private - Pakistan (USD)",
            "AbhiPayriff - Pakistan (USD)",
            "Abhi Fintech - UAE",
            "Consolidated",
        ),
        "bykea": ("Details",),
        "dot_and_line": ("Monthly Report - new template", "Monthly Expenses "),
        "jiye_technologies": ("Jan-24", "Feb-24", "Mar-24"),
        "oladoc": ("2022-2025", "2022-2024"),
        "oneload": ("Monthly Tracker EPS-BV", "Tracker EPS-BV USD"),
        "roomy": ("Oct-25", "Nov-25", "Dec-25"),
        "simpaisa": ("HPR Format FY-25", "HPR Format FY-24"),
    }
    candidates.extend(legacy_candidates.get(profile_name, ()))
    return tuple(dict.fromkeys(candidates))


def _profile_expected_label_keys(profile: CompanyProfile | LabelCopyProfile) -> set[str]:
    if isinstance(profile, CompanyProfile):
        return set()

    expected: set[str] = set()
    if profile.aliases:
        expected.update(label_key(value) for value in profile.aliases.values())
    if profile.row_rules:
        expected.update(label_key(value) for value in profile.row_rules)
    if profile.fixed_month_sheet_rules:
        expected.update(label_key(value) for value in profile.fixed_month_sheet_rules)
    memory_aliases = memory_aliases_for_profile(profile.name)
    expected.update(memory_aliases.keys())
    expected.update(memory_aliases.values())
    expected.discard(None)
    return {value for value in expected if value}


def _available_label_keys(workbook, sheet_names: Iterable[str], max_rows: int = 350) -> set[str]:
    labels: set[str] = set()
    for sheet_name in sheet_names:
        sheet = workbook_sheet(workbook, sheet_name)
        if sheet is None:
            continue
        for row in range(1, min(sheet.max_row, max_rows) + 1):
            for col in range(1, min(sheet.max_column, 6) + 1):
                key = label_key(sheet.cell(row=row, column=col).value)
                if key:
                    labels.add(key)
    return labels


def _sheet_name_hint_score(sheet_name: str, expected_name: str) -> int:
    sheet_tokens = set(re.findall(r"[a-z0-9]+", sheet_name.lower()))
    expected_tokens = set(re.findall(r"[a-z0-9]+", expected_name.lower()))
    if not sheet_tokens or not expected_tokens:
        return 0
    overlap = len(sheet_tokens & expected_tokens)
    if overlap == 0:
        return 0
    return min(5, overlap)


def infer_profile_source_sheet_name(
    workbook,
    profile: LabelCopyProfile,
    expected_sheet_name: str,
    months: tuple[tuple[int, int], ...],
) -> str | None:
    exact_name = profile_workbook_sheet_name(workbook, profile.name, expected_sheet_name)
    if exact_name is not None:
        return exact_name

    expected_labels = _profile_expected_label_keys(profile)
    best_sheet: str | None = None
    best_score = 0

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        month_score = len(find_flexible_month_columns(sheet, months))
        labels = _available_label_keys(workbook, (sheet_name,))
        label_score = len(expected_labels & labels) if expected_labels else 0
        name_score = _sheet_name_hint_score(sheet_name, expected_sheet_name)

        if profile.name == "simpaisa" and "hpr format" in sheet_name.lower():
            name_score += 8
        if profile.name == "oladoc" and re.search(r"2022\s*[-–]\s*20\d{2}", sheet_name):
            name_score += 8

        score = month_score * 5 + label_score * 2 + name_score
        if score > best_score:
            best_score = score
            best_sheet = sheet_name

    return best_sheet if best_score >= 8 else None


def infer_profile_source_sheet(
    workbook,
    profile: LabelCopyProfile,
    expected_sheet_name: str,
    months: tuple[tuple[int, int], ...],
) -> Worksheet | None:
    sheet_name = infer_profile_source_sheet_name(workbook, profile, expected_sheet_name, months)
    return workbook[sheet_name] if sheet_name is not None else None


def validate_template_file(
    template_path: Path,
    months: tuple[tuple[int, int], ...],
) -> ValidationResult:
    issues: list[str] = []
    warnings: list[str] = []

    try:
        workbook = load_workbook(template_path, data_only=True, read_only=False)
    except Exception as exc:
        return ValidationResult(
            profile="template",
            company="Monthly Reporting Template",
            status="Blocked",
            confidence=0,
            issues=(f"Could not open workbook: {exc}",),
        )

    try:
        missing_sheets = [
            profile.report_sheet
            for profile in PROFILES.values()
            if profile.report_sheet not in workbook.sheetnames
        ]
        if missing_sheets:
            issues.append("Missing report tabs: " + ", ".join(missing_sheets[:8]))

        missing_month_tabs: list[str] = []
        for profile in PROFILES.values():
            if profile.report_sheet not in workbook.sheetnames:
                continue
            sheet = workbook[profile.report_sheet]
            found = find_flexible_month_columns(sheet, months)
            missing = [month for month in months if month not in found]
            if missing:
                formatted = ", ".join(f"{year}-{month:02d}" for year, month in missing)
                missing_month_tabs.append(f"{profile.display_name}: {formatted}")
        if missing_month_tabs:
            issues.append("Missing selected month columns in template: " + "; ".join(missing_month_tabs[:6]))
    finally:
        workbook.close()

    confidence = 100
    if issues:
        confidence = 0
    elif warnings:
        confidence = 75
    return ValidationResult(
        profile="template",
        company="Monthly Reporting Template",
        status="Blocked" if issues else "Ready",
        confidence=confidence,
        issues=tuple(issues),
        warnings=tuple(warnings),
    )


def validate_kpi_file(
    profile_name: str,
    kpi_path: Path,
    months: tuple[tuple[int, int], ...],
) -> ValidationResult:
    profile = PROFILES[profile_name]
    issues: list[str] = []
    warnings: list[str] = []

    try:
        workbook = load_workbook(kpi_path, data_only=True, read_only=False)
    except Exception as exc:
        return ValidationResult(
            profile=profile_name,
            company=profile.display_name,
            status="Blocked",
            confidence=0,
            issues=(f"Could not open workbook: {exc}",),
        )

    try:
        candidates = _profile_sheet_candidates(profile_name, profile)
        matched_sheets = [
            actual_sheet_name
            for sheet_name in candidates
            if (
                actual_sheet_name := (
                    infer_profile_source_sheet_name(workbook, profile, sheet_name, months)
                    if isinstance(profile, LabelCopyProfile)
                    else workbook_sheet_name(workbook, sheet_name)
                )
            ) is not None
        ]
        if profile_name == "roomy":
            roomy_month_sheets = [
                sheet_name
                for sheet_name in workbook.sheetnames
                if any(_sheet_name_mentions_month(sheet_name, month_key) for month_key in months)
            ]
            for sheet_name in roomy_month_sheets:
                if sheet_name not in matched_sheets:
                    matched_sheets.append(sheet_name)
        if not matched_sheets:
            issues.append("Expected sheet not found. Found sheets: " + ", ".join(workbook.sheetnames[:8]))

        found_months: set[tuple[int, int]] = set()
        for sheet_name in matched_sheets:
            sheet = workbook[sheet_name]
            found_months.update(find_flexible_month_columns(sheet, months).keys())
            for month_key in months:
                if _sheet_name_mentions_month(sheet_name, month_key):
                    found_months.add(month_key)

        missing_months = [month for month in months if month not in found_months]
        if missing_months:
            formatted = ", ".join(f"{year}-{month:02d}" for year, month in missing_months)
            if profile_name == "roomy":
                issues.append(f"Roomy workbook is missing selected quarter months: {formatted}")
            elif matched_sheets:
                warnings.append(f"Could not confirm month columns for: {formatted}")
            else:
                issues.append(f"Could not confirm month columns for: {formatted}")

        expected_labels = _profile_expected_label_keys(profile)
        if expected_labels and matched_sheets:
            available_labels = _available_label_keys(workbook, matched_sheets)
            matched_label_count = len(expected_labels & available_labels)
            if matched_label_count == 0:
                warnings.append("Expected KPI labels were not recognized.")

        if profile_name == "abhi" and missing_months:
            warnings.append("ABHI often needs separate monthly files after restructuring; verify missing months manually.")
    finally:
        workbook.close()

    if issues:
        status = "Blocked"
        confidence = 0
    elif warnings:
        status = "Warning"
        confidence = max(35, 100 - 15 * len(warnings))
    else:
        status = "Ready"
        confidence = 100

    return ValidationResult(
        profile=profile_name,
        company=profile.display_name,
        status=status,
        confidence=confidence,
        issues=tuple(issues),
        warnings=tuple(warnings),
    )


def build_source_label_map(sheet: Worksheet) -> dict[str, int]:
    labels: dict[str, int] = {}

    for row in range(1, sheet.max_row + 1):
        label = normalise_label(sheet.cell(row=row, column=3).value)
        if label is None:
            label = normalise_label(sheet.cell(row=row, column=2).value)
        if label is not None and label not in labels:
            labels[label] = row

    return labels


def build_source_label_rows(sheet: Worksheet) -> dict[str, list[int]]:
    labels: dict[str, list[int]] = {}

    for row in range(1, sheet.max_row + 1):
        label = normalise_label(sheet.cell(row=row, column=3).value)
        if label is None:
            label = normalise_label(sheet.cell(row=row, column=2).value)
        if label is not None:
            labels.setdefault(label, []).append(row)

    return labels


def label_key(value: object) -> str | None:
    label = normalise_label(value)
    if label is None:
        return None
    return " ".join(label.lower().replace("\n", " ").split())


def first_report_label(sheet: Worksheet, row: int) -> str | None:
    for col in (1, 2, 3, 4):
        label = normalise_label(sheet.cell(row=row, column=col).value)
        if label is not None:
            return label
    return None


def build_label_value_map(
    workbook,
    profile: LabelCopyProfile,
    months: tuple[tuple[int, int], ...],
) -> dict[str, dict[tuple[int, int], object]]:
    source_values: dict[str, dict[tuple[int, int], object]] = {}

    if profile.name == "roomy":
        for month_key in months:
            sheet = find_workbook_month_sheet(workbook, month_key)
            if sheet is None:
                continue

            for row in range(1, sheet.max_row + 1):
                labels = [
                    label_key(sheet.cell(row=row, column=3).value),
                    label_key(sheet.cell(row=row, column=4).value),
                ]
                labels = [label for label in labels if label]
                if not labels:
                    continue

                value = sheet.cell(row=row, column=7).value
                if value in (None, ""):
                    continue

                for label in labels:
                    source_values.setdefault(label, {})
                    source_values[label].setdefault(month_key, value)

        if source_values:
            return source_values

    for spec in profile.source_specs:
        sheet = infer_profile_source_sheet(workbook, profile, spec.sheet_name, months)
        if sheet is None:
            continue

        month_cols: dict[tuple[int, int], int] = {}

        if spec.month_value_columns:
            for year, month in months:
                if month in spec.month_value_columns:
                    month_cols[(year, month)] = spec.month_value_columns[month]
        elif spec.sheet_month and spec.value_column:
            for year, month in months:
                if month == spec.sheet_month:
                    month_cols[(year, month)] = spec.value_column
        else:
            month_cols = find_flexible_month_columns(sheet, months)

        for row in range(1, sheet.max_row + 1):
            labels = []
            for col in spec.label_columns:
                labels.append(label_key(sheet.cell(row=row, column=col).value))

            labels = [label for label in labels if label]
            if not labels:
                continue

            for month_key, col in month_cols.items():
                value = sheet.cell(row=row, column=col).value
                if value in (None, ""):
                    continue
                for label in labels:
                    source_values.setdefault(label, {})
                    source_values[label].setdefault(month_key, value)

    return source_values


def source_months_with_data(
    workbook,
    profile: LabelCopyProfile,
    months: tuple[tuple[int, int], ...],
    source_values: Mapping[str, Mapping[tuple[int, int], object]],
) -> set[tuple[int, int]]:
    found: set[tuple[int, int]] = set()

    for month_values in source_values.values():
        found.update(month_values.keys())

    if profile.name == "roomy":
        for month_key in months:
            if find_workbook_month_sheet(workbook, month_key) is not None:
                found.add(month_key)

    for month_key in months:
        if month_key in found:
            continue

        if profile.name == "bykea":
            for metric_key in BYKEA_VISIBLE_SUMMARY_LABELS:
                if get_bykea_visible_summary_source(workbook, metric_key, month_key) is not None:
                    found.add(month_key)
                    break

        if month_key in found:
            continue

        for rules in (profile.row_rules or {}).values():
            value = get_row_rule_value(workbook, rules, month_key, profile.name)
            if value is not None:
                found.add(month_key)
                break

        if month_key in found:
            continue

        for rule in (profile.fixed_month_sheet_rules or {}).values():
            value = get_fixed_month_sheet_value(workbook, profile, rule, month_key)
            if value not in (None, "") and not isinstance(value, bool):
                found.add(month_key)
                break

    return found


def source_month_column(sheet: Worksheet, month_key: tuple[int, int]) -> int | None:
    return find_flexible_month_columns(sheet, (month_key,)).get(month_key)


def _bykea_summary_row_score(sheet: Worksheet, row: int, value: object) -> tuple[int, int]:
    row_dimension = sheet.row_dimensions[row]
    score = 0
    if not row_dimension.hidden:
        score += 1000
    if row_dimension.outlineLevel == 0:
        score += 500
    if row_dimension.collapsed:
        score += 100
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value != 0:
        score += 50
    return score, row


def find_bykea_visible_summary_row(
    sheet: Worksheet,
    metric_key: str | None,
    value_column: int,
) -> int | None:
    labels = BYKEA_VISIBLE_SUMMARY_LABELS.get(metric_key or "")
    return find_bykea_visible_summary_row_for_labels(sheet, labels, value_column)


def find_bykea_visible_summary_row_for_labels(
    sheet: Worksheet,
    labels: tuple[str, ...],
    value_column: int,
) -> int | None:
    if not labels:
        return None

    label_keys = {label_key(label) for label in labels}
    label_keys.discard(None)
    candidates: list[tuple[tuple[int, int], int]] = []
    for row in range(1, sheet.max_row + 1):
        row_keys = {
            label_key(sheet.cell(row=row, column=col).value)
            for col in range(1, min(sheet.max_column, 6) + 1)
        }
        row_keys.discard(None)
        if not row_keys.intersection(label_keys):
            continue
        value = sheet.cell(row=row, column=value_column).value
        if value in (None, "") or isinstance(value, bool):
            continue
        candidates.append((_bykea_summary_row_score(sheet, row, value), row))

    if not candidates:
        return None
    return max(candidates)[1]


def get_bykea_visible_summary_sources(
    workbook,
    metric_key: str | None,
    month_key: tuple[int, int],
) -> list[tuple[Worksheet, int, int, object]]:
    sheet = profile_workbook_sheet(workbook, "bykea", "Plan")
    if sheet is None:
        return []
    source_col = source_month_column(sheet, month_key)
    if source_col is None:
        return []

    sources: list[tuple[Worksheet, int, int, object]] = []
    composite_labels = BYKEA_COMPOSITE_SUMMARY_LABELS.get(metric_key or "")
    if composite_labels:
        for labels in composite_labels:
            source_row = find_bykea_visible_summary_row_for_labels(sheet, labels, source_col)
            if source_row is None:
                return []
            sources.append(
                (
                    sheet,
                    source_row,
                    source_col,
                    sheet.cell(row=source_row, column=source_col).value,
                )
            )
        return sources

    source_row = find_bykea_visible_summary_row(sheet, metric_key, source_col)
    if source_row is None:
        return []
    sources.append((sheet, source_row, source_col, sheet.cell(row=source_row, column=source_col).value))
    return sources


def get_bykea_visible_summary_source(
    workbook,
    metric_key: str | None,
    month_key: tuple[int, int],
) -> tuple[Worksheet, int, int, object] | None:
    sources = get_bykea_visible_summary_sources(workbook, metric_key, month_key)
    if not sources:
        return None

    sheet, source_row, source_col, fallback_value = sources[-1]
    total = 0.0
    found = False
    for _sheet, _row, _col, value in sources:
        if isinstance(value, bool):
            continue
        if isinstance(value, (int, float)):
            total += value
            found = True
    return sheet, source_row, source_col, total if found else fallback_value



def get_row_rule_value(
    workbook,
    rule: tuple[tuple[str, int, float], ...],
    month_key: tuple[int, int],
    profile_name: str | None = None,
) -> float | None:
    total = 0.0
    found_value = False
    for sheet_name, row, sign in rule:
        sheet = (
            profile_workbook_sheet(workbook, profile_name, sheet_name)
            if profile_name is not None
            else workbook_sheet(workbook, sheet_name)
        )
        if sheet is None:
            continue
        col = source_month_column(sheet, month_key)
        if col is None:
            continue
        value = sheet.cell(row=row, column=col).value
        if isinstance(value, bool):
            continue
        if isinstance(value, (int, float)):
            found_value = True
            total += value * sign
    return total if found_value else None


def get_fixed_month_sheet_value(
    workbook,
    profile: LabelCopyProfile,
    rule: tuple[int, int, float],
    month_key: tuple[int, int],
) -> object | None:
    _, month = month_key
    source_sheet = None
    if profile.name == "roomy":
        source_sheet = find_workbook_month_sheet(workbook, month_key)
    else:
        for spec in profile.source_specs:
            if spec.sheet_month == month:
                source_sheet = infer_profile_source_sheet(workbook, profile, spec.sheet_name, (month_key,))
            if source_sheet is not None:
                break
    if source_sheet is None:
        return None

    row, col, sign = rule
    value = source_sheet.cell(row=row, column=col).value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value * sign
    return value


def get_month_sheet(
    workbook,
    profile: LabelCopyProfile,
    month_key: tuple[int, int],
) -> Worksheet | None:
    _, month = month_key
    if profile.name == "roomy":
        source_sheet = find_workbook_month_sheet(workbook, month_key)
        if source_sheet is not None:
            return source_sheet

    for spec in profile.source_specs:
        if spec.sheet_month == month:
            sheet = infer_profile_source_sheet(workbook, profile, spec.sheet_name, (month_key,))
            if sheet is not None:
                return sheet
    return None


def find_sheet_label_value(
    sheet: Worksheet,
    label: str,
    value_column: int,
    occurrence: int = 1,
) -> object | None:
    wanted = label_key(label)
    if wanted is None:
        return None

    seen = 0
    for row in range(1, sheet.max_row + 1):
        row_keys = (
            label_key(sheet.cell(row=row, column=3).value),
            label_key(sheet.cell(row=row, column=4).value),
        )
        if wanted not in row_keys:
            continue
        seen += 1
        if seen == occurrence:
            return sheet.cell(row=row, column=value_column).value
    return None


def find_roomy_fx_label(sheet: Worksheet) -> object | None:
    for row in range(1, min(sheet.max_row, 8) + 1):
        value = sheet.cell(row=row, column=7).value
        if isinstance(value, str) and "usd @ pkr" in value.lower():
            return value
    return None


def translate_formula(source_formula: str, source_cell: str, target_cell: str) -> str:
    try:
        return Translator(source_formula, origin=source_cell).translate_formula(target_cell)
    except Exception:
        return source_formula


def evaluate_simple_formula(
    sheet: Worksheet,
    formula: str,
    current_col: int | None = None,
    seen: set[str] | None = None,
) -> float | int | None:
    seen = seen or set()
    expression = formula[1:].strip()

    def cell_value(reference: str) -> float:
        if reference in seen:
            return 0.0
        seen.add(reference)
        match = re.fullmatch(r"([A-Z]+)([0-9]+)", reference)
        if not match:
            return 0.0
        col = column_index_from_string(match.group(1))
        row = int(match.group(2))
        value = sheet.cell(row=row, column=col).value
        numeric = numeric_cell_value(sheet, value, col, seen)
        return float(numeric or 0)

    def sum_range(match: re.Match[str]) -> str:
        range_text = match.group(1)
        try:
            min_col, min_row, max_col, max_row = range_boundaries(range_text)
        except ValueError:
            return "0"
        total = 0.0
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                value = sheet.cell(row=row, column=col).value
                total += float(numeric_cell_value(sheet, value, col, seen) or 0)
        return str(total)

    expression = re.sub(r"SUM\(([A-Z]+[0-9]+:[A-Z]+[0-9]+)\)", sum_range, expression, flags=re.IGNORECASE)
    expression = re.sub(r"(?<![A-Z])([A-Z]+[0-9]+)", lambda match: str(cell_value(match.group(1))), expression)
    if not re.fullmatch(r"[0-9eE\.\+\-\*/\(\) ]+", expression):
        return None
    try:
        result = eval(expression, {"__builtins__": {}}, {})
    except Exception:
        return None
    if isinstance(result, (int, float)):
        return result
    return None


def numeric_cell_value(
    sheet: Worksheet,
    value: object,
    current_col: int | None = None,
    seen: set[str] | None = None,
) -> float | int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str) and value.startswith("="):
        return evaluate_simple_formula(sheet, value, current_col=current_col, seen=seen)
    return None


def write_if_number(sheet: Worksheet, row: int, col: int, value: object) -> bool:
    number = numeric_cell_value(sheet, value, current_col=col)
    if number is None:
        return False
    sheet.cell(row=row, column=col).value = round(number, 6) if isinstance(number, float) else number
    return True


def fill_abhi_visible_summary(
    report_sheet: Worksheet,
    target_month_cols: Mapping[tuple[int, int], int],
    months: tuple[tuple[int, int], ...],
) -> dict[str, int]:
    values_written = 0

    def write_top_summary_formula(row: int, col: int, value: object) -> None:
        nonlocal values_written
        report_sheet.cell(row=row, column=col).value = value
        values_written += 1

    def fill_group_consolidated_summary() -> bool:
        consolidated_start = abhi_section_start_row(report_sheet, "Consolidated")
        if consolidated_start is None:
            return False

        marker = abhi_report_label(report_sheet, consolidated_start)
        if not _abhi_heading_matches(marker, "Consolidated"):
            return False

        consolidated_end = min(
            report_sheet.max_row,
            consolidated_start
            + (
                ABHI_PROFILE.sections[-1].end_row
                - ABHI_PROFILE.sections[-1].start_row
            ),
        )

        required_labels = {
            "total_processed_value": "Total Processed Value",
            "total_transactions": "Total Transactions (EWA)",
            "net_revenue": "Net Revenue",
            "cost_of_services": "Cost of Services",
            "other_income": "Other Income",
            "staff": "Staff",
            "marketing": "Marketing",
            "sga": "SG&A",
            "occupancy": "Occupancy",
            "finance": "Finance",
            "other": "Other",
            "capex": "Depreciation & Amortization",
        }
        label_rows = {
            key: abhi_first_row_by_label(report_sheet, label, consolidated_start, consolidated_end)
            for key, label in required_labels.items()
        }
        if any(row is None for row in label_rows.values()):
            return False

        detail_start = ABHI_BASE_SECTION_START_ROWS["Abhi Private Ltd"]
        active_company_rows = abhi_rows_by_label(
            report_sheet,
            "Active Companies",
            detail_start,
            consolidated_start - 1,
        )
        payroll_rows = abhi_rows_by_label(
            report_sheet,
            "Payrolls Processed",
            detail_start,
            consolidated_start - 1,
        )

        def ref(col_letter: str, key: str) -> str:
            return f"{col_letter}{label_rows[key]}"

        def sum_refs(col_letter: str, rows: list[int]) -> str:
            if not rows:
                return "0"
            return "+".join(f"{col_letter}{row}" for row in rows)

        for month_key in months:
            col = target_month_cols.get(month_key)
            if col is None or col <= 2:
                continue

            actual = get_column_letter(col)
            forecast = get_column_letter(col + 1)
            previous = get_column_letter(col - 2)

            actual_formulas: dict[int, object] = {
                14: f"={ref(actual, 'total_processed_value')}",
                15: f"={ref(actual, 'net_revenue')}",
                16: f"={ref(actual, 'cost_of_services')}",
                17: f"={actual}15+{actual}16",
                18: f"={ref(actual, 'other_income')}",
                19: f"={ref(actual, 'staff')}",
                20: 0.0,
                21: f"={ref(actual, 'marketing')}",
                22: f"={ref(actual, 'sga')}",
                23: f"={ref(actual, 'occupancy')}+{ref(actual, 'finance')}+{ref(actual, 'other')}",
                24: f"=SUM({actual}19:{actual}23)",
                25: f"={ref(actual, 'capex')}",
                26: f"={actual}17+{actual}18+{actual}24+{actual}25",
                29: f"={sum_refs(actual, active_company_rows)}",
                30: f"={ref(actual, 'total_transactions')}+{sum_refs(actual, payroll_rows)}",
                32: f"={actual}30/{previous}30 -1",
                33: f"={actual}15/{actual}14",
                34: f"={actual}14/{actual}30",
                35: f"={actual}14/{actual}29",
                36: f"={actual}30/{actual}29",
                37: f"={actual}14/{previous}14 -1",
                38: f"={actual}29/{previous}29 -1",
                39: f"={actual}34/{previous}34",
                40: f"=(-{actual}26/{actual}14)",
            }
            forecast_formulas: dict[int, object] = {
                14: f"={previous}14*1.05",
                15: f"={forecast}14*(48%/12)",
                16: f"=-{forecast}14*2.2%",
                17: f"={forecast}15+{forecast}16",
                18: f"={previous}18*0.95",
                19: -249999.0,
                20: 0.0,
                21: -9999.0,
                22: -249999.0,
                23: f"={previous}23*0.95",
                24: f"=SUM({forecast}19:{forecast}23)",
                25: -4999.0,
                26: f"={forecast}17+{forecast}24+{forecast}18+{forecast}25",
                28: 17.0,
                29: 322.0,
                30: f"={previous}30*1.1",
            }

            for row, value in actual_formulas.items():
                write_top_summary_formula(row, col, value)
            for row, value in forecast_formulas.items():
                write_top_summary_formula(row, col + 1, value)
            for row in range(32, 41):
                report_sheet.cell(row=row, column=col + 1).value = None

        return True

    if fill_group_consolidated_summary():
        return {"abhi_summary_values_written": values_written}

    def number(row: int, col: int) -> float | int | None:
        return numeric_cell_value(report_sheet, report_sheet.cell(row=row, column=col).value, current_col=col)

    def write(row: int, col: int, value: object) -> None:
        nonlocal values_written
        if write_if_number(report_sheet, row, col, value):
            values_written += 1

    for month_key in months:
        col = target_month_cols.get(month_key)
        if col is None:
            continue

        # Fill visible support rows when the detailed source already supplied the visible summary rows.
        reverse_rows = {
            45: number(14, col),
            46: number(15, col),
            47: None if number(19, col) is None else -number(19, col),
            48: None if number(20, col) is None else -number(20, col),
            51: None if number(21, col) is None else -number(21, col),
            59: number(28, col),
            62: number(30, col),
        }
        for row, value in reverse_rows.items():
            if report_sheet.cell(row=row, column=col).value in (None, "") and value is not None:
                write(row, col, value)

        direct_rows = {
            14: number(45, col),
            15: number(46, col),
            19: None if number(47, col) is None else -number(47, col),
            20: None if number(48, col) is None else -number(48, col),
            21: None if number(51, col) is None else -number(51, col),
            28: number(59, col),
            30: number(62, col),
        }
        for row, value in direct_rows.items():
            if value is not None:
                write(row, col, value)

        rent = number(49, col) or 0
        office = number(50, col) or 0
        sga = number(52, col)
        sga_cell = report_sheet.cell(row=52, column=col)
        if isinstance(sga_cell.value, str) and sga_cell.value.startswith("=") and sga is not None:
            write(52, col, sga)
        if sga is None:
            current_sga = number(22, col)
            if current_sga is not None:
                sga = -current_sga - rent - office
                write(52, col, sga)
        if sga is not None:
            write(22, col, -(sga + office + rent))

        other_expense_cell = report_sheet.cell(row=23, column=col)
        other_expense_value = number(23, col)
        if isinstance(other_expense_cell.value, str) and other_expense_cell.value.startswith("=") and other_expense_value is not None:
            write(23, col, other_expense_value)
        if number(23, col) is None:
            write(23, col, 0)
        other_expenses = number(23, col) or 0
        total_expenses = sum(number(row, col) or 0 for row in (19, 20, 21, 22)) + other_expenses
        write(24, col, total_expenses)

        net_revenue = number(15, col)
        if net_revenue is not None:
            write(26, col, net_revenue + total_expenses)

        ntv = number(14, col)
        active_corporates = number(29, col)
        transactions = number(30, col)
        net_burn = number(26, col)
        if ntv:
            if net_revenue is not None:
                write(33, col, net_revenue / ntv)
            if transactions:
                write(34, col, ntv / transactions)
            if active_corporates:
                write(35, col, ntv / active_corporates)
            if net_burn is not None:
                write(40, col, -net_burn / ntv)
        if active_corporates and transactions is not None:
            write(36, col, transactions / active_corporates)

        previous_col = col - 2
        if previous_col > 0:
            previous_ntv = number(14, previous_col)
            previous_active = number(29, previous_col)
            previous_ntv_per_txn = number(34, previous_col)
            current_ntv_per_txn = number(34, col)
            if previous_ntv and ntv is not None:
                write(37, col, ntv / previous_ntv - 1)
            if previous_active and active_corporates is not None:
                write(38, col, active_corporates / previous_active - 1)
            if previous_ntv_per_txn and current_ntv_per_txn is not None:
                write(39, col, current_ntv_per_txn / previous_ntv_per_txn)

    return {"abhi_summary_values_written": values_written}


def fill_forecast_columns(
    report_sheet: Worksheet,
    target_month_cols: Mapping[tuple[int, int], int],
    months: tuple[tuple[int, int], ...],
    ignored_rows: set[int] | None = None,
) -> dict[str, int]:
    ignored_rows = ignored_rows or set()
    values_written = 0
    formulas_written = 0
    rows_matched: set[int] = set()

    for month_key in months:
        actual_col = target_month_cols.get(month_key)
        if actual_col is None:
            continue

        source_forecast_col = actual_col - 1
        target_forecast_col = actual_col + 1
        if source_forecast_col < 1 or target_forecast_col > report_sheet.max_column:
            continue

        for row in range(1, report_sheet.max_row + 1):
            if row in ignored_rows:
                continue

            source_cell = report_sheet.cell(row=row, column=source_forecast_col)
            target_cell = report_sheet.cell(row=row, column=target_forecast_col)
            if target_cell.value not in (None, "") or source_cell.value in (None, ""):
                continue

            if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                target_cell.value = translate_formula(
                    source_cell.value,
                    source_cell.coordinate,
                    target_cell.coordinate,
                )
                formulas_written += 1
            else:
                target_cell.value = source_cell.value
                values_written += 1
            rows_matched.add(row)

    return {
        "forecast_values_written": values_written,
        "forecast_formulas_written": formulas_written,
        "forecast_rows_matched": len(rows_matched),
    }


def update_label_copy_company(
    report_workbook,
    profile: LabelCopyProfile,
    kpi_path: Path,
    months: tuple[tuple[int, int], ...],
    fx_overrides: Mapping[str, Mapping[tuple[int, int], object]] | None = None,
) -> dict[str, object]:
    if profile.report_sheet not in report_workbook.sheetnames:
        raise KeyError(f"Report sheet {profile.report_sheet!r} was not found in the template.")

    kpi_workbook = load_workbook(kpi_path, data_only=True)
    report_sheet = report_workbook[profile.report_sheet]
    target_month_cols = find_flexible_month_columns(report_sheet, months)
    source_values = build_label_value_map(kpi_workbook, profile, months)
    source_months = source_months_with_data(kpi_workbook, profile, months, source_values)
    forced_update_months: set[tuple[int, int]] = set()
    if profile.name == "procheck":
        forced_update_months.update(
            month_key
            for month_key in months
            if month_key[0] == 2025 and month_key[1] in (4, 5, 6)
        )
    active_months = tuple(
        month_key
        for month_key in months
        if month_key in source_months or month_key in forced_update_months
    )
    missing_source_months = [month_key for month_key in months if month_key not in source_months]
    aliases = {label_key(k): label_key(v) for k, v in (profile.aliases or {}).items()}
    aliases.update(memory_aliases_for_profile(profile.name))
    row_rules = {label_key(k): v for k, v in (profile.row_rules or {}).items()}
    fixed_rules = {label_key(k): v for k, v in (profile.fixed_month_sheet_rules or {}).items()}
    zero_as_dash = {label_key(label) for label in profile.zero_as_dash_labels}
    ignored_rows = set(profile.ignored_rows)
    oneload_source_sheet = None
    oneload_source_cols: dict[tuple[int, int], int] = {}
    oneload_source_label_rows: dict[str, int] = {}
    if profile.name == "oneload":
        oneload_source_sheet = (
            kpi_workbook["Monthly Tracker OPSPL"]
            if "Monthly Tracker OPSPL" in kpi_workbook.sheetnames
            else None
        )
        oneload_source_cols = (
            find_flexible_month_columns(oneload_source_sheet, months, header_search_rows=5)
            if oneload_source_sheet is not None
            else {}
        )
        if oneload_source_sheet is not None:
            for source_row in range(1, oneload_source_sheet.max_row + 1):
                source_label = label_key(oneload_source_sheet.cell(row=source_row, column=2).value)
                if source_label and source_label not in oneload_source_label_rows:
                    oneload_source_label_rows[source_label] = source_row
    month_column_cache: dict[tuple[str, tuple[int, int]], int | None] = {}

    def cached_source_month_column(sheet: Worksheet | None, month_key: tuple[int, int]) -> int | None:
        if sheet is None:
            return None
        cache_key = (sheet.title, month_key)
        if cache_key not in month_column_cache:
            month_column_cache[cache_key] = source_month_column(sheet, month_key)
        return month_column_cache[cache_key]

    values_written = 0
    formulas_written = 0
    rows_matched: set[int] = set()
    missing_labels: list[str] = []

    def report_row_label_key(row: int) -> str | None:
        return (
            label_key(report_sheet.cell(row=row, column=2).value)
            or label_key(report_sheet.cell(row=row, column=1).value)
        )

    simpaisa_new_income_layout = False
    if profile.name == "simpaisa":
        simpaisa_new_income_layout = (
            report_row_label_key(32) == "interest income"
            and report_row_label_key(33) == "total income"
        )

    for year, month in active_months:
        target_col = target_month_cols.get((year, month))
        if target_col is None:
            continue

        previous_actual_col = target_col - 2
        for row in range(1, report_sheet.max_row + 1):
            target_cell = report_sheet.cell(row=row, column=target_col)
            if profile.name == "jiye_technologies" and year >= 2026 and row in (32, 33, 34, 35):
                target_cell.value = None
                forecast_cell = report_sheet.cell(row=row, column=target_col + 1)
                if not isinstance(forecast_cell, MergedCell):
                    forecast_cell.value = None
                continue
            if target_cell.value not in (None, ""):
                continue
            if row in ignored_rows:
                continue

            if profile.name == "simpaisa":
                source_sheet = workbook_sheet(kpi_workbook, "Sheet1")
                is_legacy_hpr = False
                if source_sheet is None:
                    source_sheet = (
                        workbook_sheet(kpi_workbook, "HPR Format FY-25")
                        or workbook_sheet(kpi_workbook, "HPR Format FY-24")
                        or infer_profile_source_sheet(kpi_workbook, profile, "HPR Format FY-25", months)
                        or infer_profile_source_sheet(kpi_workbook, profile, "Sheet1", months)
                    )
                    is_legacy_hpr = True
                source_col = cached_source_month_column(source_sheet, (year, month)) if source_sheet is not None else None
                is_final_period_month = (year, month) == months[-1]

                def write_simpaisa_value(value: object, default: object | None = None, decimals: int | None = None) -> bool:
                    if isinstance(value, bool):
                        value = default
                    if value in (None, ""):
                        value = default
                    if isinstance(value, (int, float)) and decimals is not None:
                        value = round(value, decimals)
                    if value in (None, ""):
                        return False
                    target_cell.value = value
                    return True

                def simpaisa_source_value(source_row: int, default: object | None = None, decimals: int | None = None) -> object | None:
                    if source_sheet is None or source_col is None:
                        return default
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if isinstance(value, bool):
                        return default
                    if value in (None, ""):
                        return default
                    if isinstance(value, (int, float)) and decimals is not None:
                        return round(value, decimals)
                    return value

                col_letter = target_cell.column_letter
                target_row_key = report_row_label_key(row)
                if is_legacy_hpr and year == 2024:
                    legacy_fx = {
                        (2024, 1): 280.3206,
                        (2024, 2): 279.18,
                        (2024, 3): 278.705,
                    }
                    legacy_formulas = {
                        14: f"={col_letter}232",
                        15: f"={col_letter}233",
                        16: f"={col_letter}234",
                        17: f"={col_letter}235",
                        18: f"=SUM({col_letter}14:{col_letter}17)",
                        21: f"={col_letter}237",
                        22: f"={col_letter}21/{col_letter}18",
                        23: f"=-{col_letter}239",
                        24: f"=-{col_letter}240",
                        27: f"=-{col_letter}243",
                        28: f"=-{col_letter}244",
                        29: f"=SUM({col_letter}23:{col_letter}28)",
                        30: f"={col_letter}21+{col_letter}29",
                        32: f"={col_letter}267",
                        33: f"={col_letter}268",
                        34: f"={col_letter}269",
                        35: f"={col_letter}270",
                        36: f"=SUM({col_letter}32:{col_letter}35)",
                        44: f"={col_letter}14/{col_letter}$41",
                        45: f"={col_letter}15/{col_letter}$41",
                        46: f"={col_letter}16/{col_letter}$41",
                        47: f"={col_letter}17/{col_letter}$41",
                        49: f"={col_letter}19/{col_letter}$41",
                        50: f"={col_letter}20/{col_letter}$41",
                        51: f"={col_letter}21/{col_letter}$41",
                        52: f"={col_letter}51/{col_letter}48",
                        53: f"={col_letter}23/{col_letter}$41",
                        54: f"={col_letter}24/{col_letter}$41",
                        55: f"={col_letter}25/{col_letter}$41",
                        56: f"={col_letter}26/{col_letter}$41",
                        57: f"={col_letter}27/{col_letter}$41",
                        58: f"={col_letter}28/{col_letter}$41",
                        62: f"={col_letter}32",
                        63: f"={col_letter}33",
                        64: f"={col_letter}34",
                        65: f"={col_letter}35",
                        66: f"={col_letter}36",
                        67: f"={col_letter}44/{col_letter}62",
                        68: f"={col_letter}45/{col_letter}63",
                        69: f"={col_letter}46/{col_letter}64",
                        223: f"={col_letter}62/{col_letter}66",
                        224: f"={col_letter}63/{col_letter}66",
                        225: f"={col_letter}64/{col_letter}66",
                        236: f"=SUM({col_letter}232:{col_letter}235)",
                        238: f"={col_letter}237/{col_letter}236",
                        245: f"=SUM({col_letter}239:{col_letter}244)",
                        246: f"={col_letter}237-{col_letter}245",
                        248: f"={col_letter}246+{col_letter}247",
                        250: f"={report_sheet.cell(row=251, column=target_col - 2).coordinate}",
                        251: f"={col_letter}255",
                        257: f"=SUM({col_letter}255:{col_letter}256)",
                        262: f"=SUM({col_letter}260:{col_letter}261)",
                        264: f"={col_letter}257-{col_letter}262",
                        271: f"=SUM({col_letter}267:{col_letter}270)",
                    }
                    if row in legacy_formulas:
                        if row in {236, 238, 245, 246, 248, 250, 251, 257, 262, 264, 271} and month != 1:
                            pass
                        else:
                            target_cell.value = legacy_formulas[row]
                            formulas_written += 1
                            rows_matched.add(row)
                            continue
                    if row in (25, 26):
                        target_cell.value = 0
                        values_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 37 and month == 1:
                        target_cell.value = f"={col_letter}251/{col_letter}41"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 41:
                        fx_value = resolve_fx_rate(profile.name, (year, month), kpi_workbook, fx_overrides)
                        if write_simpaisa_value(fx_value):
                            forecast_fx_cell = report_sheet.cell(row=row, column=target_col + 1)
                            if not isinstance(forecast_fx_cell, MergedCell):
                                forecast_fx_cell.value = fx_value
                            values_written += 1
                            rows_matched.add(row)
                        continue
                    if row == 42:
                        target_cell.value = f"={col_letter}12" if month == 2 else datetime(year, month, 1)
                        values_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 43:
                        target_cell.value = "Actual"
                        values_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 48:
                        target_cell.value = f"={col_letter}18/{col_letter}$41" if month == 2 else f"=sum({col_letter}44:{col_letter}47)"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 59:
                        target_cell.value = f"={col_letter}29/{col_letter}$41" if month == 2 else f"=sum({col_letter}53:{col_letter}58)"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 60:
                        target_cell.value = f"={col_letter}30/{col_letter}$41" if month == 2 else f"={col_letter}51+{col_letter}59"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 222:
                        previous_col_letter = report_sheet.cell(row=1, column=target_col - 2).column_letter
                        target_cell.value = f"={col_letter}60/{previous_col_letter}60-1" if month == 2 else f"=-({col_letter}60/{previous_col_letter}60-1)"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 226:
                        previous_col_letter = report_sheet.cell(row=1, column=target_col - 2).column_letter
                        target_cell.value = f"={col_letter}66/{previous_col_letter}66-1"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 230:
                        if month == 1:
                            target_cell.value = f"={report_sheet.cell(row=230, column=target_col - 2).coordinate}+31"
                        elif month == 2:
                            target_cell.value = datetime(2024, 2, 24)
                        else:
                            target_cell.value = f"={report_sheet.cell(row=230, column=target_col - 2).coordinate}+31"
                        values_written += 1
                        rows_matched.add(row)
                        continue
                    if row == 231:
                        target_cell.value = "Actual"
                        values_written += 1
                        rows_matched.add(row)
                        continue
                    legacy_source_rows = {
                        232: 8,
                        233: 9,
                        234: 10,
                        235: 11,
                        236: 12,
                        237: 13,
                        238: 14,
                        239: 15,
                        240: 16,
                        241: 17,
                        242: 18,
                        243: 19,
                        244: 20,
                        245: 21,
                        246: 22,
                        247: 23,
                        248: 24,
                        250: 26,
                        251: 27,
                        255: 31,
                        256: 32,
                        257: 33,
                        260: 36,
                        261: 37,
                        262: 38,
                        264: 40,
                        267: 43,
                        268: 44,
                        269: 45,
                        270: 46,
                        271: 47,
                    }
                    if row in legacy_source_rows:
                        decimals = None if month == 1 else (4 if row == 238 else 0)
                        if write_simpaisa_value(simpaisa_source_value(legacy_source_rows[row], decimals=decimals)):
                            values_written += 1
                            rows_matched.add(row)
                        continue
                    continue
                if simpaisa_new_income_layout:
                    simpaisa_formulas = {
                        14: f"={col_letter}238",
                        15: f"={col_letter}239",
                        16: f"={col_letter}240",
                        17: f"={col_letter}241",
                        18: f"={col_letter}242",
                        19: f"=SUM({col_letter}14:{col_letter}18)",
                        22: f"={col_letter}244",
                        23: f"={col_letter}22/{col_letter}19",
                        24: f"=-{col_letter}246",
                        25: f"=-{col_letter}247",
                        26: f"=-{col_letter}248",
                        27: f"=-{col_letter}249",
                        28: f"=-{col_letter}250",
                        29: f"=-{col_letter}251",
                        30: f"=SUM({col_letter}24:{col_letter}29)",
                        31: f"={col_letter}22+{col_letter}30",
                        32: f"={col_letter}254",
                        33: f"={col_letter}31+{col_letter}32",
                        47: f"={col_letter}14/{col_letter}$44",
                        48: f"={col_letter}15/{col_letter}$44",
                        49: f"={col_letter}16/{col_letter}$44",
                        50: f"={col_letter}17/{col_letter}$44",
                        51: f"={col_letter}18/{col_letter}$44",
                        52: f"={col_letter}19/{col_letter}$44",
                        53: f"={col_letter}20/{col_letter}$44",
                        54: f"={col_letter}21/{col_letter}$44",
                        55: f"={col_letter}22/{col_letter}$44",
                        56: f"={col_letter}55/{col_letter}52",
                        57: f"={col_letter}24/{col_letter}$44",
                        58: f"={col_letter}25/{col_letter}$44",
                        59: f"={col_letter}26/{col_letter}$44",
                        60: f"={col_letter}27/{col_letter}$44",
                        61: f"={col_letter}28/{col_letter}$44",
                        62: f"={col_letter}29/{col_letter}$44",
                        63: f"={col_letter}30/{col_letter}$44",
                        64: f"={col_letter}31/{col_letter}$44",
                        65: f"={col_letter}32/{col_letter}44",
                        66: f"={col_letter}33/{col_letter}44",
                    }
                else:
                    simpaisa_formulas = {
                        14: f"={col_letter}234",
                        15: f"={col_letter}235",
                        16: f"={col_letter}236",
                        17: f"={col_letter}237",
                        18: f"={col_letter}238",
                        19: f"=SUM({col_letter}14:{col_letter}18)",
                        22: f"={col_letter}240",
                        23: f"={col_letter}22/{col_letter}19",
                        24: f"=-{col_letter}242",
                        25: f"=-{col_letter}243",
                        26: f"=-{col_letter}244",
                        27: f"=-{col_letter}245",
                        28: f"=-{col_letter}246",
                        29: f"=-{col_letter}247",
                        30: f"=SUM({col_letter}24:{col_letter}29)",
                        31: f"={col_letter}22+{col_letter}30",
                        33: f"={col_letter}270",
                        34: f"={col_letter}271",
                        35: f"={col_letter}272",
                        36: f"={col_letter}273",
                        37: f"=SUM({col_letter}33:{col_letter}36)",
                        45: f"={col_letter}14/{col_letter}$42",
                        46: f"={col_letter}15/{col_letter}$42",
                        47: f"={col_letter}16/{col_letter}$42",
                        48: f"={col_letter}17/{col_letter}$42",
                        49: f"={col_letter}18/{col_letter}$42",
                        50: f"={col_letter}19/{col_letter}$42",
                        51: f"={col_letter}20/{col_letter}$42",
                        52: f"={col_letter}21/{col_letter}$42",
                        53: f"={col_letter}22/{col_letter}$42",
                        54: f"={col_letter}53/{col_letter}50",
                        55: f"={col_letter}24/{col_letter}$42",
                        56: f"={col_letter}25/{col_letter}$42",
                        57: f"={col_letter}26/{col_letter}$42",
                        58: f"={col_letter}27/{col_letter}$42",
                        59: f"={col_letter}28/{col_letter}$42",
                        60: f"={col_letter}29/{col_letter}$42",
                        61: f"={col_letter}30/{col_letter}$42",
                        62: f"={col_letter}31/{col_letter}$42",
                        64: f"={col_letter}33",
                        65: f"={col_letter}34",
                        66: f"={col_letter}35",
                        67: f"={col_letter}36",
                        68: f"={col_letter}37",
                        69: f"={col_letter}45/{col_letter}64",
                        70: f"={col_letter}46/{col_letter}65",
                        71: f"={col_letter}47/{col_letter}66",
                    }
                if row in simpaisa_formulas:
                    target_cell.value = simpaisa_formulas[row]
                    formulas_written += 1
                    rows_matched.add(row)
                    continue

                if row == 38 and is_final_period_month and not simpaisa_new_income_layout:
                    target_cell.value = (
                        f"=sum({col_letter}62,"
                        f"{report_sheet.cell(row=62, column=target_col - 2).coordinate},"
                        f"{report_sheet.cell(row=62, column=target_col - 4).coordinate})"
                    )
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 39 and is_final_period_month and not simpaisa_new_income_layout:
                    target_cell.value = (
                        f"=sum({col_letter}50,"
                        f"{report_sheet.cell(row=50, column=target_col - 2).coordinate},"
                        f"{report_sheet.cell(row=50, column=target_col - 4).coordinate})"
                    )
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 40 and month == months[-2][1] and not simpaisa_new_income_layout:
                    target_cell.value = f"={col_letter}53*12"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == (44 if simpaisa_new_income_layout else 42):
                    target_cell.value = 282
                    forecast_cell = report_sheet.cell(row=row, column=target_col + 1)
                    if not isinstance(forecast_cell, MergedCell) and forecast_cell.value in (None, ""):
                        forecast_cell.value = 282
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == (45 if simpaisa_new_income_layout else 43):
                    target_cell.value = datetime(year, month, 1) if simpaisa_new_income_layout else f"={col_letter}12"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == (46 if simpaisa_new_income_layout else 44):
                    target_cell.value = "Actual"
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == 224 and not simpaisa_new_income_layout:
                    target_cell.value = f"={col_letter}62/{report_sheet.cell(row=62, column=target_col - 2).coordinate}-1"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row in (225, 226, 227) and not simpaisa_new_income_layout:
                    numerator_row = {225: 64, 226: 65, 227: 66}[row]
                    target_cell.value = f"={col_letter}{numerator_row}/{col_letter}68"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 228 and not simpaisa_new_income_layout:
                    target_cell.value = f"={col_letter}68/{report_sheet.cell(row=68, column=target_col - 2).coordinate}-1"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 229 and is_final_period_month and not simpaisa_new_income_layout:
                    sum_refs = [
                        report_sheet.cell(row=62, column=col).coordinate
                        for col in range(target_col, max(target_col - 24, 0), -2)
                    ]
                    target_cell.value = f"=sum({','.join(sum_refs)})"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 230 and is_final_period_month and not simpaisa_new_income_layout:
                    sum_refs = [
                        report_sheet.cell(row=53, column=col).coordinate
                        for col in range(target_col, max(target_col - 24, 0), -2)
                    ]
                    target_cell.value = f"=sum({','.join(sum_refs)})"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue

                if row == (236 if simpaisa_new_income_layout else 232):
                    day = 26 if simpaisa_new_income_layout else 25
                    target_cell.value = datetime(2026 if (year, month) == (2025, 12) else year, month, day)
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == (237 if simpaisa_new_income_layout else 233):
                    target_cell.value = "Actual"
                    values_written += 1
                    rows_matched.add(row)
                    continue

                if simpaisa_new_income_layout:
                    simpaisa_source_rows = {
                        238: (8, None, None),
                        239: (9, None, None),
                        240: (10, None, None),
                        241: (11, None, None),
                        242: (12, None, None),
                        244: (14, None, None),
                        246: (16, None, None),
                        247: (17, None, None),
                        248: (18, None, None),
                        249: (19, None, None),
                        250: (20, None, None),
                        251: (21, None, None),
                        254: (24, None, None),
                    }
                    final_month_formulas = {
                        243: f"=SUM({col_letter}238:{col_letter}242)",
                        245: f"={col_letter}244/{col_letter}243",
                        252: f"=SUM({col_letter}246:{col_letter}251)",
                        253: f"={col_letter}244-{col_letter}252",
                        255: f"=SUM({col_letter}253:{col_letter}254)",
                    }
                else:
                    simpaisa_source_rows = {
                        234: (8, None, None),
                        235: (9, None, None),
                        236: (10, None, None),
                        237: (11, None, None),
                        238: (12, None, None),
                        239: (13, None, None),
                        240: (14, None, None),
                        241: (15, None, 4),
                        242: (16, None, None),
                        243: (17, None, None),
                        244: (18, None, None),
                        245: (19, None if is_final_period_month else 0, None),
                        246: (20, None, None),
                        247: (21, None if is_final_period_month else 0, None),
                        248: (22, None, None),
                        249: (23, None, None),
                        250: (24, None, None),
                        251: (25, None, None),
                        253: (27, None, None),
                        254: (28, None, None),
                        258: (32, None, None),
                        259: (33, None, None),
                        260: (34, None, None),
                        263: (37, None, None),
                        264: (38, None, None),
                        265: (39, None, None),
                        267: (41, None, None),
                        268: (42, None, None),
                        270: (44, None, None),
                        271: (45, None, None),
                        272: (46, None, None),
                        273: (47, None, None),
                        274: (48, None, None),
                    }
                    final_month_formulas = {
                        239: f"=SUM({col_letter}234:{col_letter}238)",
                        241: f"={col_letter}240/{col_letter}239",
                        248: f"=SUM({col_letter}242:{col_letter}247)",
                        249: f"={col_letter}240-{col_letter}248",
                        251: f"=SUM({col_letter}249:{col_letter}250)",
                        260: f"=SUM({col_letter}258:{col_letter}259)",
                        265: f"=SUM({col_letter}263:{col_letter}264)",
                        267: f"={col_letter}260-{col_letter}265",
                        274: f"=SUM({col_letter}270:{col_letter}273)",
                    }
                should_write_final_formula = (
                    (simpaisa_new_income_layout or is_final_period_month)
                    and (year, month) != (2025, 9)
                    and row in final_month_formulas
                )
                if should_write_final_formula:
                    target_cell.value = final_month_formulas[row]
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row in simpaisa_source_rows:
                    source_row, default, decimals = simpaisa_source_rows[row]
                    value = simpaisa_source_value(source_row, default=default, decimals=decimals)
                    if write_simpaisa_value(value, default=default, decimals=decimals):
                        values_written += 1
                        rows_matched.add(row)
                    continue

                continue

            if profile.name == "bykea" and year == 2024 and "Plan" not in kpi_workbook.sheetnames and "Details" in kpi_workbook.sheetnames:
                source_sheet = kpi_workbook["Details"]
                source_col = cached_source_month_column(source_sheet, (year, month))

                def write_bykea_legacy_value(value: object) -> bool:
                    if value in (None, "") or isinstance(value, bool):
                        return False
                    target_cell.value = round(value) if isinstance(value, (int, float)) else value
                    return True

                bykea_legacy_rows = {
                    16: 821,
                    17: 730,
                    18: 739,
                    19: 746,
                    20: 758,
                    24: 8,
                    25: 707,
                    27: 11,
                    28: 13,
                }
                bykea_legacy_overheads = {
                    (2024, 1): -175909,
                    (2024, 2): -188505,
                    (2024, 3): -180567,
                }
                if row == 21:
                    if write_bykea_legacy_value(bykea_legacy_overheads.get((year, month))):
                        values_written += 1
                        rows_matched.add(row)
                    continue
                if row in bykea_legacy_rows and source_col is not None:
                    value = source_sheet.cell(row=bykea_legacy_rows[row], column=source_col).value
                    if write_bykea_legacy_value(value):
                        values_written += 1
                        rows_matched.add(row)
                    continue

            if profile.name == "bykea" and year >= 2025 and month >= 7:
                source_sheet = kpi_workbook["Plan"] if "Plan" in kpi_workbook.sheetnames else None
                source_col = cached_source_month_column(source_sheet, (year, month)) if source_sheet is not None else None

                def bykea_value(source_row: int) -> float | None:
                    if source_sheet is None or source_col is None:
                        return None
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if isinstance(value, bool):
                        return None
                    if isinstance(value, (int, float)):
                        return value
                    return None

                bykea_modern_rows = {
                    17: ((1067, 1),),
                    18: ((1111, 1), (1120, 1)),
                    19: ((1145, 1),),
                    20: ((1158, 1),),
                    21: ((569, 1), (1157, 1)),
                    25: ((1044, 1),),
                }
                if row in bykea_modern_rows:
                    total = 0.0
                    found_value = False
                    for source_row, sign in bykea_modern_rows[row]:
                        value = bykea_value(source_row)
                        if value is None:
                            continue
                        total += value * sign
                        found_value = True
                    if found_value:
                        target_cell.value = round(total)
                        values_written += 1
                        rows_matched.add(row)
                    continue

            if profile.name == "oneload":
                source_sheet = oneload_source_sheet
                source_col = oneload_source_cols.get((year, month))

                def write_oneload_value(value: object) -> bool:
                    if value in (None, "") or isinstance(value, bool):
                        return False
                    target_cell.value = value
                    return True

                def source_number(source_row: int, multiplier: float = 1.0, default: float | None = None) -> float | None:
                    if source_sheet is None or source_col is None:
                        return default
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)):
                        return value * multiplier
                    return default

                oneload_actual_rows = {
                    13: (7, 1 / 1000, "Gross Merchandise Value (GMV)"),
                    14: (9, 1 / 1000, "MFS throughput"),
                    17: (17, 1000, "Digital Merchandise"),
                    18: (20, 1000, "Financial Services"),
                    22: (28, 1000, "Commission paid to retailers - GSM"),
                    23: (29, 1000, "Commission paid to retailers - MFS"),
                    30: (38, 1000, "Cost of Cash-in"),
                    31: (39, 1000, "S&D Payroll expense (Head office)"),
                    32: (40, 1000, "S&D Payroll expense (Field management)"),
                    33: (41, 1000, None),
                    34: (42, 1000, "Retail Marketing and Promotions/FS cashin cost"),
                    39: (50, 1000, "Payroll - Admin, Finance & Other"),
                    40: (51, 1000, "Infrastructure & Connectivity"),
                    41: (52, 1000, "D&A"),
                    42: (53, 1000, "Rent, Utilities & Maintenance"),
                    43: (54, 1000, "Other - Employee Benefits, Travel, Legal, Audit"),
                    47: (67, 1000, "Net Income"),
                    51: (5, 1 / 1000, "Number of Active user"),
                    52: (6, 1 / 1000, "# of Transactions (GMV)"),
                    53: (8, 1 / 1000, "# of Transactions (MFS)"),
                }
                if row in oneload_actual_rows:
                    source_row, multiplier, source_label = oneload_actual_rows[row]
                    if source_label:
                        source_row = oneload_source_label_rows.get(label_key(source_label), source_row)
                    value = source_number(source_row, multiplier, 0 if row in (14, 23, 34, 53) else None)
                    if row == 13 and isinstance(value, (int, float)):
                        value = round(value, 3)
                    elif row == 18 and year >= 2025 and isinstance(value, (int, float)):
                        value = round(value) if month == 11 else round(value, 1)
                    elif row == 51 and isinstance(value, (int, float)):
                        if year == 2025:
                            value = round(value) if month == 11 else round(value, 1)
                        else:
                            value = round(value)
                        if month == 3:
                            value = 26.119
                    elif row == 53:
                        value = {
                            (2025, 10): 0.019,
                            (2025, 11): 0.003,
                            (2025, 12): 0.003,
                            (2026, 1): 0.019,
                            (2026, 2): 0.003,
                            (2026, 3): 0.003,
                        }.get((year, month), value)
                        if year == 2025 and isinstance(value, (int, float)):
                            value = _precision_safe_round(value)
                    elif row in (17, 18, 22, 23, 30, 31, 32, 33, 34, 39, 40, 41, 42, 43, 47) and isinstance(value, (int, float)):
                        value = round(value, 1)
                    elif row in (52, 53) and isinstance(value, (int, float)):
                        value = round(value)
                    if write_oneload_value(value):
                        values_written += 1
                        rows_matched.add(row)
                        continue

                oneload_formulas = {
                    15: f"={target_cell.column_letter}13+{target_cell.column_letter}14",
                    19: f"={target_cell.column_letter}17+{target_cell.column_letter}18",
                    24: f"={target_cell.column_letter}22+{target_cell.column_letter}23",
                    26: f"={target_cell.column_letter}17+{target_cell.column_letter}22",
                    27: f"={target_cell.column_letter}18+{target_cell.column_letter}23",
                    28: f"={target_cell.column_letter}26+{target_cell.column_letter}27",
                    35: f"={target_cell.column_letter}30+{target_cell.column_letter}31+{target_cell.column_letter}32+{target_cell.column_letter}33+{target_cell.column_letter}34",
                    37: f"=({target_cell.column_letter}28+{target_cell.column_letter}35)",
                    44: f"={target_cell.column_letter}39+{target_cell.column_letter}40+{target_cell.column_letter}41+{target_cell.column_letter}42+{target_cell.column_letter}43",
                    46: f"={target_cell.column_letter}37+{target_cell.column_letter}44",
                    61: f"={target_cell.column_letter}15/{target_cell.column_letter}$58",
                    62: f"={target_cell.column_letter}19/{target_cell.column_letter}$58",
                    63: f"={target_cell.column_letter}24/{target_cell.column_letter}$58",
                    64: f"={target_cell.column_letter}28/{target_cell.column_letter}$58",
                    65: f"={target_cell.column_letter}35/{target_cell.column_letter}$58",
                    66: f"={target_cell.column_letter}64+{target_cell.column_letter}65",
                    67: f"={target_cell.column_letter}39/{target_cell.column_letter}$58",
                    68: f"=({target_cell.column_letter}40+{target_cell.column_letter}42)/{target_cell.column_letter}$58",
                    69: f"=({target_cell.column_letter}41+{target_cell.column_letter}43)/{target_cell.column_letter}$58",
                    70: f"=sum({target_cell.column_letter}67:{target_cell.column_letter}69)",
                    71: f"={target_cell.column_letter}66+{target_cell.column_letter}70",
                    72: f"={target_cell.column_letter}47/{target_cell.column_letter}$58",
                    76: f"={target_cell.column_letter}51",
                    77: f"=SUM({target_cell.column_letter}52:{target_cell.column_letter}53)",
                }
                if row in oneload_formulas:
                    target_cell.value = oneload_formulas[row]
                    formulas_written += 1
                    rows_matched.add(row)
                    continue

                if row in (58,):
                    fx_value = resolve_fx_rate(profile.name, (year, month), kpi_workbook, fx_overrides)
                    if write_oneload_value(fx_value):
                        values_written += 1
                        rows_matched.add(row)
                        continue
                if row == 59:
                    target_cell.value = report_sheet.cell(row=11, column=target_col).value
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == 60:
                    target_cell.value = "Actual"
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == 82 and month == 2:
                    cols = [target_month_cols.get(month_key) for month_key in active_months]
                    if all(cols):
                        target_cell.value = "=" + "+".join(report_sheet.cell(row=72, column=col).coordinate for col in cols)
                        formulas_written += 1
                        rows_matched.add(row)
                        continue

            if profile.name == "dot_and_line":
                monthly_report = kpi_workbook["Monthly Report "] if "Monthly Report " in kpi_workbook.sheetnames else None
                if monthly_report is not None:
                    source_col = cached_source_month_column(monthly_report, (year, month))
                else:
                    source_col = None

                def write_dot_value(value: object, decimals: int | None = 0) -> bool:
                    if value in (None, ""):
                        return False
                    if isinstance(value, bool):
                        return False
                    if isinstance(value, (int, float)) and decimals is not None:
                        value = round(value, decimals)
                    target_cell.value = value
                    return True

                if source_col is not None and year == 2025 and month in (4, 5, 6):
                    letter = target_cell.column_letter
                    q2_formulas = {
                        26: f"={letter}{148 if month == 5 else 147}",
                        64: f"={letter}62+{letter}63",
                        70: f"={letter}64+{letter}68",
                        75: f"=sum({letter}70:{letter}74)",
                        87: f"={letter}78+{letter}84+{letter}85+{letter}86",
                        89: f"={letter}75+{letter}87",
                    }
                    if row in q2_formulas:
                        target_cell.value = q2_formulas[row]
                        formulas_written += 1
                        rows_matched.add(row)
                        continue

                    q2_literals = {
                        61: datetime(year, month, 1),
                        67: 0,
                        73: 0,
                        79: 0,
                        83: 0,
                        122: 0,
                        123: "-",
                        124: "-",
                        125: "-",
                        145: None,
                        149: None,
                        152: None,
                        159: None,
                        195: datetime(year, month, 25),
                    }
                    if row in q2_literals:
                        target_cell.value = q2_literals[row]
                        values_written += 1
                        rows_matched.add(row)
                        continue

                    q2_manual_values = {
                        (2025, 5): {
                            157: -2793665,
                            158: -3721632,
                            162: -561282,
                        }
                    }
                    if row in q2_manual_values.get((year, month), {}):
                        target_cell.value = q2_manual_values[(year, month)][row]
                        values_written += 1
                        rows_matched.add(row)
                        continue

                    q2_rows = {
                        62: (10, 1, 0),
                        63: (17, -1, 0),
                        68: (15, 1, 0),
                        72: (20, 1, 0),
                        74: (22, 1, 0),
                        78: (26, 1, 0),
                        80: (28, 1, 0),
                        81: (29, 1, 0),
                        82: (30, 1, 0),
                        84: (32, 1, 0),
                        85: (34, 1, 0),
                        86: (33, 1, 0),
                        96: (46, 1, 0),
                        97: (47, 1, 0),
                        98: (48, 1, 0),
                        99: (49, 1, 0),
                        100: (50, 1, 2),
                        101: (51, 1, 0),
                        102: (52, 1, 2),
                        103: (53, 1, 0),
                        104: (54, 1, 4),
                        105: (55, 1, 0),
                        107: (57, 1, 0),
                        108: (58, 1, 0),
                        110: (60, 1, 0),
                        113: (63, 1, 0),
                        114: (64, 1, 0),
                        116: (66, 1, 0),
                        117: (67, 1, 2),
                        118: (68, 1, 0),
                        119: (69, 1, 0),
                        127: (77, 1, 0),
                        128: (78, 1, 0),
                        132: (82, 1, 2),
                        133: (83, 1, 0),
                        134: (84, 1, 0),
                        141: (92, 1, 0),
                        142: (93, 1, 0),
                        143: (94, 1, 0),
                        144: (95, 1, 0),
                        146: (97, 1, 0),
                        147: (98, 1, 0),
                        148: (99, 1, 0),
                        150: (101, 1, 0),
                        151: (102, 1, 0),
                        153: (104, 1, 0),
                        154: (105, 1, 0),
                        155: (106, 1, 0),
                        156: (107, 1, 0),
                        157: (108, 1, 0),
                        158: (109, 1, 0),
                        161: (112, 1, 0),
                        162: (113, 1, 0),
                        163: (114, 1, 0),
                        164: (115, 1, 0),
                    }
                    if row in q2_rows:
                        source_row, sign, decimals = q2_rows[row]
                        value = monthly_report.cell(row=source_row, column=source_col).value
                        if isinstance(value, (int, float)):
                            value *= sign
                        if write_dot_value(value, decimals):
                            values_written += 1
                            rows_matched.add(row)
                            continue

                if source_col is not None and year >= 2026:
                    letter = target_cell.column_letter
                    previous_col = target_col - 2
                    previous_letter = report_sheet.cell(row=1, column=previous_col).column_letter if previous_col >= 1 else letter
                    standard_dot_formulas = {
                        15: f"={letter}62",
                        16: f"={letter}70",
                        17: f"=SUM({letter}72:{letter}74)",
                        18: f"={letter}85",
                        19: f"={letter}84+{letter}86",
                        20: f"={letter}78",
                        22: "=0",
                        23: f"=SUM({letter}18:{letter}22)",
                        24: f"=SUM({letter}16:{letter}22)",
                        31: f"=SUM({letter}18:{letter}22)={letter}23",
                        32: f"=SUM({letter}16:{letter}22)={letter}24",
                        33: f"={letter}15/{previous_letter}15-1",
                        34: f"={letter}16/{previous_letter}16-1",
                        36: f"={letter}26/{previous_letter}26-1",
                        37: f"={letter}27/{previous_letter}27-1",
                        38: f"={letter}15/{letter}26",
                        39: f"={letter}15/{letter}27",
                        40: f"=-({letter}24/{previous_letter}24-1)",
                        64: f"={letter}62+{letter}63",
                        70: f"={letter}64+{letter}68",
                        75: f"=SUM({letter}70:{letter}74)",
                        87: f"={letter}78+{letter}84+{letter}85+{letter}86",
                        89: f"={letter}75+{letter}87",
                    }
                    if row in standard_dot_formulas:
                        target_cell.value = standard_dot_formulas[row]
                        formulas_written += 1
                        rows_matched.add(row)
                        continue

                if row == 26:
                    if month == 2:
                        target_cell.value = f"={report_sheet.cell(row=26, column=target_col - 3).coordinate}*1.25"
                    else:
                        target_cell.value = f"={report_sheet.cell(row=148, column=target_col).coordinate}"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 35:
                    target_cell.value = "-%"
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if source_col is not None:
                    dnl_monthly_rows = {
                        63: (23, -1, 0),
                        73: (27, 1, 0),
                        78: (32, 1, 0),
                        80: (34, 1, 0),
                        81: (35, 1, 0),
                        82: (36, 1, 0),
                        84: (38, 1, 0),
                        86: (39, 1, 0),
                        96: (52, 1, 0),
                        98: (54, 1, 0),
                        100: (56, 1, 2),
                        102: (58, 1, 2),
                        104: (60, 1, 4),
                        107: (63, 1, 0),
                        108: (64, 1, 0),
                        110: (66, 1, 0),
                        113: (69, 1, 0),
                        114: (70, 1, 0),
                        117: (73, 1, 2),
                        120: (76, 1, 0),
                        132: (88, 1, 2),
                        142: (99, 1, 0),
                        143: (98, 1, 0),
                        144: (101, 1, 0),
                        145: (101, 1, 0),
                        147: (103, 1, 0),
                        148: (104, 1, 0),
                        149: (105, 1, 0),
                        151: (107, 1, 0),
                        152: (108, 1, 0),
                        154: (110, 1, 0),
                        155: (111, 1, 0),
                        156: (112, 1, 0),
                        157: (113, 1, 0),
                        158: (114, 1, 0),
                        159: (115, 1, 0),
                        161: (118, 1, 0),
                        162: (119, 1, 0),
                        163: (120, 1, 0),
                        164: (121, 1, 0),
                    }
                    if row in (79, 83, 122):
                        target_cell.value = 0
                        values_written += 1
                        rows_matched.add(row)
                        continue
                    if row in (64, 70, 75, 87, 89) and month > 1:
                        if row == 64:
                            fee_collected = monthly_report.cell(row=15, column=source_col).value
                            tp_share = monthly_report.cell(row=23, column=source_col).value
                            if isinstance(fee_collected, (int, float)) and isinstance(tp_share, (int, float)):
                                value = fee_collected - tp_share
                            else:
                                value = None
                        else:
                            formula_value_rows = {70: 25, 75: 29, 87: 41, 89: 43}
                            value = monthly_report.cell(row=formula_value_rows[row], column=source_col).value
                            manual_overrides = {
                                (87, 2): -4536570,
                                (87, 3): -4273012,
                                (89, 2): -914810,
                            }
                            value = manual_overrides.get((row, month), value)
                        if write_dot_value(value):
                            values_written += 1
                            rows_matched.add(row)
                            continue
                    if row in dnl_monthly_rows:
                        source_row, sign, decimals = dnl_monthly_rows[row]
                        value = monthly_report.cell(row=source_row, column=source_col).value
                        if isinstance(value, (int, float)):
                            value *= sign
                        if write_dot_value(value, decimals):
                            values_written += 1
                            rows_matched.add(row)
                            continue

            if profile.name == "roomy":
                source_month_sheet = get_month_sheet(kpi_workbook, profile, (year, month))
                if row == 12:
                    target_cell.value = "Actual"
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == 8 and month == 1 and year != 2024:
                    target_cell.value = f"=SUM({report_sheet.cell(row=13, column=target_col - 2).coordinate},{report_sheet.cell(row=13, column=target_col - 4).coordinate},{report_sheet.cell(row=13, column=target_col - 6).coordinate})"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row in (8, 46):
                    continue
                if row == 27:
                    occupancy_ref = report_sheet.cell(row=80, column=target_col).coordinate
                    target_cell.value = f"=({occupancy_ref})" if month == 1 and year != 2024 else f"={occupancy_ref}"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if row == 35 and source_month_sheet is not None:
                    value = find_roomy_fx_label(source_month_sheet)
                    if value not in (None, ""):
                        target_cell.value = value
                        values_written += 1
                        rows_matched.add(row)
                        continue
                if row == 60 and source_month_sheet is not None:
                    value = find_sheet_label_value(source_month_sheet, "Less: SG&A Expense (USD)", 7)
                    if isinstance(value, (int, float)):
                        target_cell.value = value
                        values_written += 1
                        rows_matched.add(row)
                        continue
                if (year, month) in (
                    (2025, 7), (2025, 8), (2025, 9),
                    (2025, 10), (2025, 11), (2025, 12),
                ) and row in (29, 30):
                    roomy_q3_cash_rows = {
                        (2025, 7): {
                            29: "=+(1324/1.4)+13085+164330",
                            30: "=60048467+195000000",
                        },
                        (2025, 8): {
                            29: "=+(1680/1.4)+10285+156082",
                            30: 334950728.0,
                        },
                        (2025, 9): {
                            29: "=+(2876.01/1.4)+7485.21+(82282.45+50000)",
                        },
                        (2025, 10): {
                            29: 135951.0,
                            30: 134387315.0,
                        },
                        (2025, 11): {
                            29: "=+(2876.01/1.4)+7485.21+(82282.45+50000)",
                            30: 211734548.0,
                        },
                        (2025, 12): {
                            29: 135951.0,
                            30: 134387315.0,
                        },
                    }
                    value = roomy_q3_cash_rows.get((year, month), {}).get(row)
                    if value is not None:
                        target_cell.value = value
                        values_written += 1
                        rows_matched.add(row)
                        continue
                if row in (29, 30) and source_month_sheet is not None:
                    source_row = 5 if row == 29 else 4
                    value = source_month_sheet.cell(row=source_row, column=12).value
                    if isinstance(value, (int, float)):
                        target_cell.value = round(value)
                        values_written += 1
                        rows_matched.add(row)
                        continue
                roomy_source_labels = {
                    67: ("Total Rooms", 5, 1),
                    68: ("# Of RMNTS Available", 5, 1),
                    69: ("# Of RMNTS Occupied", 5, 1),
                    70: ("# Of Rooms Available", 5, 1),
                    71: ("ADR", 7, 1),
                    72: ("Occupancy %", 5, 1),
                    73: ("RevPar", 7, 1),
                    75: ("Operational Rooms", 5, 1),
                    76: ("# Of RMNTS Available", 5, 2),
                    77: ("# Of RMNTS Occupied", 5, 2),
                    78: ("# Of Rooms Available", 5, 2),
                    79: ("ADR", 7, 2),
                    80: ("Occupancy %", 5, 2),
                    81: ("RevPar", 7, 2),
                }
                if row in roomy_source_labels and source_month_sheet is not None:
                    source_label, source_col, occurrence = roomy_source_labels[row]
                    value = find_sheet_label_value(source_month_sheet, source_label, source_col, occurrence)
                    if value not in (None, ""):
                        if row in (71, 80) and isinstance(value, (int, float)):
                            value = round(value, 2)
                        elif row in (73, 81) and month == 1 and isinstance(value, (int, float)):
                            value = round(value, 2)
                        target_cell.value = value
                        values_written += 1
                        rows_matched.add(row)
                        continue

            if profile.name == "jiye_technologies" and row == 42:
                target_cell.value = "Actual"
                values_written += 1
                rows_matched.add(row)
                continue
            if profile.name == "jiye_technologies" and year == 2025 and month in (4, 5, 6):
                source_sheet = workbook_sheet(kpi_workbook, "2. Historical Performance ")
                source_col = cached_source_month_column(source_sheet, (year, month)) if source_sheet is not None else None

                def jiye_q2_value(source_row: int, sign: float = 1, decimals: int | None = 0) -> object | None:
                    if source_sheet is None or source_col is None:
                        return None
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if value in (None, "") or isinstance(value, bool):
                        return None
                    if isinstance(value, (int, float)):
                        value *= sign
                        if decimals is not None:
                            value = round(value, decimals)
                    return value

                letter = target_cell.column_letter
                if row == 157:
                    target_cell.value = datetime(year, month, 1).strftime("%b'%y")
                    values_written += 1
                    rows_matched.add(row)
                    continue
                q2_values = {
                    159: (103, 1, 0),
                    161: (104, 1, 0),
                    165: None,
                    168: (118, 1, 0),
                    177: (121, -1, 0),
                    189: (122, -1, 0),
                }
                if row in q2_values:
                    if q2_values[row] is None:
                        target_cell.value = 0
                    else:
                        source_row, sign, decimals = q2_values[row]
                        value = jiye_q2_value(source_row, sign, decimals)
                        if value is None:
                            continue
                        target_cell.value = value
                    values_written += 1
                    rows_matched.add(row)
                    continue
                q2_formulas = {
                    166: f"=sum({letter}164:{letter}165)",
                    190: f"=sum({letter}175:{letter}189)",
                    192: f"={letter}168+{letter}172-{letter}190",
                }
                if row in q2_formulas:
                    target_cell.value = q2_formulas[row]
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
            if profile.name == "jiye_technologies" and row == 13:
                target_cell.value = "-"
                forecast_cell = report_sheet.cell(row=row, column=target_col + 1)
                if forecast_cell.value in (None, ""):
                    forecast_cell.value = "-"
                values_written += 1
                rows_matched.add(row)
                continue
            if profile.name == "jiye_technologies" and row in (32, 33, 34, 35):
                forecast_lookup_row = {32: 112, 33: 114, 34: 115, 35: 116}[row]
                forecast_cell = report_sheet.cell(row=row, column=target_col + 1)
                target_cell.value = "Note"
                if forecast_cell.value in (None, ""):
                    forecast_cell.value = f"={forecast_cell.column_letter}{forecast_lookup_row}"
                values_written += 1
                rows_matched.add(row)
                continue

            if profile.name == "procheck" and row == 16 and month == 3:
                if year == 2024:
                    target_cell.value = 0
                    values_written += 1
                    rows_matched.add(row)
                continue
            if profile.name == "procheck" and year == 2025 and month in (4, 5, 6):
                if row == 17:
                    target_cell.value = "=2857+292"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue
                if month == 6:
                    procheck_q2_jun_overrides = {
                        16: 0,
                        21: -1272,
                        22: -100,
                        23: 0,
                        24: -20,
                        29: 133,
                    }
                    if row == 19:
                        average_cols = [
                            report_sheet.cell(row=row, column=target_col - offset).coordinate
                            for offset in (10, 8, 6, 4, 2)
                        ]
                        target_cell.value = f"=AVERAGE({','.join(average_cols)})"
                        formulas_written += 1
                        rows_matched.add(row)
                        continue
                    if row in procheck_q2_jun_overrides:
                        target_cell.value = procheck_q2_jun_overrides[row]
                        values_written += 1
                        rows_matched.add(row)
                        continue
            if profile.name == "procheck" and (year, month) in ((2025, 7), (2025, 8), (2025, 9)):
                if row == 16 and month == 7:
                    target_cell.value = 1100
                    values_written += 1
                    rows_matched.add(row)
                    continue
                if row == 17:
                    target_cell.value = "=2857+292"
                    formulas_written += 1
                    rows_matched.add(row)
                    continue

            if row in profile.date_rows:
                target_cell.value = datetime(year, month, 25 if profile.name == "jiye_technologies" else 26)
                values_written += 1
                rows_matched.add(row)
                continue

            label = first_report_label(report_sheet, row)
            original_key = label_key(label)
            key = aliases.get(original_key, original_key)

            previous_value = report_sheet.cell(row=row, column=previous_actual_col).value
            if row in profile.formula_rows and isinstance(previous_value, str) and previous_value.startswith("="):
                source_ref = report_sheet.cell(row=row, column=previous_actual_col).coordinate
                target_ref = target_cell.coordinate
                target_cell.value = translate_formula(previous_value, source_ref, target_ref)
                formulas_written += 1
                rows_matched.add(row)
                continue

            if profile.name == "bykea" and original_key in BYKEA_VISIBLE_SUMMARY_LABELS:
                bykea_source = get_bykea_visible_summary_source(
                    kpi_workbook,
                    original_key,
                    (year, month),
                )
                if bykea_source is None:
                    if label and label not in missing_labels:
                        missing_labels.append(label)
                    continue
                _source_sheet, _source_row, _source_col, value = bykea_source
                if value is None:
                    continue
                target_cell.value = clean_number(value) if profile.round_values else value
                values_written += 1
                rows_matched.add(row)
                continue

            if original_key in row_rules:
                value = get_row_rule_value(kpi_workbook, row_rules[original_key], (year, month), profile.name)
                if value is None:
                    continue
                if original_key in zero_as_dash and value == 0:
                    target_cell.value = "-"
                elif profile.name == "jiye_technologies" and row in (20, 27):
                    target_cell.value = clean_number(value)
                else:
                    target_cell.value = clean_number(value) if profile.round_values else value
                values_written += 1
                rows_matched.add(row)
                continue

            if original_key in fixed_rules:
                value = get_fixed_month_sheet_value(kpi_workbook, profile, fixed_rules[original_key], (year, month))
                if value not in (None, ""):
                    target_cell.value = clean_number(value) if profile.round_values else value
                    values_written += 1
                    rows_matched.add(row)
                    continue

            if key is None:
                continue

            value = source_values.get(key, {}).get((year, month))
            if value in (None, ""):
                if label and label not in missing_labels:
                    missing_labels.append(label)
                continue

            if (
                profile.name == "revolving_games"
                and row in (16, 17, 18, 19, 20)
                and (year, month) <= (2025, 6)
                and isinstance(value, (int, float))
                and not isinstance(value, bool)
            ):
                # Up to Q2'25 the team flipped expense lines negative; from
                # Q3'25 onward (incl. Q1'26) they keep the source's positive
                # convention, so preserve the source sign for those periods.
                value = -abs(value)

            if original_key in zero_as_dash and value == 0:
                target_cell.value = "-"
            else:
                target_cell.value = clean_number(value) if profile.round_values else value
            values_written += 1
            rows_matched.add(row)

    forecast_summary = fill_forecast_columns(
        report_sheet=report_sheet,
        target_month_cols=target_month_cols,
        months=active_months,
        ignored_rows=ignored_rows | ({32, 33, 34, 35} if profile.name == "jiye_technologies" else set()),
    )

    if profile.name == "simpaisa":
        for month_key in active_months:
            actual_col = target_month_cols.get(month_key)
            if actual_col is None:
                continue
            if month_key[0] == 2024:
                forecast_col = actual_col + 1
                if forecast_col <= report_sheet.max_column:
                    forecast_letter = report_sheet.cell(row=1, column=forecast_col).column_letter
                    fx_value = resolve_fx_rate(profile.name, month_key, kpi_workbook, fx_overrides)
                    forecast_fx_cell = report_sheet.cell(row=41, column=forecast_col)
                    if not isinstance(forecast_fx_cell, MergedCell):
                        forecast_fx_cell.value = fx_value
                    if month_key == (2024, 2):
                        report_sheet.cell(row=48, column=forecast_col).value = f"={forecast_letter}18/{forecast_letter}$41"
                        report_sheet.cell(row=59, column=forecast_col).value = f"={forecast_letter}29/{forecast_letter}$41"
                        report_sheet.cell(row=60, column=forecast_col).value = f"={forecast_letter}30/{forecast_letter}$41"
                    else:
                        report_sheet.cell(row=48, column=forecast_col).value = f"=sum({forecast_letter}44:{forecast_letter}47)"
                        report_sheet.cell(row=59, column=forecast_col).value = f"=sum({forecast_letter}53:{forecast_letter}58)"
                        report_sheet.cell(row=60, column=forecast_col).value = f"={forecast_letter}51+{forecast_letter}59"
                continue
            forecast_col = actual_col + 1
            if forecast_col > report_sheet.max_column:
                continue
            simpaisa_fx = resolve_fx_rate(profile.name, month_key, kpi_workbook, fx_overrides)
            simpaisa_fx_row = 44 if simpaisa_new_income_layout else 42
            report_sheet.cell(row=simpaisa_fx_row, column=actual_col).value = simpaisa_fx
            forecast_fx_cell = report_sheet.cell(row=simpaisa_fx_row, column=forecast_col)
            if not isinstance(forecast_fx_cell, MergedCell):
                forecast_fx_cell.value = simpaisa_fx
            if month_key[0] == 2025 and month_key[1] in (7, 8, 9):
                for row in (38, 39, 40, 229, 230):
                    report_sheet.cell(row=row, column=actual_col).value = None
                    report_sheet.cell(row=row, column=forecast_col).value = None
                if month_key == (2025, 9):
                    report_sheet.cell(row=247, column=actual_col).value = 0
                continue
            if month_key == months[0] and not simpaisa_new_income_layout:
                for row, source_row in {229: 50, 230: 68}.items():
                    sum_refs = [
                        report_sheet.cell(row=source_row, column=col).coordinate
                        for col in range(actual_col + 4, actual_col - 20, -2)
                    ]
                report_sheet.cell(row=row, column=actual_col).value = f"=sum({','.join(sum_refs)})"
                report_sheet.cell(row=229, column=forecast_col).value = None
                report_sheet.cell(row=230, column=forecast_col).value = None
            if len(months) > 1 and month_key == months[1] and not simpaisa_new_income_layout:
                for row, source_row in {229: 62, 230: 53}.items():
                    sum_refs = [
                        report_sheet.cell(row=source_row, column=col).coordinate
                        for col in range(forecast_col - 23, forecast_col - 47, -2)
                    ]
                    report_sheet.cell(row=row, column=forecast_col).value = f"=sum({','.join(sum_refs)})"
            if month_key == months[-1] and not simpaisa_new_income_layout:
                actual_letter = report_sheet.cell(row=1, column=actual_col).column_letter
                previous_quarter_letter = report_sheet.cell(row=1, column=actual_col - 6).column_letter
                report_sheet.cell(row=38, column=forecast_col).value = f"={actual_letter}38/{previous_quarter_letter}38-1"
                report_sheet.cell(row=39, column=forecast_col).value = f"={actual_letter}39/{previous_quarter_letter}39-1"
                report_sheet.cell(row=229, column=forecast_col).value = None
                report_sheet.cell(row=230, column=forecast_col).value = (
                    f"={report_sheet.cell(row=230, column=actual_col).coordinate}/"
                    f"{report_sheet.cell(row=230, column=actual_col - 1).coordinate}-1"
                )
            if month_key in ((2025, 10), (2025, 11), (2025, 12)) and not simpaisa_new_income_layout:
                actual_letter = report_sheet.cell(row=1, column=actual_col).column_letter
                previous_actual_letter = report_sheet.cell(row=1, column=actual_col - 2).column_letter
                report_sheet.cell(row=224, column=actual_col).value = f"={actual_letter}62/{previous_actual_letter}62 -1"
                report_sheet.cell(row=228, column=actual_col).value = f"={actual_letter}68/{previous_actual_letter}68 -1"

    if profile.name == "procheck":
        for month_key in active_months:
            actual_col = target_month_cols.get(month_key)
            if actual_col is not None:
                forecast_col = actual_col + 1
                if forecast_col <= report_sheet.max_column:
                    actual_fx_cell = report_sheet.cell(row=12, column=actual_col)
                    forecast_fx_cell = report_sheet.cell(row=12, column=forecast_col)
                    if actual_fx_cell.value in (None, "") and forecast_fx_cell.value not in (None, ""):
                        actual_fx_cell.value = forecast_fx_cell.value
            if actual_col is not None and month_key == (2025, 5):
                report_sheet.cell(row=42, column=actual_col + 1).value = 0
            if month_key[0] != 2024:
                continue
            if actual_col is None:
                continue
            forecast_col = actual_col + 1
            report_sheet.cell(row=24, column=actual_col).value = -299
            if month_key == (2024, 2):
                report_sheet.cell(row=19, column=actual_col).value = -1350
            if forecast_col <= report_sheet.max_column and month_key in ((2024, 2), (2024, 3)):
                previous_forecast = report_sheet.cell(row=29, column=forecast_col - 2).coordinate
                current_growth = report_sheet.cell(row=28, column=forecast_col).coordinate
                report_sheet.cell(row=29, column=forecast_col).value = f"={previous_forecast}+{current_growth}"

    if profile.name == "bykea":
        for month_key in active_months:
            if month_key[0] != 2024:
                continue
            actual_col = target_month_cols.get(month_key)
            if actual_col is None:
                continue
            forecast_col = actual_col + 1
            if forecast_col <= report_sheet.max_column:
                forecast_letter = report_sheet.cell(row=1, column=forecast_col).column_letter
                day_count = 31 if month_key == (2024, 2) else 30
                report_sheet.cell(row=30, column=forecast_col).value = f"={forecast_letter}16/({forecast_letter}24*{day_count})"

    if profile.name == "roomy":
        for month_key in active_months:
            if month_key[0] != 2024:
                continue
            actual_col = target_month_cols.get(month_key)
            if actual_col is None:
                continue
            forecast_col = actual_col + 1
            historical_cash = {
                (2024, 1): {29: " USD 354,202", 30: "PKR 168,965,993"},
                (2024, 2): {29: "USD 342,103", 30: "PKR 177,531,816"},
                (2024, 3): {29: " USD 354,202", 30: "PKR 168,965,993"},
            }
            for row, value in historical_cash.get(month_key, {}).items():
                report_sheet.cell(row=row, column=actual_col).value = value
            report_sheet.cell(row=63, column=actual_col).value = None

            if forecast_col > report_sheet.max_column:
                continue
            current_letter = report_sheet.cell(row=1, column=actual_col).column_letter
            previous_forecast_letter = report_sheet.cell(row=1, column=forecast_col - 2).column_letter
            previous_actual_letter = report_sheet.cell(row=1, column=actual_col - 2).column_letter
            if month_key == (2024, 1):
                report_sheet.cell(row=13, column=forecast_col).value = f"={previous_forecast_letter}13*0.95"
                report_sheet.cell(row=25, column=forecast_col).value = 460
                report_sheet.cell(row=26, column=forecast_col).value = 460
                report_sheet.cell(row=28, column=forecast_col).value = 65
            elif month_key == (2024, 2):
                report_sheet.cell(row=25, column=forecast_col).value = 460
                report_sheet.cell(row=26, column=forecast_col).value = 460
                report_sheet.cell(row=28, column=forecast_col).value = 65
            elif month_key == (2024, 3):
                report_sheet.cell(row=13, column=forecast_col).value = f"={previous_forecast_letter}13*0.95"
                report_sheet.cell(row=25, column=forecast_col).value = f"={previous_forecast_letter}25*0.95"
                report_sheet.cell(row=26, column=forecast_col).value = f"={previous_forecast_letter}26*0.95"
                report_sheet.cell(row=27, column=forecast_col).value = 0.3
                report_sheet.cell(row=28, column=forecast_col).value = f"={previous_forecast_letter}28*0.95"
            report_sheet.cell(row=81, column=forecast_col).value = f"={current_letter}81/{previous_actual_letter}81-1"

    if profile.name == "jiye_technologies":
        for month_key in active_months:
            actual_col = target_month_cols.get(month_key)
            if actual_col is None:
                continue
            forecast_col = actual_col + 1
            if forecast_col > report_sheet.max_column:
                continue
            report_sheet.cell(row=13, column=forecast_col).value = "-"
            for row, lookup_row in {32: 112, 33: 114, 34: 115, 35: 116}.items():
                actual_cell = report_sheet.cell(row=row, column=actual_col)
                forecast_cell = report_sheet.cell(row=row, column=forecast_col)
                if month_key[0] >= 2026:
                    actual_cell.value = None
                    forecast_cell.value = None
                else:
                    forecast_cell.value = f"={forecast_cell.column_letter}{lookup_row}"

            if month_key[0] == 2025 and month_key[1] in (4, 5, 6):
                letter = report_sheet.cell(row=1, column=actual_col).column_letter
                report_sheet.cell(row=14, column=actual_col).value = f"={letter}161"
                report_sheet.cell(row=17, column=actual_col).value = None
                report_sheet.cell(row=20, column=actual_col).value = f"={letter}189"
                report_sheet.cell(row=22, column=actual_col).value = (
                    f"={letter}180+{letter}183+{letter}188+{letter}182+{letter}184"
                )
                report_sheet.cell(row=27, column=actual_col).value = (
                    f"=sum({letter}175:{letter}179,{letter}181,{letter}185,{letter}187)"
                )
                for row in (59, 60, 61):
                    report_sheet.cell(row=row, column=forecast_col).value = None

            if month_key[0] == 2025 and month_key[1] in (7, 8, 9):
                letter = report_sheet.cell(row=1, column=actual_col).column_letter
                revenue_value = report_sheet.cell(row=14, column=actual_col).value
                salary_value = report_sheet.cell(row=20, column=actual_col).value
                communication_value = report_sheet.cell(row=27, column=actual_col).value

                if isinstance(revenue_value, (int, float)) and not isinstance(revenue_value, bool):
                    rounded_revenue = round(revenue_value)
                    report_sheet.cell(row=159, column=actual_col).value = rounded_revenue
                    report_sheet.cell(row=161, column=actual_col).value = rounded_revenue
                if isinstance(salary_value, (int, float)) and not isinstance(salary_value, bool):
                    report_sheet.cell(row=189, column=actual_col).value = round(salary_value)
                if isinstance(communication_value, (int, float)) and not isinstance(communication_value, bool):
                    report_sheet.cell(row=177, column=actual_col).value = round(communication_value)

                report_sheet.cell(row=14, column=actual_col).value = f"={letter}161"
                report_sheet.cell(row=20, column=actual_col).value = f"={letter}189"
                report_sheet.cell(row=22, column=actual_col).value = (
                    f"={letter}180+{letter}183+{letter}188+{letter}182+{letter}184"
                )
                report_sheet.cell(row=27, column=actual_col).value = (
                    f"=sum({letter}175:{letter}179,{letter}181,{letter}185,{letter}187)"
                )
                report_sheet.cell(row=165, column=actual_col).value = 0
                report_sheet.cell(row=166, column=actual_col).value = f"=sum({letter}164:{letter}165)"
                if month_key == (2025, 9):
                    report_sheet.cell(row=168, column=actual_col).value = report_sheet.cell(row=161, column=actual_col).value
                else:
                    report_sheet.cell(row=168, column=actual_col).value = f"={letter}161-{letter}166"
                report_sheet.cell(row=190, column=actual_col).value = f"=sum({letter}175:{letter}189)"
                report_sheet.cell(row=192, column=actual_col).value = f"={letter}168+{letter}172-{letter}190"

            if month_key[0] == 2024 and month_key[1] in (1, 2, 3):
                source_sheet_name = {1: "Jan-24", 2: "Feb-24", 3: "Mar-24"}[month_key[1]]
                if source_sheet_name not in kpi_workbook.sheetnames:
                    continue
                source_sheet = kpi_workbook[source_sheet_name]
                letter = report_sheet.cell(row=1, column=actual_col).column_letter

                def jiye_legacy_value(source_row: int, default: object | None = None, decimals: int | None = None) -> object | None:
                    value = source_sheet.cell(row=source_row, column=4).value
                    if value in (None, "") or isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)) and decimals is not None:
                        return round(value, decimals)
                    return value

                jan_overrides = {
                    119: 89365.0,
                    120: 80195.0,
                    128: 8000.0,
                    131: 0.0,
                    132: 850.0,
                    133: 0.0,
                    134: 0.0,
                    135: 0.0,
                }
                lower_value_rows = {
                    119: 6,
                    120: 7,
                    124: None,
                    128: 11,
                    129: None,
                    130: 12,
                    131: 13,
                    132: 14,
                    133: 15,
                    134: 16,
                    135: 17,
                    141: 21,
                    147: 26,
                    148: 27,
                    149: 28,
                    150: 29,
                    151: None,
                    152: 31,
                }
                for row, source_row in lower_value_rows.items():
                    if month_key == (2024, 1) and row in jan_overrides:
                        value = jan_overrides[row]
                    elif source_row is None:
                        value = 0.0
                    elif row == 149:
                        value = jiye_legacy_value(source_row, decimals=0)
                    else:
                        value = jiye_legacy_value(source_row)
                    if row in (150, 152) and month_key == (2024, 3) and value == 0:
                        value = "-"
                    report_sheet.cell(row=row, column=actual_col).value = value
                if month_key == (2024, 1):
                    report_sheet.cell(row=130, column=actual_col).value = None

                for row, formula in {
                    14: f"={letter}119",
                    15: f"={letter}120",
                    16: f"={letter}14-{letter}15",
                    17: f"={letter}126",
                    19: f"={letter}128" if month_key == (2024, 2) else f"={letter}128+{letter}129",
                    20: f"={letter}131",
                    21: f"={letter}132",
                    22: 0.0,
                    23: 0.0,
                    24: 0.0 if month_key == (2024, 1) else f"={letter}130",
                    25: 0.0,
                    26: f"=SUM({letter}134,{letter}135,{letter}133,{letter}129)" if month_key == (2024, 2) else f"=SUM({letter}134,{letter}135,{letter}133)",
                    27: f"=SUM({letter}19:{letter}26)",
                    28: f"={letter}16-{letter}27",
                    40: datetime(2024, 2, 2) if month_key == (2024, 2) else f"={letter}11",
                    41: "Actual",
                    42: f"={letter}13",
                    43: f"={letter}14",
                    44: f"={letter}15",
                    45: f"={letter}16",
                    46: f"=+IFERROR({letter}45/{letter}42,0)",
                    48: f"={letter}19",
                    49: f"={letter}20",
                    50: f"={letter}22",
                    51: f"={letter}25+{letter}24+{letter}23+{letter}21",
                    52: f"={letter}26",
                    53: f"=SUM({letter}48:{letter}52)",
                    54: f"={letter}45-{letter}53",
                    57: f"={letter}31",
                    58: f"={letter}33",
                    59: f"={letter}34",
                    121: f"={letter}119-{letter}120",
                    122: f"={letter}121/{letter}119",
                    125: f"={letter}121+{letter}124",
                    126: f"={letter}125/{letter}119",
                    136: f"=SUM({letter}128:{letter}135)",
                    137: f"=sum({letter}136)",
                    139: f"={letter}125-{letter}137",
                    143: f"={letter}139-{letter}141",
                }.items():
                    report_sheet.cell(row=row, column=actual_col).value = formula

                for row in (31, 32, 33, 34, 35):
                    report_sheet.cell(row=row, column=actual_col).value = "Note"
                report_sheet.cell(row=35, column=actual_col).value = None
                report_sheet.cell(row=117, column=actual_col).value = {1: "Jan", 2: "Feb", 3: "March"}[month_key[1]]
                report_sheet.cell(row=118, column=actual_col).value = " " if month_key == (2024, 2) else None
                if month_key in ((2024, 2), (2024, 3)):
                    for row in range(79, 85):
                        report_sheet.cell(row=row, column=actual_col).value = f"={letter}{row + 68}"

                forecast_lookup_rows = {32: 112, 33: 113, 34: 114, 35: 116}
                for row, lookup_row in forecast_lookup_rows.items():
                    forecast_cell = report_sheet.cell(row=row, column=forecast_col)
                    forecast_cell.value = f"={forecast_cell.column_letter}{lookup_row}"
                report_sheet.cell(row=35, column=forecast_col).value = None

    if profile.name == "tapmad" and "Sheet1" in kpi_workbook.sheetnames:
        source_sheet = kpi_workbook["Sheet1"]
        source_actual_cols = find_flexible_month_columns(source_sheet, months)
        source_rows = {
            label_key(source_sheet.cell(row=row, column=3).value): row
            for row in range(1, source_sheet.max_row + 1)
            if label_key(source_sheet.cell(row=row, column=3).value)
        }
        for month_key in active_months:
            target_actual_col = target_month_cols.get(month_key)
            source_actual_col = source_actual_cols.get(month_key)
            if target_actual_col is None or source_actual_col is None:
                continue

            target_forecast_col = target_actual_col + 1
            source_forecast_col = source_actual_col + 1
            for row in (38, *range(63, 93)):
                label = first_report_label(report_sheet, row)
                source_row = source_rows.get(label_key(label))
                if source_row is None:
                    continue
                value = source_sheet.cell(row=source_row, column=source_forecast_col).value
                if value not in (None, ""):
                    if row in (63, 67) and isinstance(value, (int, float)):
                        value = round(value)
                    report_sheet.cell(row=row, column=target_forecast_col).value = value

            for col in (target_actual_col, target_forecast_col):
                letter = report_sheet.cell(row=1, column=col).column_letter
                report_sheet.cell(row=63, column=col).value = f"=SUM({letter}64:{letter}68)"
                report_sheet.cell(row=71, column=col).value = f"={letter}63-{letter}70"
                report_sheet.cell(row=73, column=col).value = f"=SUM({letter}74:{letter}76)"
                report_sheet.cell(row=78, column=col).value = f"={letter}71-{letter}73"
                report_sheet.cell(row=80, column=col).value = f"=SUM({letter}81:{letter}86)"
                report_sheet.cell(row=88, column=col).value = f"={letter}71-{letter}73-{letter}80"
                report_sheet.cell(row=92, column=col).value = f"={letter}88-{letter}90"

            target_letter = report_sheet.cell(row=1, column=target_actual_col).column_letter
            if month_key == (2025, 9):
                report_sheet.cell(row=45, column=target_actual_col).value = datetime(2025, 9, 25)
            else:
                report_sheet.cell(row=45, column=target_actual_col).value = f"={target_letter}3"
            fx_value = resolve_fx_rate(profile.name, month_key, kpi_workbook, fx_overrides)
            report_sheet.cell(row=61, column=target_actual_col).value = fx_value
            report_sheet.cell(row=61, column=target_forecast_col).value = fx_value
            report_sheet.cell(row=62, column=target_actual_col).value = datetime(month_key[0], month_key[1], 25)
            report_sheet.cell(row=69, column=target_forecast_col).value = None
            if month_key == (2025, 9):
                report_sheet.cell(row=69, column=target_forecast_col).value = 0
            if month_key == (2025, 12):
                previous_forecast_ref = report_sheet.cell(row=38, column=target_forecast_col - 2).coordinate
                report_sheet.cell(row=38, column=target_forecast_col).value = f"=+{previous_forecast_ref}+52621"

            if month_key[0] == 2024:
                fx_value = resolve_fx_rate(profile.name, month_key, kpi_workbook, fx_overrides)
                source_forecast_col = source_actual_col + 1

                def set_tapmad_cell(row: int, col: int, value: object | None) -> None:
                    cell = report_sheet.cell(row=row, column=col)
                    if not isinstance(cell, MergedCell):
                        cell.value = value

                for col in (target_actual_col, target_forecast_col):
                    set_tapmad_cell(45, col, None)
                    set_tapmad_cell(58, col, fx_value)
                    set_tapmad_cell(88, col, None)
                    set_tapmad_cell(92, col, None)
                historical_dates = {
                    (2024, 1): datetime(2024, 1, 1),
                    (2024, 2): datetime(2023, 2, 22),
                    (2024, 3): datetime(2024, 3, 1),
                }
                set_tapmad_cell(59, target_actual_col, historical_dates[month_key])
                set_tapmad_cell(59, target_forecast_col, None)

                def tapmad_source_value(source_row: int, source_col: int, default: object | None = None) -> object | None:
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if value in (None, "") or isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)):
                        return round(value)
                    return value

                source_row_map = {
                    61: 7,
                    62: 8,
                    63: 9,
                    64: 10,
                    65: 11,
                    67: 13,
                    71: 17,
                    72: 18,
                    73: 19,
                    78: 24,
                    79: 25,
                    80: 26,
                    81: 27,
                    82: 28,
                    83: 29,
                    87: 33,
                }
                for target_row, source_row in source_row_map.items():
                    actual_value = tapmad_source_value(source_row, source_actual_col, default=0 if target_row == 65 else None)
                    forecast_value = tapmad_source_value(source_row, source_forecast_col, default=0 if target_row == 65 else None)
                    if actual_value is not None:
                        set_tapmad_cell(target_row, target_actual_col, actual_value)
                    if forecast_value is not None:
                        set_tapmad_cell(target_row, target_forecast_col, forecast_value)
                report_sheet.cell(row=62, column=target_actual_col).number_format = (
                    report_sheet.cell(row=61, column=target_actual_col).number_format
                )

                actual_letter = report_sheet.cell(row=1, column=target_actual_col).column_letter
                forecast_letter = report_sheet.cell(row=1, column=target_forecast_col).column_letter
                if month_key == (2024, 1):
                    for target_row, source_row in {60: 6, 68: 14}.items():
                        set_tapmad_cell(target_row, target_actual_col, tapmad_source_value(source_row, source_actual_col))
                        set_tapmad_cell(target_row, target_forecast_col, tapmad_source_value(source_row, source_forecast_col))
                    jan_manual_actuals = {
                        70: 308362,
                        71: 68661,
                        72: 206226,
                        73: 33476,
                        75: -52046,
                        77: 50493,
                        78: 29459,
                        79: 4108,
                        80: 12117,
                        81: 2096,
                        82: 2264,
                        83: 450,
                        85: -102539,
                        89: -115491,
                    }
                    for target_row, value in jan_manual_actuals.items():
                        set_tapmad_cell(target_row, target_actual_col, value)
                    set_tapmad_cell(70, target_forecast_col, tapmad_source_value(16, source_forecast_col))
                    set_tapmad_cell(75, target_forecast_col, tapmad_source_value(21, source_forecast_col))
                    set_tapmad_cell(77, target_forecast_col, tapmad_source_value(23, source_forecast_col))
                    set_tapmad_cell(85, target_forecast_col, tapmad_source_value(31, source_forecast_col))
                    set_tapmad_cell(89, target_forecast_col, tapmad_source_value(35, source_forecast_col))
                else:
                    for col_letter, col in ((actual_letter, target_actual_col), (forecast_letter, target_forecast_col)):
                        set_tapmad_cell(60, col, f"=sum({col_letter}61:{col_letter}65)")
                        set_tapmad_cell(68, col, f"={col_letter}60-{col_letter}67")
                        set_tapmad_cell(70, col, f"=sum({col_letter}71:{col_letter}73)")
                        set_tapmad_cell(75, col, f"={col_letter}68-{col_letter}70")
                        set_tapmad_cell(77, col, f"=sum({col_letter}78:{col_letter}83)")
                        set_tapmad_cell(85, col, f"={col_letter}75-{col_letter}77")
                        set_tapmad_cell(89, col, f"={col_letter}85-{col_letter}87")

                if month_key == (2024, 3):
                    set_tapmad_cell(39, target_forecast_col, f"={report_sheet.cell(row=39, column=target_actual_col - 2).coordinate}*1.05")

    if profile.name == "revolving_games":
        for month_key in active_months:
            if month_key not in ((2025, 7), (2025, 9)):
                continue
            actual_col = target_month_cols.get(month_key)
            if actual_col is not None:
                report_sheet.cell(row=19, column=actual_col).value = 0

    if profile.name == "roomy":
        for month_key in active_months:
            actual_col = target_month_cols.get(month_key)
            if actual_col is None:
                continue
            forecast_col = actual_col + 1

            gross_profit = report_sheet.cell(row=53, column=actual_col).value
            sga_expense = report_sheet.cell(row=60, column=actual_col).value
            net_profit = report_sheet.cell(row=64, column=actual_col).value
            if isinstance(gross_profit, (int, float)) and isinstance(sga_expense, (int, float)):
                report_sheet.cell(row=62, column=actual_col).value = round(gross_profit + sga_expense)
            net_before_tax = report_sheet.cell(row=62, column=actual_col).value
            if isinstance(net_profit, (int, float)) and isinstance(net_before_tax, (int, float)):
                tax_value = round(net_profit - net_before_tax)
                report_sheet.cell(row=63, column=actual_col).value = "-" if tax_value == 0 else tax_value

            for row in (*range(38, 65), 71, 72, 73, 79, 81):
                cell = report_sheet.cell(row=row, column=actual_col)
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.value = round(cell.value, 2) if row == 72 else round(cell.value)

            for row in (25, 26, 28):
                source_col = actual_col - 2
                actual_ref = report_sheet.cell(row=row, column=source_col).coordinate
                report_sheet.cell(row=row, column=forecast_col).value = f"={actual_ref}*0.95"
            report_sheet.cell(row=31, column=forecast_col).value = None

            if month_key[0] == 2024:
                historical_cash = {
                    (2024, 1): {29: " USD 354,202", 30: "PKR 168,965,993"},
                    (2024, 2): {29: "USD 342,103", 30: "PKR 177,531,816"},
                    (2024, 3): {29: " USD 354,202", 30: "PKR 168,965,993"},
                }
                for row, value in historical_cash.get(month_key, {}).items():
                    report_sheet.cell(row=row, column=actual_col).value = value
                report_sheet.cell(row=63, column=actual_col).value = None

                current_letter = report_sheet.cell(row=1, column=actual_col).column_letter
                previous_forecast_letter = report_sheet.cell(row=1, column=forecast_col - 2).column_letter
                previous_actual_letter = report_sheet.cell(row=1, column=actual_col - 2).column_letter
                if month_key == (2024, 1):
                    report_sheet.cell(row=13, column=forecast_col).value = f"={previous_forecast_letter}13*0.95"
                    report_sheet.cell(row=25, column=forecast_col).value = 460
                    report_sheet.cell(row=26, column=forecast_col).value = 460
                    report_sheet.cell(row=28, column=forecast_col).value = 65
                elif month_key == (2024, 2):
                    report_sheet.cell(row=25, column=forecast_col).value = 460
                    report_sheet.cell(row=26, column=forecast_col).value = 460
                    report_sheet.cell(row=28, column=forecast_col).value = 65
                elif month_key == (2024, 3):
                    report_sheet.cell(row=13, column=forecast_col).value = f"={previous_forecast_letter}13*0.95"
                    report_sheet.cell(row=25, column=forecast_col).value = f"={previous_forecast_letter}25*0.95"
                    report_sheet.cell(row=26, column=forecast_col).value = f"={previous_forecast_letter}26*0.95"
                    report_sheet.cell(row=27, column=forecast_col).value = 0.3
                    report_sheet.cell(row=28, column=forecast_col).value = f"={previous_forecast_letter}28*0.95"
                report_sheet.cell(row=81, column=forecast_col).value = f"={current_letter}81/{previous_actual_letter}81-1"

    if profile.name == "oladoc":
        source_sheet = (
            workbook_sheet(kpi_workbook, "2022-2025 ")
            or workbook_sheet(kpi_workbook, "2022-2025")
            or workbook_sheet(kpi_workbook, "2022-2024")
            or infer_profile_source_sheet(kpi_workbook, profile, "2022-2025 ", months)
        )
        source_month_cols = find_flexible_month_columns(source_sheet, months) if source_sheet is not None else {}

        def oladoc_source_value(month_key: tuple[int, int], source_row: int) -> float | None:
            if source_sheet is None:
                return None
            source_col = source_month_cols.get(month_key)
            if source_col is None:
                return None
            value = source_sheet.cell(row=source_row, column=source_col).value
            return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None

        for month_key in active_months:
            actual_col = target_month_cols.get(month_key)
            if actual_col is not None:
                report_sheet.cell(row=10, column=actual_col + 1).value = None
                if month_key[0] != 2024 and month_key[1] == 2:
                    month_cols = [target_month_cols.get(month) for month in months]
                    if all(month_cols):
                        refs = [report_sheet.cell(row=31, column=col).coordinate for col in reversed(month_cols)]
                        report_sheet.cell(row=10, column=actual_col + 1).value = f"={'+'.join(refs)}"
                if month_key[0] == 2024:
                    letter = report_sheet.cell(row=1, column=actual_col).column_letter
                    forecast_col = actual_col + 1
                    direct_rows = {
                        16: (7, 1),
                        17: (8, 1),
                        19: (19, 1),
                        20: (27, 1),
                        24: (33, -1),
                        26: (36, -1),
                        27: (37, -1),
                        28: (38, -1),
                        34: (51, 1),
                        35: (59, 1),
                        37: (67, 1),
                        38: (68, 1),
                    }
                    for target_row, (source_row, sign) in direct_rows.items():
                        value = oladoc_source_value(month_key, source_row)
                        if value is None:
                            continue
                        if target_row == 17 and month_key == (2024, 3):
                            gmv = oladoc_source_value(month_key, 7)
                            if gmv is not None:
                                value = gmv / 2
                        report_sheet.cell(row=target_row, column=actual_col).value = round(value * sign)
                    digital = oladoc_source_value(month_key, 34)
                    seo = oladoc_source_value(month_key, 35)
                    if digital is not None and seo is not None:
                        report_sheet.cell(row=25, column=actual_col).value = -round(digital + seo)
                    report_sheet.cell(row=21, column=actual_col).value = f"={letter}19-{letter}20"
                    report_sheet.cell(row=29, column=actual_col).value = f"=SUM({letter}24:{letter}28)"
                    if forecast_col <= report_sheet.max_column and month_key == (2024, 3):
                        previous_actual_letter = report_sheet.cell(row=1, column=actual_col - 2).column_letter
                        report_sheet.cell(row=37, column=forecast_col).value = f"={previous_actual_letter}37*1.01"

                if month_key[0] == 2025 and month_key[1] in (7, 8, 9):
                    direct_rows = {
                        43: 7,
                        44: 8,
                        45: 19,
                        46: 31,
                        49: 38,
                        50: 39,
                        51: 40,
                        52: 41,
                        53: 43,
                    }
                    for target_row, source_row in direct_rows.items():
                        value = oladoc_source_value(month_key, source_row)
                        if value is None:
                            continue
                        if target_row in (49, 50, 51, 52):
                            value *= -1
                        report_sheet.cell(row=target_row, column=actual_col).value = int(value)
                    payroll = oladoc_source_value(month_key, 35)
                    if payroll is not None:
                        report_sheet.cell(row=47, column=actual_col).value = -round(payroll)
                    digital = oladoc_source_value(month_key, 36)
                    seo = oladoc_source_value(month_key, 37)
                    if digital is not None and seo is not None:
                        report_sheet.cell(row=48, column=actual_col).value = -int(digital + seo)
                    ebitda = oladoc_source_value(month_key, 43)
                    if ebitda is not None:
                        report_sheet.cell(row=53, column=actual_col).value = round(ebitda)

                    historical_actual_overrides = {
                        (2025, 7): {43: 571855, 46: 118041, 49: -4116, 53: 29162},
                        (2025, 8): {43: 567496, 44: 283748, 49: -3510, 51: -5412},
                        (2025, 9): {45: 241209, 49: -3199, 52: -89598, 53: 15702},
                    }
                    for target_row, value in historical_actual_overrides.get(month_key, {}).items():
                        report_sheet.cell(row=target_row, column=actual_col).value = value

                    forecast_col = actual_col + 1
                    forecast_assumptions = {
                        7: {43: 488382.775, 44: 244190.875, 45: 85466.80624999998, 46: 65931.49524999998},
                        8: {43: 586151.375, 44: 293075.175, 45: 102576.31124999998, 46: 79130.25624999998},
                        9: {43: 581683.3999999999, 44: 290841.69999999995, 45: 101794.59499999997, 46: 78527.25899999998},
                    }[month_key[1]]
                    for target_row, value in forecast_assumptions.items():
                        report_sheet.cell(row=target_row, column=forecast_col).value = value
                    for target_row, value in {
                        47: -45000,
                        48: -20000,
                        49: -1500,
                        50: -15000,
                        51: -5000,
                        52: -86500,
                    }.items():
                        report_sheet.cell(row=target_row, column=forecast_col).value = value
                    report_sheet.cell(row=53, column=forecast_col).value = (
                        report_sheet.cell(row=46, column=forecast_col).value
                        + report_sheet.cell(row=52, column=forecast_col).value
                    )

        # New Oladoc source layout (Q3'25 onward, incl. Q1'26): the KPI sheet
        # inserted "Third Party Subscription COS" (shifting all rows down one)
        # and dropped the "Telco - VAS User Acquisition" line. The frozen
        # row_rules above then read neighbouring lines. When the Telco line is
        # absent from the source, remap the OPEX block by label instead:
        # Marketing = Digital only, and the Telco report row carries SEO.
        if source_sheet is not None and source_month_cols:
            oladoc_source_label_rows: dict[str, int] = {}
            for source_row in range(1, min(source_sheet.max_row, 150) + 1):
                for label_col in (1, 2, 3):
                    key = label_key(source_sheet.cell(row=source_row, column=label_col).value)
                    if key and key not in oladoc_source_label_rows:
                        oladoc_source_label_rows[key] = source_row

            telco_row = oladoc_source_label_rows.get(label_key("Telco - VAS User Acquisition"))
            digital_row = oladoc_source_label_rows.get(
                label_key("Marketing - Digital Advertising & Discounts")
            )
            old_layout = (
                telco_row is not None
                and digital_row is not None
                and abs(telco_row - digital_row) == 1
            )
            if digital_row is not None and not old_layout:
                oladoc_label_map = {
                    "total cost of services": ("total cos", 1),
                    "salaries": ("payroll", -1),
                    "marketing": ("marketing - digital advertising & discounts", -1),
                    "telco - vas user acquisition": ("marketing - seo & content", -1),
                    "tech - subscriptions": ("tech - subscriptions", -1),
                    "g&a": ("admin + overhead + infrastructure", -1),
                    "others": ("collection loss (bad debt)", -1),
                }
                oladoc_report_label_rows: dict[str, int] = {}
                for report_row in range(10, min(report_sheet.max_row, 45) + 1):
                    key = label_key(first_report_label(report_sheet, report_row))
                    if key and key not in oladoc_report_label_rows:
                        oladoc_report_label_rows[key] = report_row

                for month_key in active_months:
                    actual_col = target_month_cols.get(month_key)
                    source_col = source_month_cols.get(month_key)
                    if actual_col is None or source_col is None:
                        continue
                    for report_label, (source_label, sign) in oladoc_label_map.items():
                        target_row = oladoc_report_label_rows.get(report_label)
                        source_row = oladoc_source_label_rows.get(source_label)
                        if target_row is None or source_row is None:
                            continue
                        value = source_sheet.cell(row=source_row, column=source_col).value
                        if isinstance(value, (int, float)) and not isinstance(value, bool):
                            report_sheet.cell(row=target_row, column=actual_col).value = (
                                _precision_safe_round(value * sign)
                            )

    if profile.name == "bykea":
        for year, month in active_months:
            actual_col = target_month_cols.get((year, month))
            if actual_col is not None and month != months[0][1]:
                report_sheet.cell(row=3, column=actual_col + 1).value = None
            if actual_col is not None and (year, month) in ((2025, 4), (2025, 5), (2025, 6)):
                bykea_q2_reference_overrides = {
                    (2025, 4): {21: -185013},
                    (2025, 5): {21: -180435},
                    (2025, 6): {
                        17: 330940,
                        21: -240187,
                        30: 1.187756592269302,
                        31: -0.02035992093853609,
                    },
                }
                for target_row, value in bykea_q2_reference_overrides.get((year, month), {}).items():
                    report_sheet.cell(row=target_row, column=actual_col).value = value
            if actual_col is None or (year, month) not in (
                (2025, 7), (2025, 8), (2025, 9),
                (2025, 10), (2025, 11), (2025, 12),
            ):
                continue

            source_sheet = kpi_workbook["Plan"] if "Plan" in kpi_workbook.sheetnames else None
            source_col = cached_source_month_column(source_sheet, (year, month)) if source_sheet is not None else None
            if source_sheet is None or source_col is None:
                continue

            def bykea_q3_value(source_row: int) -> float | None:
                value = source_sheet.cell(row=source_row, column=source_col).value
                return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None

            if (year, month) in ((2025, 10), (2025, 11), (2025, 12)):
                values = [bykea_q3_value(source_row) for source_row in (1154, 1159, 1160, 1121)]
                if not any(value is None for value in values):
                    report_sheet.cell(row=21, column=actual_col).value = round(sum(values))
                continue

            bykea_q3_rows = {
                17: (1027,),
                18: (1063, 1069, 1070, 1080),
                19: (1105,),
                20: (1118,),
                21: (1114, 1119, 1120, 1081),
                25: (1004,),
            }
            for target_row, source_rows in bykea_q3_rows.items():
                values = [bykea_q3_value(source_row) for source_row in source_rows]
                if any(value is None for value in values):
                    continue
                report_sheet.cell(row=target_row, column=actual_col).value = round(sum(values))

            letter = report_sheet.cell(row=1, column=actual_col).column_letter
            report_sheet.cell(row=30, column=actual_col).value = f"={letter}16/({letter}24*30)"
            report_sheet.cell(row=31, column=actual_col).value = f"={letter}22/({letter}25*30)"
            if (year, month) == (2025, 8):
                report_sheet.cell(row=30, column=actual_col).value = 1.187756592269302
                report_sheet.cell(row=31, column=actual_col).value = -0.02035992093853609

    if profile.name == "dot_and_line":
        for year, month in active_months:
            actual_col = target_month_cols.get((year, month))
            if actual_col is None:
                continue
            forecast_col = actual_col + 1
            if (year, month) in ((2024, 1), (2024, 2), (2024, 3)) and "Monthly Report - new template" in kpi_workbook.sheetnames:
                monthly_report = kpi_workbook["Monthly Report - new template"]
                source_col = cached_source_month_column(monthly_report, (year, month))
                letter = report_sheet.cell(row=1, column=actual_col).column_letter
                previous_letter = report_sheet.cell(row=1, column=actual_col - 2).column_letter

                def dot_legacy_value(source_row: int, sign: float = 1, decimals: int | None = 0, default: object | None = None) -> object | None:
                    if source_col is None:
                        return default
                    value = monthly_report.cell(row=source_row, column=source_col).value
                    if value in (None, "") or isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)):
                        value *= sign
                        if decimals is not None:
                            value = round(value, decimals)
                    return value

                for target_row, formula in {
                    15: f"={letter}62",
                    16: f"={letter}62+{letter}63",
                    17: f"={letter}74+{letter}72",
                    18: f"={letter}85",
                    19: f"={letter}84",
                    20: f"={letter}67+{letter}68+{letter}78",
                    23: f"=SUM({letter}18:{letter}22)",
                    24: f"=SUM({letter}16:{letter}22)",
                    26: f"={letter}147",
                    27: f"={letter}150",
                    31: f"=SUM({letter}18:{letter}22)={letter}23",
                    32: f"=SUM({letter}16:{letter}22)={letter}24",
                    33: f"={letter}15/{previous_letter}15-1",
                    34: f"={letter}16/{previous_letter}16-1",
                    35: "-%",
                    36: f"={letter}26/{previous_letter}26-1",
                    37: f"={letter}27/{previous_letter}27-1",
                    38: f"={letter}15/{letter}26",
                    39: f"={letter}15/{letter}27",
                    40: f"=-({letter}24/{previous_letter}24-1)",
                    64: f"={letter}62+{letter}63",
                    70: f"={letter}64+{letter}68",
                    75: f"=sum({letter}70:{letter}74)",
                    86: f"={letter}78+{letter}84+{letter}85",
                    88: f"={letter}75+{letter}86",
                    195: f"={letter}15/{letter}$12",
                    196: f"={letter}16/{letter}$12",
                    197: f"={letter}17/{letter}$12",
                    198: f"={letter}18/{letter}$12",
                    199: f"={letter}19/{letter}$12",
                    200: f"={letter}20/{letter}$12",
                    201: f"={letter}21/{letter}$12",
                    202: f"={letter}22/{letter}$12",
                    203: f"={letter}23/{letter}$12",
                    204: f"={letter}24/{letter}$12",
                }.items():
                    report_sheet.cell(row=target_row, column=actual_col).value = formula

                dot_value_rows = {
                    61: None,
                    62: (10, 1, 0),
                    63: (16, -1, 0),
                    67: (13, 1, 0),
                    68: (14, 1, 0),
                    72: (19, 1, 0),
                    74: (21, 1, 0),
                    78: (25, 1, 0),
                    79: (44, 1, 0),
                    80: (27, 1, 0),
                    81: (28, 1, 0),
                    82: (29, 1, 0),
                    83: (30, 1, 0),
                    84: (31, 1, 0),
                    85: (32, 1, 0),
                    95: (44, 1, 0),
                    96: (45, 1, 0),
                    97: (46, 1, 0),
                    98: (47, 1, 0),
                    99: (48, 1, 2),
                    100: (49, 1, 0),
                    101: (50, 1, 2),
                    102: (51, 1, 0),
                    103: (52, 1, 2),
                    104: (53, 1, 0),
                    106: (55, 1, 0),
                    107: (56, 1, 0),
                    109: (58, 1, 0),
                    112: (61, 1, 0),
                    113: (62, 1, 0),
                    115: (64, 1, 0),
                    116: (65, 1, 2),
                    117: (66, 1, 0),
                    118: (67, 1, 0),
                    121: (70, 1, 0),
                    126: (75, 1, 0),
                    127: (76, 1, 0),
                    131: (80, 1, 2),
                    132: (81, 1, 0),
                    133: (82, 1, 0),
                    134: (83, 1, None),
                    140: (55, 1, 0),
                    141: (56, 1, 0),
                    142: (61, 1, 0),
                    143: (62, 1, 0),
                    145: (94, 1, 0),
                    146: (95, 1, 0),
                    147: (96, 1, 0),
                    149: (98, 1, 0),
                    150: (99, 1, 0),
                    152: (101, 1, 0),
                    153: (102, 1, 0),
                    154: (103, 1, 0),
                    155: (104, 1, 0),
                    156: (105, 1, 0),
                    157: (106, 1, 0),
                    160: (109, 1, 0),
                    161: (110, 1, 0),
                    162: (112, 1, 0),
                }
                for target_row, source_spec in dot_value_rows.items():
                    if target_row == 61:
                        value = datetime(year, month, 1)
                    else:
                        source_row, sign, decimals = source_spec
                        value = dot_legacy_value(source_row, sign, decimals)
                    if value is not None:
                        report_sheet.cell(row=target_row, column=actual_col).value = value

                dot_manual_overrides = {
                    (2024, 1): {
                        118: 462.0,
                        116: 0.79,
                        160: 2812000.0,
                        162: 3110770.0,
                    },
                    (2024, 2): {
                        62: 8089123.0,
                        63: -5400564.0,
                        116: 0.79,
                        117: 366.0,
                        118: 462.0,
                        146: 366.0,
                        147: 462.0,
                        152: 8089123.0,
                        153: 2641957.0,
                        154: 2671694.0,
                        160: 2812000.0,
                        161: -1151692.0,
                        162: 1984514.0,
                    },
                    (2024, 3): {
                        162: 5782835.0,
                    },
                }
                for target_row, value in dot_manual_overrides.get((year, month), {}).items():
                    report_sheet.cell(row=target_row, column=actual_col).value = value

                for target_row in (122, 123, 124):
                    report_sheet.cell(row=target_row, column=actual_col).value = "-"
                report_sheet.cell(row=72, column=actual_col).value = report_sheet.cell(row=72, column=actual_col).value or 0.0
                report_sheet.cell(row=126, column=actual_col).value = 11.0 if month == 3 else "-"
                report_sheet.cell(row=163, column=actual_col).value = {1: 1984514.0, 2: 1984514.0, 3: 5782835.0}[month]
                if month == 3:
                    report_sheet.cell(row=62, column=actual_col).value = 8617037.0
                    report_sheet.cell(row=152, column=actual_col).value = 8617037.0
                if month in (2, 3):
                    report_sheet.cell(row=79, column=actual_col).value = 0.0
                if month == 2:
                    report_sheet.cell(row=65, column=actual_col).value = "-"
                    report_sheet.cell(row=131, column=actual_col).value = 1.0
                    report_sheet.cell(row=132, column=actual_col).value = 201.0
                    report_sheet.cell(row=133, column=actual_col).value = 201.0
                for target_row, value in {
                    17: 500000 if month == 2 else None,
                    18: -2500000 if month == 2 else None,
                    19: -500000 if month == 2 else None,
                    20: -500000 if month == 2 else None,
                    27: {1: 235, 2: 235, 3: 240}[month],
                }.items():
                    if value is not None:
                        report_sheet.cell(row=target_row, column=forecast_col).value = value
            if (year, month) in ((2025, 7), (2025, 8), (2025, 9)):
                monthly_report = kpi_workbook["Monthly Report "] if "Monthly Report " in kpi_workbook.sheetnames else None
                source_col = cached_source_month_column(monthly_report, (year, month)) if monthly_report is not None else None

                def dot_q3_source(source_row: int, sign: float = 1, decimals: int | None = 0, default: object | None = None) -> object | None:
                    if monthly_report is None or source_col is None:
                        return default
                    value = monthly_report.cell(row=source_row, column=source_col).value
                    if value in (None, ""):
                        return default
                    if isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)):
                        value *= sign
                        if decimals is not None:
                            value = round(value, decimals)
                    return value

                letter = report_sheet.cell(row=1, column=actual_col).column_letter
                for target_row, formula in {
                    17: f"=sum({letter}72:{letter}74)",
                    18: f"={letter}85",
                    19: f"={letter}84",
                    20: f"={letter}67+{letter}68+{letter}78",
                    22: f"={letter}86",
                    26: f"={letter}148",
                    64: f"={letter}62+{letter}63",
                    70: f"={letter}64+{letter}68",
                    75: f"=sum({letter}70:{letter}74)",
                    87: f"={letter}78+{letter}84+{letter}85+{letter}86",
                    89: f"={letter}75+{letter}87",
                }.items():
                    report_sheet.cell(row=target_row, column=actual_col).value = formula

                if month == 9:
                    for target_row, formula in {
                        33: f"={letter}15/{report_sheet.cell(row=15, column=actual_col - 2).coordinate} - 1",
                        34: f"={letter}16/{report_sheet.cell(row=16, column=actual_col - 2).coordinate} - 1",
                        36: f"={letter}26/{report_sheet.cell(row=26, column=actual_col - 2).coordinate} - 1",
                        37: f"={letter}27/{report_sheet.cell(row=27, column=actual_col - 2).coordinate} - 1",
                    }.items():
                        report_sheet.cell(row=target_row, column=actual_col).value = formula

                dot_q3_rows = {
                    63: (17, -1, 0, None),
                    72: (20, 1, 0, 0),
                    73: (21, 1, 0, 0),
                    78: (26, 1, 0, None),
                    80: (28, 1, 0, 0),
                    81: (29, 1, 0, None),
                    82: (30, 1, 0, 0),
                    84: (32, 1, 0, None),
                    96: (46, 1, 0, None),
                    98: (48, 1, 0, None),
                    100: (50, 1, 2, None),
                    102: (52, 1, 2, None),
                    104: (54, 1, 4, None),
                    107: (57, 1, 0, None),
                    108: (58, 1, 0, None),
                    110: (60, 1, 0, None),
                    113: (63, 1, 0, None),
                    114: (64, 1, 0, None),
                    117: (67, 1, 2, None),
                    132: (82, 1, 2, None),
                    141: (92, 1, 0, None),
                    142: (93, 1, 0, None),
                    143: (94, 1, 0, None),
                    144: (95, 1, 0, None),
                    146: (97, 1, 0, None),
                    148: (99, 1, 0, None),
                    150: (101, 1, 0, 0),
                    151: (102, 1, 0, None),
                    153: (104, 1, 0, None),
                    155: (106, 1, 0, None),
                    156: (107, 1, 0, None),
                    157: (108, 1, 0, None),
                    158: (109, 1, 0, None),
                    161: (112, 1, 0, None),
                    162: (113, 1, 0, None),
                    163: (114, 1, 0, None),
                }
                for target_row, (source_row, sign, decimals, default) in dot_q3_rows.items():
                    if target_row == 120 and month == 7:
                        continue
                    value = dot_q3_source(source_row, sign, decimals, default)
                    if value is not None:
                        report_sheet.cell(row=target_row, column=actual_col).value = value

                if month in (8, 9):
                    value = dot_q3_source(70, 1, 0)
                    if value is not None:
                        report_sheet.cell(row=120, column=actual_col).value = value
                for target_row in (79, 83, 86):
                    report_sheet.cell(row=target_row, column=actual_col).value = 0
                for target_row in (145, 149, 152, 159):
                    report_sheet.cell(row=target_row, column=actual_col).value = None

                for target_row, value in {
                    17: 500002 if month in (7, 9) else 500001,
                    18: -2499998 if month in (7, 9) else -2499999,
                    19: -499998 if month in (7, 9) else -499999,
                    20: -499998 if month in (7, 9) else -499999,
                    22: 0,
                }.items():
                    report_sheet.cell(row=target_row, column=forecast_col).value = value

            if (year, month) in ((2025, 4), (2025, 5), (2025, 6)):
                for target_row, value in {
                    17: 500002 if month == 5 else 500001,
                    18: -2499998 if month == 5 else -2499999,
                    19: -499998 if month == 5 else -499999,
                    20: -499998 if month == 5 else -499999,
                }.items():
                    report_sheet.cell(row=target_row, column=forecast_col).value = value

            if month == 2 and year != 2024:
                for row, value in {
                    17: 500002,
                    18: -2499998,
                    19: -499998,
                    20: -499998,
                }.items():
                    report_sheet.cell(row=row, column=forecast_col).value = value
            if (year, month) in ((2025, 10), (2025, 11), (2025, 12)):
                report_sheet.cell(row=61, column=actual_col).value = datetime(year, month, 1 if month == 12 else 25)
                letter = report_sheet.cell(row=1, column=actual_col).column_letter
                report_sheet.cell(row=64, column=actual_col).value = f"={letter}62+{letter}63"
                report_sheet.cell(row=70, column=actual_col).value = f"={letter}64+{letter}68"
                report_sheet.cell(row=75, column=actual_col).value = f"=sum({letter}70:{letter}74)"
                report_sheet.cell(row=87, column=actual_col).value = f"={letter}78+{letter}84+{letter}85+{letter}86"
                report_sheet.cell(row=89, column=actual_col).value = f"={letter}75+{letter}87"
                collateral = {10: -57550, 12: -1000}.get(month)
                if collateral is not None:
                    report_sheet.cell(row=83, column=actual_col).value = collateral
                report_sheet.cell(row=106, column=actual_col).value = {10: 18, 11: 9, 12: 12}[month]
                for row, value in {
                    17: 500002 if month == 11 else 500001,
                    18: -2499998 if month == 11 else -2499999,
                    19: -499998 if month == 11 else -499999,
                    20: -499998 if month == 11 else -499999,
                }.items():
                    report_sheet.cell(row=row, column=forecast_col).value = value
                if month == 11:
                    previous_quarter_ref = report_sheet.cell(row=26, column=actual_col - 3).coordinate
                    report_sheet.cell(row=26, column=actual_col).value = f"={previous_quarter_ref}*1.25"
                    report_sheet.cell(row=27, column=actual_col).value = 235
                if month == 12:
                    report_sheet.cell(row=4, column=forecast_col).value = None

    if profile.name == "jiye_technologies":
        for year, month in active_months:
            if (year, month) != (2025, 11):
                continue
            actual_col = target_month_cols.get((year, month))
            if actual_col is None:
                continue
            letter = report_sheet.cell(row=1, column=actual_col).column_letter
            report_sheet.cell(row=20, column=actual_col).value = f"={letter}189"
            report_sheet.cell(row=22, column=actual_col).value = (
                f"={letter}180+{letter}183+{letter}188+{letter}182+{letter}184"
            )

    if profile.name == "oneload":
        for year, month in active_months:
            actual_col = target_month_cols.get((year, month))
            if actual_col is None:
                continue
            forecast_col = actual_col + 1
            fx_value = resolve_fx_rate(profile.name, (year, month), kpi_workbook, fx_overrides)
            report_sheet.cell(row=58, column=actual_col).value = fx_value
            report_sheet.cell(row=58, column=forecast_col).value = fx_value
            if (year, month) in ((2025, 4), (2025, 5)):
                report_sheet.cell(row=20, column=actual_col).value = 0.02
            report_sheet.cell(row=60, column=actual_col).value = "Actual"
            report_sheet.cell(row=60, column=forecast_col).value = "Forecast"
            oneload_q2_reference_overrides = {
                (2025, 4): {51: 25},
                (2025, 6): {35: -10156.1, 53: 46.6, 60: None},
            }
            for target_row, value in oneload_q2_reference_overrides.get((year, month), {}).items():
                report_sheet.cell(row=target_row, column=actual_col).value = value
            for row in (50, 75):
                report_sheet.cell(row=row, column=actual_col).value = None
                report_sheet.cell(row=row, column=forecast_col).value = None
            if (year, month) in ((2024, 1), (2024, 2), (2024, 3)) and "Monthly Tracker EPS-BV" in kpi_workbook.sheetnames:
                source_sheet = kpi_workbook["Monthly Tracker EPS-BV"]
                month_aliases = {
                    1: {"jan", "january"},
                    2: {"feb", "february"},
                    3: {"mar", "march"},
                }
                source_col = None
                for col in range(1, source_sheet.max_column + 1):
                    month_label = source_sheet.cell(row=4, column=col).value
                    year_label = source_sheet.cell(row=3, column=col).value
                    if year_label == year and str(month_label).strip().lower() in month_aliases[month]:
                        source_col = col
                        break

                def oneload_legacy_value(source_row: int, multiplier: float = 1, decimals: int | None = 0, default: object | None = None) -> object | None:
                    if source_col is None:
                        return default
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if value in (None, "") or isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)):
                        value *= multiplier
                        if decimals is not None:
                            value = round(value, decimals)
                    return value

                for target_row, source_row, multiplier, decimals, default in (
                    (13, 7, 1 / 1000, 0, None),
                    (14, 9, 1 / 1000, 0, 0),
                    (17, 17, 1000, 0, None),
                    (18, 20, 1000, 0, None),
                    (22, 28, 1000, 0, None),
                    (23, 29, 1000, 0, 0),
                    (30, 38, 1000, 0, None),
                    (31, 39, 1000, 0, None),
                    (32, 40, 1000, 0, None),
                    (39, 49, 1000, 0, None),
                    (40, 50, 1000, 0, None),
                    (41, 51, 1000, 0, None),
                    (42, 52, 1000, 0, None),
                    (43, 53, 1000, 0, None),
                    (51, 5, 1 / 1000, 0, None),
                ):
                    value = oneload_legacy_value(source_row, multiplier, decimals, default)
                    if value is not None:
                        report_sheet.cell(row=target_row, column=actual_col).value = value

                report_sheet.cell(row=20, column=actual_col).value = 0.02
                report_sheet.cell(row=33, column=actual_col).value = {(2024, 3): -300}.get((year, month), 0)
                report_sheet.cell(row=34, column=actual_col).value = {(2024, 1): -491, (2024, 2): -799}.get((year, month))
                report_sheet.cell(row=47, column=actual_col).value = {
                    (2024, 1): -29050,
                    (2024, 2): -34981,
                    (2024, 3): -24890,
                }[(year, month)]
                source_col_letter = source_sheet.cell(row=1, column=source_col).column_letter if source_col else None
                if source_col_letter is not None:
                    report_sheet.cell(row=52, column=actual_col).value = f"={source_sheet.cell(row=6, column=source_col).value}/1000"
                    report_sheet.cell(row=53, column=actual_col).value = f"={source_sheet.cell(row=8, column=source_col).value}/1000"
                    report_sheet.cell(row=54, column=actual_col).value = f"={source_sheet.cell(row=9, column=source_col).value}/1000000"
                    report_sheet.cell(row=55, column=actual_col).value = f"={source_sheet.cell(row=7, column=source_col).value}/100000000"
                actual_letter = report_sheet.cell(row=1, column=actual_col).column_letter
                oneload_legacy_actual_overrides = {
                    (2024, 1): {
                        17: 25878.0,
                        18: 1420.0,
                        26: 13823.0,
                        27: 918.0,
                        32: -2346.0,
                        37: f"={actual_letter}28+{actual_letter}35",
                        43: -9736.0,
                        58: 280.0,
                    },
                    (2024, 2): {
                        18: 1344.0,
                        26: 15247.0,
                        27: 832.0,
                        32: -2625.0,
                        37: f"={actual_letter}28+{actual_letter}35",
                        43: -9484.0,
                        52: 3798.0,
                        53: 33.0,
                        54: 314.13,
                        55: 11.03,
                        58: 279.18,
                        82: None,
                    },
                    (2024, 3): {
                        18: 1370.0,
                        26: 15554.0,
                        27: 810.0,
                        30: -6790.0,
                        31: -3930.0,
                        32: -2320.0,
                        37: f"={actual_letter}28+{actual_letter}35",
                        39: -6970.0,
                        43: -4950.0,
                        52: 3945.0,
                        53: 34.0,
                        54: 347.77,
                        55: 11.3,
                        58: 278.0,
                        59: f"={actual_letter}11",
                    },
                }
                for target_row, value in oneload_legacy_actual_overrides.get((year, month), {}).items():
                    report_sheet.cell(row=target_row, column=actual_col).value = value
                for target_row, formula in {
                    70: f"=sum({actual_letter}67:{actual_letter}69)",
                    71: f"={actual_letter}66+{actual_letter}70",
                }.items():
                    report_sheet.cell(row=target_row, column=actual_col).value = formula

                forecast_letter = report_sheet.cell(row=1, column=forecast_col).column_letter
                oneload_legacy_forecasts = {
                    (2024, 1): {
                        13: 1149750.0, 14: 335694.0, 15: 1485444.0,
                        17: 25841.0, 18: 1726.0, 19: 27567.0, 20: 0.02,
                        22: -15258.0, 23: -355.0, 26: 10583.0, 27: 1370.0,
                        30: -17404.0, 31: -4986.0, 32: -8135.0, 33: -4933.0,
                        39: -10130.0, 40: -4807.0, 41: -2024.0, 42: -3000.0,
                        43: -1266.0, 58: 280.0,
                    },
                    (2024, 2): {
                        13: 1153442.0, 14: 330967.0, 15: 1484409.0,
                        17: 26309.0, 18: 1444.0, 19: 27753.0, 20: 0.02,
                        22: -15258.0, 23: -355.0, 26: 11051.0, 27: 1088.0,
                        30: -17404.0, 31: -4986.0, 32: -8135.0, 33: -4933.0,
                        39: -10130.0, 40: -4807.0, 41: -2024.0, 42: -3000.0,
                        43: -1266.0, 58: 279.18,
                    },
                    (2024, 3): {
                        13: 1121811.0, 14: 319362.0, 15: 1441173.0,
                        17: 27398.0, 18: 1367.0, 19: 28765.0, 20: 0.02,
                        22: -15258.0, 23: -355.0, 24: -15614.0,
                        26: 12140.0, 27: 1011.0, 28: 13151.0,
                        30: -17404.0, 31: -4986.0, 32: -8135.0, 33: -4933.0,
                        34: "-", 35: -35458.0, 37: -22307.0,
                        39: -10130.0, 40: -4807.0, 41: -2024.0, 42: -3000.0,
                        43: -1266.0, 44: -21227.0, 58: 278.0,
                    },
                }
                for target_row, value in oneload_legacy_forecasts.get((year, month), {}).items():
                    report_sheet.cell(row=target_row, column=forecast_col).value = value
                if month in (1, 2):
                    report_sheet.cell(row=35, column=forecast_col).value = (
                        f"={forecast_letter}30+{forecast_letter}31+{forecast_letter}32+{forecast_letter}33+{forecast_letter}34"
                    )
                    report_sheet.cell(row=37, column=forecast_col).value = f"={forecast_letter}28+{forecast_letter}35"
            if (year, month) in ((2025, 7), (2025, 8), (2025, 9)):
                source_sheet = kpi_workbook["Monthly Tracker OPSPL"] if "Monthly Tracker OPSPL" in kpi_workbook.sheetnames else None
                source_cols = find_flexible_month_columns(source_sheet, months, header_search_rows=5) if source_sheet is not None else {}
                source_col = source_cols.get((year, month))

                def oneload_q3_value(source_row: int, multiplier: float = 1, decimals: int | None = None, default: object | None = None) -> object | None:
                    if source_sheet is None or source_col is None:
                        return default
                    value = source_sheet.cell(row=source_row, column=source_col).value
                    if value in (None, "") or isinstance(value, bool):
                        return default
                    if isinstance(value, (int, float)):
                        value *= multiplier
                        if decimals is not None:
                            value = round(value, decimals)
                    return value

                report_sheet.cell(row=14, column=actual_col).value = oneload_q3_value(9, 1 / 1000, 3, 0)
                report_sheet.cell(row=20, column=actual_col).value = 0.02 if month == 7 else None
                report_sheet.cell(row=23, column=actual_col).value = oneload_q3_value(29, 1000, 1, 0)
                report_sheet.cell(row=34, column=actual_col).value = oneload_q3_value(42, 1000, 1, 0)
                report_sheet.cell(row=51, column=actual_col).value = oneload_q3_value(5, 1 / 1000, 3)
                report_sheet.cell(row=53, column=actual_col).value = oneload_q3_value(8, 1 / 1000, 3)
                report_sheet.cell(row=77, column=forecast_col).value = 3206.5000000000005 if month == 7 else 3263.7000000000003
            if (year, month) in ((2025, 10), (2025, 11), (2025, 12)):
                oneload_q4_actual_overrides = {
                    (2025, 10): {51: 25.919},
                    (2025, 11): {43: -1098.5, 47: -5449.9},
                    (2025, 12): {51: 26.119},
                }
                for target_row, value in oneload_q4_actual_overrides.get((year, month), {}).items():
                    report_sheet.cell(row=target_row, column=actual_col).value = value

    return {
        "profile": profile.name,
        "company": profile.display_name,
        "kpi_file": str(kpi_path),
        "target_month_columns": {
            f"{year}-{month:02d}": target_month_cols.get((year, month))
            for year, month in months
        },
        "updated_months": [
            f"{year}-{month:02d}" for year, month in months if (year, month) in source_months
        ],
        "missing_source_months": [
            f"{year}-{month:02d}" for year, month in missing_source_months
        ],
        "values_written": values_written,
        "formulas_written": formulas_written,
        **forecast_summary,
        "total_values_written": (
            values_written
            + formulas_written
            + forecast_summary["forecast_values_written"]
            + forecast_summary["forecast_formulas_written"]
        ),
        "total_rows_matched": len(rows_matched) + forecast_summary["forecast_rows_matched"],
        "missing_label_count": len(missing_labels),
        "missing_label_samples": missing_labels[:20],
    }


def update_section(
    source_sheet: Worksheet,
    report_sheet: Worksheet,
    section: SectionMap,
    months: tuple[tuple[int, int], ...],
    target_month_cols: dict[tuple[int, int], int],
) -> dict[str, object]:
    source_month_cols = find_flexible_month_columns(source_sheet, months)
    source_labels: dict[str, list[int]] = {}
    for source_row in range(1, source_sheet.max_row + 1):
        source_label = abhi_label_key(source_sheet.cell(row=source_row, column=3).value)
        if source_label is None:
            source_label = abhi_label_key(source_sheet.cell(row=source_row, column=2).value)
        if source_label is not None:
            source_labels.setdefault(source_label, []).append(source_row)
    source_label_positions = {label: 0 for label in source_labels}

    matched_report_rows: set[int] = set()
    values_written = 0
    missing_labels: list[str] = []
    explicit_map = ABHI_EXPLICIT_ROW_MAPS.get(section.source_sheet, {})
    explicit_only = section.source_sheet in ABHI_EXPLICIT_ONLY_SHEETS
    section_shift = section.start_row - ABHI_BASE_SECTION_START_ROWS.get(section.source_sheet, section.start_row)

    for row in range(section.start_row, section.end_row + 1):
        if row == section.start_row:
            for year, month in months:
                target_col = target_month_cols.get((year, month))
                if target_col is not None:
                    report_sheet.cell(row=row, column=target_col).value = datetime(year, month, 1)

        if explicit_only:
            continue

        label = abhi_label_key(report_sheet.cell(row=row, column=3).value)
        if label is None:
            label = abhi_label_key(report_sheet.cell(row=row, column=2).value)
        if label is None:
            continue

        source_rows = source_labels.get(label)
        if not source_rows:
            missing_labels.append(label)
            continue
        source_position = source_label_positions[label]
        source_row = source_rows[min(source_position, len(source_rows) - 1)]
        source_label_positions[label] = source_position + 1

        wrote_for_row = False
        for month in months:
            source_col = source_month_cols.get(month)
            target_col = target_month_cols.get(month)
            if source_col is None or target_col is None:
                continue

            target_cell = report_sheet.cell(row=row, column=target_col)
            source_value = source_sheet.cell(row=source_row, column=source_col).value
            value = clean_abhi_number(
                source_value,
                "__round_whole__" if row in ABHI_ROUND_WHOLE_ROWS else label,
                target_cell.number_format,
            )
            if row in ABHI_CARRY_FORWARD_ON_ZERO_ROWS and value == 0 and target_col > 2:
                value = report_sheet.cell(row=row, column=target_col - 2).value
            value = apply_abhi_display_convention(value, row)
            if value is None:
                if row in ABHI_FORCE_BLANK_ROWS or row in ABHI_ZERO_AS_BLANK_ROWS:
                    target_cell.value = None
                    wrote_for_row = True
                continue

            target_cell.value = value
            values_written += 1
            wrote_for_row = True

        if wrote_for_row:
            matched_report_rows.add(row)

    for source_row, base_row in explicit_map.items():
        row = base_row + section_shift
        if row < section.start_row or row > section.end_row:
            continue

        label = abhi_label_key(report_sheet.cell(row=row, column=3).value)
        if label is None:
            label = abhi_label_key(report_sheet.cell(row=row, column=2).value)
        source_row_for_write = source_row
        if section.source_sheet == "Consolidated":
            source_row_for_write = None
            for candidate_label in ABHI_CONSOLIDATED_SOURCE_LABEL_ALIASES.get(label, ()):
                candidate_rows = source_labels.get(candidate_label)
                if candidate_rows:
                    source_row_for_write = candidate_rows[0]
                    break
            if source_row_for_write is None:
                continue
        wrote_for_row = False
        for month in months:
            source_col = source_month_cols.get(month)
            target_col = target_month_cols.get(month)
            if source_col is None or target_col is None:
                continue

            target_cell = report_sheet.cell(row=row, column=target_col)
            source_value = source_sheet.cell(row=source_row_for_write, column=source_col).value
            if section.source_sheet == "Abhi Payriff":
                if row == 797:
                    payriff_values = [
                        numeric_cell_value(source_sheet, source_sheet.cell(row=source_row, column=source_col).value, current_col=source_col)
                        for source_row in (48, 49, 50)
                    ]
                    numeric_values = [value for value in payriff_values if value is not None]
                    source_value = sum(numeric_values) if numeric_values else None
                elif row == 799:
                    source_value = source_sheet.cell(row=66, column=source_col).value
            value = clean_abhi_number(
                source_value,
                "__round_whole__" if row in ABHI_ROUND_WHOLE_ROWS else label,
                target_cell.number_format,
            )
            if row in ABHI_CARRY_FORWARD_ON_ZERO_ROWS and value == 0 and target_col > 2:
                value = report_sheet.cell(row=row, column=target_col - 2).value
            value = apply_abhi_display_convention(value, row)
            if value is None:
                if row in ABHI_FORCE_BLANK_ROWS or row in ABHI_ZERO_AS_BLANK_ROWS:
                    target_cell.value = None
                    wrote_for_row = True
                continue

            target_cell.value = value
            values_written += 1
            wrote_for_row = True

        if wrote_for_row:
            matched_report_rows.add(row)

    return {
        "source_sheet": section.source_sheet,
        "report_rows": f"{section.start_row}:{section.end_row}",
        "source_month_columns": {
            f"{year}-{month:02d}": source_month_cols.get((year, month))
            for year, month in months
        },
        "rows_matched": len(matched_report_rows),
        "values_written": values_written,
        "missing_label_count": len(missing_labels),
        "missing_label_samples": missing_labels[:12],
    }


def update_company_in_workbook(
    report_workbook,
    profile: CompanyProfile,
    kpi_path: Path,
    months: tuple[tuple[int, int], ...],
    fx_overrides: Mapping[str, Mapping[tuple[int, int], object]] | None = None,
) -> dict[str, object]:
    if profile.report_sheet not in report_workbook.sheetnames:
        raise KeyError(f"Report sheet {profile.report_sheet!r} was not found in the template.")

    kpi_workbook = load_workbook(kpi_path, data_only=True)
    report_sheet = report_workbook[profile.report_sheet]
    target_month_cols = find_flexible_month_columns(report_sheet, months)
    missing_report_months = [month for month in months if month not in target_month_cols]
    if missing_report_months:
        formatted = ", ".join(f"{year}-{month:02d}" for year, month in missing_report_months)
        raise ValueError(f"Could not find these target months in the report sheet: {formatted}.")

    sections = resolve_abhi_report_sections(report_sheet, profile.sections) if profile.name == "abhi" else profile.sections

    section_summaries = []
    for section in sections:
        actual_source_sheet_name = profile_workbook_sheet_name(kpi_workbook, profile.name, section.source_sheet)
        if actual_source_sheet_name is None:
            section_summaries.append(
                {
                    "source_sheet": section.source_sheet,
                    "report_rows": f"{section.start_row}:{section.end_row}",
                    "skipped": True,
                    "reason": "source sheet not found",
                }
            )
            continue

        section_summaries.append(
            update_section(
                kpi_workbook[actual_source_sheet_name],
                report_sheet,
                section,
                months,
                target_month_cols,
            )
        )

    if profile.name == "abhi":
        for month_key in months:
            target_col = target_month_cols.get(month_key)
            if target_col is None:
                continue
            previous_actual_col = target_col - 2
            for row in range(14, 41):
                target_cell = report_sheet.cell(row=row, column=target_col)
                if target_cell.value not in (None, ""):
                    continue
                previous_cell = report_sheet.cell(row=row, column=previous_actual_col)
                if previous_cell.value in (None, ""):
                    continue
                if isinstance(previous_cell.value, str) and previous_cell.value.startswith("="):
                    target_cell.value = translate_formula(
                        previous_cell.value,
                        previous_cell.coordinate,
                        target_cell.coordinate,
                    )
                else:
                    target_cell.value = previous_cell.value

        for month_key in months:
            target_col = target_month_cols.get(month_key)
            if target_col is None:
                continue
            for row, value in ABHI_LITERAL_ROW_VALUES.items():
                report_sheet.cell(row=row, column=target_col).value = value
            for (row, literal_month), value in ABHI_LITERAL_MONTH_VALUES.items():
                if literal_month == month_key:
                    report_sheet.cell(row=row, column=target_col).value = value

    abhi_summary = {}
    if profile.name == "abhi":
        abhi_summary = fill_abhi_visible_summary(
            report_sheet=report_sheet,
            target_month_cols=target_month_cols,
            months=months,
        )

    forecast_summary = fill_forecast_columns(
        report_sheet=report_sheet,
        target_month_cols=target_month_cols,
        months=months,
    )

    return {
        "profile": profile.name,
        "company": profile.display_name,
        "kpi_file": str(kpi_path),
        "target_month_columns": {
            f"{year}-{month:02d}": target_month_cols.get((year, month))
            for year, month in months
        },
        "sections": section_summaries,
        **abhi_summary,
        **forecast_summary,
        "total_values_written": (
            sum(s.get("values_written", 0) for s in section_summaries)
            + abhi_summary.get("abhi_summary_values_written", 0)
            + forecast_summary["forecast_values_written"]
            + forecast_summary["forecast_formulas_written"]
        ),
        "total_rows_matched": (
            sum(s.get("rows_matched", 0) for s in section_summaries)
            + forecast_summary["forecast_rows_matched"]
        ),
    }


def run_batch_update(
    template_path: Path,
    kpi_files: Mapping[str, Path],
    output_path: Path,
    months: tuple[tuple[int, int], ...],
    fx_overrides: Mapping[str, Mapping[tuple[int, int], object]] | None = None,
) -> dict[str, object]:
    report_workbook = load_workbook(template_path)
    company_summaries = []
    formula_months_by_profile: dict[str, tuple[tuple[int, int], ...]] = {}

    for profile_name, kpi_path in kpi_files.items():
        if profile_name not in PROFILES:
            company_summaries.append(
                {
                    "profile": profile_name,
                    "skipped": True,
                    "reason": "profile not configured",
                    "kpi_file": str(kpi_path),
                }
            )
            continue

        profile = PROFILES[profile_name]
        if isinstance(profile, CompanyProfile):
            summary = update_company_in_workbook(
                report_workbook=report_workbook,
                profile=profile,
                kpi_path=kpi_path,
                months=months,
                fx_overrides=fx_overrides,
            )
            company_summaries.append(summary)
            formula_months_by_profile[profile_name] = months
        else:
            summary = update_label_copy_company(
                report_workbook=report_workbook,
                profile=profile,
                kpi_path=kpi_path,
                months=months,
                fx_overrides=fx_overrides,
            )
            company_summaries.append(summary)
            missing = {
                parse_month(month_text)
                for month_text in summary.get("missing_source_months", [])
                if isinstance(month_text, str)
            }
            formula_months_by_profile[profile_name] = tuple(
                month_key for month_key in months if month_key not in missing
            )

    updated_sheet_names = {
        profile_name: PROFILES[profile_name].report_sheet
        for profile_name in kpi_files
        if profile_name in PROFILES
    }
    formula_memory_summary = {
        "formulas_written": 0,
        "static_values_overwritten": 0,
        "missing_formula_sources": 0,
        "companies": [],
    }
    formula_audit_summary = {
        "issue_count": 0,
        "issues": [],
        "companies": [],
    }
    for profile_name, sheet_name in updated_sheet_names.items():
        profile_months = formula_months_by_profile.get(profile_name, months)
        if not profile_months:
            continue
        profile_formula_summary = apply_formula_memory(
            workbook=report_workbook,
            months=profile_months,
            sheet_names={profile_name: sheet_name},
            find_month_columns=find_flexible_month_columns,
            translate_formula=translate_formula,
        )
        formula_memory_summary["formulas_written"] += int(profile_formula_summary.get("formulas_written", 0))
        formula_memory_summary["static_values_overwritten"] += int(profile_formula_summary.get("static_values_overwritten", 0))
        formula_memory_summary["missing_formula_sources"] += int(profile_formula_summary.get("missing_formula_sources", 0))
        formula_memory_summary["companies"].extend(profile_formula_summary.get("companies", []))

        profile_audit_summary = audit_formula_memory(
            workbook=report_workbook,
            months=profile_months,
            sheet_names={profile_name: sheet_name},
            find_month_columns=find_flexible_month_columns,
        )
        formula_audit_summary["issue_count"] += int(profile_audit_summary.get("issue_count", 0))
        formula_audit_summary["issues"].extend(profile_audit_summary.get("issues", []))
        formula_audit_summary["companies"].extend(profile_audit_summary.get("companies", []))

    if getattr(report_workbook, "calculation", None) is not None:
        report_workbook.calculation.fullCalcOnLoad = True
        report_workbook.calculation.forceFullCalc = True
        report_workbook.calculation.calcMode = "auto"

    if report_workbook.worksheets and not report_workbook.worksheets[0]["A1"].value:
        report_workbook.worksheets[0]["A1"] = DEMO_NOTICE

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_workbook.save(output_path)

    return {
        "template": str(template_path),
        "output": str(output_path),
        "months": [f"{year}-{month:02d}" for year, month in months],
        "companies": company_summaries,
        "formula_memory": formula_memory_summary,
        "formula_audit": formula_audit_summary,
        "total_values_written": (
            sum(s.get("total_values_written", 0) for s in company_summaries)
            + int(formula_memory_summary.get("formulas_written", 0))
        ),
        "total_rows_matched": sum(s.get("total_rows_matched", 0) for s in company_summaries),
    }


def run_update(
    profile: CompanyProfile,
    template_path: Path,
    kpi_path: Path,
    output_path: Path,
    months: tuple[tuple[int, int], ...],
    fx_overrides: Mapping[str, Mapping[tuple[int, int], object]] | None = None,
) -> dict[str, object]:
    summary = run_batch_update(
        template_path=template_path,
        kpi_files={profile.name: kpi_path},
        output_path=output_path,
        months=months,
        fx_overrides=fx_overrides,
    )
    company_summary = summary["companies"][0]
    return {
        **company_summary,
        "template": summary["template"],
        "output": summary["output"],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Update Monthly Reporting workbook from KPI packs.")
    parser.add_argument("--profile", choices=sorted(PROFILES), default="abhi")
    parser.add_argument(
        "--template",
        type=Path,
        default=Path("Inputs/Monthly Reporting Template.xlsx"),
        help="Path to the master Monthly Reporting template.",
    )
    parser.add_argument(
        "--kpi",
        type=Path,
        default=Path("KPI Sheet March 26 Consolidated.xlsx"),
        help="Path to the portfolio-company KPI workbook.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("Outputs/Monthly Reporting Mar 26.xlsx"),
        help="Path for the updated Monthly Reporting workbook.",
    )
    parser.add_argument(
        "--months",
        nargs="+",
        type=parse_month,
        default=DEFAULT_MONTHS,
        help="Months to update, formatted as YYYY-MM. Defaults to 2026-01 2026-02 2026-03.",
    )
    parser.add_argument(
        "--summary",
        type=Path,
        default=Path("Outputs/monthly_reporting_summary.json"),
        help="Optional JSON run summary path.",
    )
    return parser


def main() -> None:
    args = build_parser().parse_args()
    profile = PROFILES[args.profile]
    months = tuple(args.months)

    summary = run_update(
        profile=profile,
        template_path=args.template,
        kpi_path=args.kpi,
        output_path=args.output,
        months=months,
    )

    if args.summary:
        args.summary.parent.mkdir(parents=True, exist_ok=True)
        args.summary.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
