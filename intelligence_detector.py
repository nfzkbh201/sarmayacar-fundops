from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from monthly_reporting_automation import (
    ABHI_BASE_SECTION_START_ROWS,
    ABHI_EXPLICIT_ROW_MAPS,
    CompanyProfile,
    LabelCopyProfile,
    PROFILES,
    _available_label_keys,
    _profile_expected_label_keys,
    _profile_sheet_candidates,
    _sheet_name_mentions_month,
    _string_mentions_month,
    abhi_label_key,
    abhi_report_label,
    abhi_section_start_row,
    find_bykea_visible_summary_row,
    find_flexible_month_columns,
    infer_profile_source_sheet_name,
    label_key,
    profile_workbook_sheet_name,
    workbook_sheet_name,
)


@dataclass(frozen=True)
class IntelligenceFinding:
    company: str
    severity: str
    check: str
    finding: str
    action: str
    location: str = ""
    suggested_fix: str = ""
    confidence: int = 0
    approval_required: bool = False

    def row(self) -> dict[str, str]:
        return asdict(self)


def _month_label(month_key: tuple[int, int]) -> str:
    year, month = month_key
    return datetime(year, month, 1).strftime("%b %y")


def _format_months(months: list[tuple[int, int]] | tuple[tuple[int, int], ...]) -> str:
    return ", ".join(_month_label(month) for month in months)


def _expected_missing_column(
    found_cols: Mapping[tuple[int, int], int],
    missing_month: tuple[int, int],
) -> str:
    previous_months = [
        month_key
        for month_key, col in found_cols.items()
        if month_key < missing_month and col is not None
    ]
    if not previous_months:
        return "Header row"

    previous_month = max(previous_months)
    previous_col = found_cols[previous_month]
    expected_col = previous_col + 2
    return f"Expected around {get_column_letter(expected_col)} after {_month_label(previous_month)}"


def _sheet_label_keys(sheet, max_rows: int | None = None) -> set[str]:
    labels: set[str] = set()
    row_limit = min(sheet.max_row, max_rows) if max_rows else sheet.max_row
    for row in range(1, row_limit + 1):
        for col in range(1, min(sheet.max_column, 6) + 1):
            key = label_key(sheet.cell(row=row, column=col).value)
            if key:
                labels.add(key)
    return labels


def _abhi_report_labels(report_sheet, start_row: int, end_row: int) -> set[str]:
    labels: set[str] = set()
    ignored_fragments = (
        "financial statements",
        "income statement",
        "balance sheet",
        "payments",
        "summary",
        "check",
    )
    for row in range(start_row, end_row + 1):
        label = abhi_label_key(report_sheet.cell(row=row, column=3).value)
        if label is None:
            label = abhi_label_key(report_sheet.cell(row=row, column=2).value)
        if not label:
            continue
        if any(fragment in label for fragment in ignored_fragments):
            continue
        labels.add(label)
    return labels


def _abhi_source_labels(source_sheet) -> set[str]:
    labels: set[str] = set()
    for row in range(1, source_sheet.max_row + 1):
        label = abhi_label_key(source_sheet.cell(row=row, column=3).value)
        if label is None:
            label = abhi_label_key(source_sheet.cell(row=row, column=2).value)
        if label:
            labels.add(label)
    return labels


def _month_header_columns(sheet, month_key: tuple[int, int]) -> set[int]:
    """All columns whose header mentions the given month (dates or text)."""
    year, month = month_key
    columns: set[int] = set()
    for row in range(1, min(25, sheet.max_row) + 1):
        for col in range(1, min(sheet.max_column, 300) + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, (datetime, date)):
                if value.month == month and (value.year == year or value.day == year % 100):
                    columns.add(col)
            elif _string_mentions_month(value, month_key):
                columns.add(col)
    return columns


def _duplicate_month_findings(
    company: str,
    sheet,
    sheet_name: str,
    months: tuple[tuple[int, int], ...],
) -> list[IntelligenceFinding]:
    findings: list[IntelligenceFinding] = []
    chosen_cols = find_flexible_month_columns(sheet, months)
    for month_key in months:
        columns = _month_header_columns(sheet, month_key)
        if len(columns) <= 1:
            continue
        chosen = chosen_cols.get(month_key)
        column_letters = ", ".join(get_column_letter(col) for col in sorted(columns))
        chosen_letter = get_column_letter(chosen) if chosen else "unknown"
        findings.append(
            IntelligenceFinding(
                company=company,
                severity="Review",
                check="Duplicate month columns",
                finding=(
                    f"'{sheet_name}' has {len(columns)} columns labelled {_month_label(month_key)} "
                    f"(columns {column_letters})."
                ),
                action="Confirm the correct column is used (e.g. actual vs budget vs YTD).",
                location=f"{sheet_name} columns {column_letters}",
                suggested_fix=f"App will read column {chosen_letter} for {_month_label(month_key)}.",
                confidence=70,
                approval_required=False,
            )
        )
    return findings


def _label_rows_map(sheet, max_rows: int | None = 400) -> dict[str, list[int]]:
    labels: dict[str, list[int]] = {}
    row_limit = min(sheet.max_row, max_rows) if max_rows is not None else sheet.max_row
    for row in range(1, row_limit + 1):
        for col in range(1, min(sheet.max_column, 6) + 1):
            key = label_key(sheet.cell(row=row, column=col).value)
            if key:
                rows = labels.setdefault(key, [])
                if row not in rows:
                    rows.append(row)
    return labels


def _row_rule_source_rows(
    sheet_name: str,
    row_rules: Mapping[str, tuple[tuple[str, int, float], ...]] | None,
) -> dict[str, list[int]]:
    source_rows: dict[str, list[int]] = {}
    for metric, rules in (row_rules or {}).items():
        key = label_key(metric)
        if not key:
            continue
        rows = [
            row
            for rule_sheet_name, row, _sign in rules
            if label_key(rule_sheet_name) == label_key(sheet_name)
        ]
        if rows:
            source_rows[key] = rows
    return source_rows


def _source_row_suggestion(sheet, rows: list[int], row_rules: list[int] | None = None) -> str:
    if row_rules:
        visible_rule_rows = [row for row in row_rules if not sheet.row_dimensions[row].hidden]
        chosen_rows = visible_rule_rows or row_rules
    else:
        visible_rows = [row for row in rows if not sheet.row_dimensions[row].hidden]
        chosen_rows = visible_rows or rows

    if not chosen_rows:
        return "Confirm the correct source row."

    row_text = ", ".join(str(row) for row in chosen_rows)
    prefix = "visible " if all(not sheet.row_dimensions[row].hidden for row in chosen_rows) else ""
    noun = "row" if len(chosen_rows) == 1 else "rows"
    return f"App will use {prefix}{noun} {row_text}."


def _duplicate_label_finding_text(
    key: str,
    sheet,
    sheet_name: str,
    rows: list[int],
    rule_rows: list[int] | None = None,
) -> str:
    shown_rows = rows[:6]
    hidden_rows = [row for row in shown_rows if sheet.row_dimensions[row].hidden]
    visible_rows = [row for row in shown_rows if not sheet.row_dimensions[row].hidden]
    row_text = ", ".join(str(row) for row in shown_rows)
    if hidden_rows and not visible_rows:
        base = f"'{key}' appears on hidden detail rows {row_text} of '{sheet_name}' with different values."
    elif hidden_rows:
        base = f"'{key}' appears on rows {row_text} of '{sheet_name}' with different values, including hidden detail rows."
    else:
        base = f"'{key}' appears on visible rows {row_text} of '{sheet_name}' with different values."

    if rule_rows:
        visible_rule_rows = [row for row in rule_rows if not sheet.row_dimensions[row].hidden]
        chosen_rows = visible_rule_rows or rule_rows
        chosen_text = ", ".join(str(row) for row in chosen_rows)
        prefix = "visible collapsed " if visible_rule_rows else ""
        noun = "row" if len(chosen_rows) == 1 else "rows"
        return f"{base} The configured source is {prefix}{noun} {chosen_text}."

    return base


def _report_metric_keys(report_sheet, max_rows: int = 400) -> set[str]:
    keys: set[str] = set()
    for row in range(1, min(report_sheet.max_row, max_rows) + 1):
        for col in (1, 2, 3, 4):
            key = label_key(report_sheet.cell(row=row, column=col).value)
            if key:
                keys.add(key)
    return keys


def _duplicate_label_findings(
    company: str,
    sheet,
    sheet_name: str,
    mapped_keys: set[str],
    months: tuple[tuple[int, int], ...],
    row_rules: Mapping[str, tuple[tuple[str, int, float], ...]] | None = None,
) -> list[IntelligenceFinding]:
    findings: list[IntelligenceFinding] = []
    month_cols = find_flexible_month_columns(sheet, months)
    if not month_cols:
        return findings

    generic_keys = {"total", "sub-total", "subtotal", "grand total", "other", "others"}
    label_rows = _label_rows_map(sheet, max_rows=None)
    rule_rows_by_key = _row_rule_source_rows(sheet_name, row_rules)
    for key, rows in label_rows.items():
        if key not in mapped_keys or len(rows) < 2:
            continue
        if key in generic_keys or len(key) < 4:
            continue
        # Only flag when the duplicate rows would produce different numbers.
        differing = False
        for col in month_cols.values():
            values = {
                sheet.cell(row=row, column=col).value
                for row in rows
                if sheet.cell(row=row, column=col).value not in (None, "")
            }
            if len(values) > 1:
                differing = True
                break
        if not differing:
            continue
        dynamic_rule_rows = None
        if company == "Bykea":
            first_month_col = next(iter(month_cols.values()))
            dynamic_row = find_bykea_visible_summary_row(sheet, key, first_month_col)
            if dynamic_row is not None:
                dynamic_rule_rows = [dynamic_row]
        configured_rows = dynamic_rule_rows or rule_rows_by_key.get(key)
        findings.append(
            IntelligenceFinding(
                company=company,
                severity="Review",
                check="Duplicate source label",
                finding=_duplicate_label_finding_text(key, sheet, sheet_name, rows, configured_rows),
                action="Confirm which row is the right source line for this metric.",
                location=f"{sheet_name} rows {', '.join(str(row) for row in rows[:6])}",
                suggested_fix=_source_row_suggestion(sheet, rows, configured_rows),
                confidence=60,
                approval_required=False,
            )
        )
        if len(findings) >= 5:
            break
    return findings


def _magnitude_findings(
    company: str,
    sheet,
    sheet_name: str,
    mapped_keys: set[str],
    months: tuple[tuple[int, int], ...],
) -> list[IntelligenceFinding]:
    """Flag mapped metrics whose new values jump ~100x vs the row's history,
    which usually means a currency or unit change (PKR vs USD, 000s vs units)."""
    findings: list[IntelligenceFinding] = []
    month_cols = find_flexible_month_columns(sheet, months)
    if not month_cols:
        return findings
    selected_cols = set(month_cols.values())

    label_rows = _label_rows_map(sheet)
    for key, rows in label_rows.items():
        if key not in mapped_keys:
            continue
        row = rows[0]
        selected_values = [
            abs(float(value))
            for col in sorted(selected_cols)
            if isinstance((value := sheet.cell(row=row, column=col).value), (int, float))
            and not isinstance(value, bool)
            and value != 0
        ]
        if not selected_values:
            continue
        history_values = [
            abs(float(value))
            for col in range(2, min(sheet.max_column, 100) + 1)
            if col not in selected_cols
            and isinstance((value := sheet.cell(row=row, column=col).value), (int, float))
            and not isinstance(value, bool)
            and value != 0
        ]
        if len(history_values) < 2:
            continue
        history_values.sort()
        baseline = history_values[len(history_values) // 2]
        if baseline <= 0:
            continue
        ratios = [value / baseline for value in selected_values]
        if all(ratio >= 100 for ratio in ratios) and max(selected_values) >= 1000:
            direction = "larger"
        elif all(ratio <= 0.01 for ratio in ratios) and baseline >= 1000:
            direction = "smaller"
        else:
            continue
        findings.append(
            IntelligenceFinding(
                company=company,
                severity="Review",
                check="Possible unit/currency change",
                finding=(
                    f"'{key}' values for the selected months are ~100x {direction} "
                    f"than this row's history in '{sheet_name}'."
                ),
                action="Check whether the company changed currency or units (e.g. PKR vs USD, 000s).",
                location=f"{sheet_name} row {row}",
                suggested_fix="Confirm the intended currency/unit before generating.",
                confidence=65,
                approval_required=False,
            )
        )
        if len(findings) >= 5:
            break
    return findings


_CONSOLIDATED_HEADER_PATTERN = re.compile(r"\bq[1-4]\b|\bfy\s?'?\d{2,4}\b|quarter|annual|consolidat|\bytd\b", re.IGNORECASE)


def _consolidated_finding(
    company: str,
    kpi_workbook,
    matched_sheets: list[str],
    missing_months: list[tuple[int, int]],
) -> IntelligenceFinding | None:
    if not missing_months:
        return None
    evidence: list[str] = []
    for sheet_name in matched_sheets:
        if _CONSOLIDATED_HEADER_PATTERN.search(sheet_name):
            evidence.append(f"tab '{sheet_name}'")
            continue
        sheet = kpi_workbook[sheet_name]
        for row in range(1, min(15, sheet.max_row) + 1):
            for col in range(1, min(sheet.max_column, 100) + 1):
                value = sheet.cell(row=row, column=col).value
                if isinstance(value, str) and _CONSOLIDATED_HEADER_PATTERN.search(value):
                    evidence.append(f"'{sheet_name}' header '{value.strip()[:40]}'")
                    break
            if evidence and evidence[-1].startswith(f"'{sheet_name}'"):
                break
    if not evidence:
        return None
    return IntelligenceFinding(
        company=company,
        severity="Review",
        check="Consolidated file suspected",
        finding=(
            f"Monthly columns for {_format_months(missing_months)} were not found, but the file "
            f"contains quarterly/annual style headers ({'; '.join(evidence[:3])})."
        ),
        action="Confirm whether the company sent a consolidated/quarterly file instead of monthly KPIs.",
        location=", ".join(matched_sheets[:6]),
        suggested_fix="Request the monthly breakdown or approve using the consolidated figures manually.",
        confidence=60,
        approval_required=True,
    )


def _check_abhi(
    template_workbook,
    kpi_workbook,
    profile: CompanyProfile,
    months: tuple[tuple[int, int], ...],
) -> list[IntelligenceFinding]:
    findings: list[IntelligenceFinding] = []
    report_sheet = template_workbook[profile.report_sheet]

    for section in profile.sections:
        actual_source_sheet_name = profile_workbook_sheet_name(kpi_workbook, profile.name, section.source_sheet)
        if actual_source_sheet_name is None:
            findings.append(
                IntelligenceFinding(
                    company=profile.display_name,
                    severity="Warning",
                    check="Missing source sheet",
                    finding=f"ABHI source sheet '{section.source_sheet}' was not found.",
                    action="Ask ABHI/team for the missing sheet or expect this section to remain blank.",
                    location=section.source_sheet,
                    suggested_fix="Confirm whether ABHI renamed or removed this subsidiary sheet.",
                    confidence=85,
                )
            )
            continue

        source_sheet = kpi_workbook[actual_source_sheet_name]
        findings.extend(
            _duplicate_month_findings(profile.display_name, source_sheet, actual_source_sheet_name, months)
        )
        month_cols = find_flexible_month_columns(source_sheet, months)
        missing_months = [month for month in months if month not in month_cols]
        for missing_month in missing_months:
            findings.append(
                IntelligenceFinding(
                    company=profile.display_name,
                    severity="Warning",
                    check="Missing source month",
                    finding=f"'{section.source_sheet}' is missing {_month_label(missing_month)} source data.",
                    action="Confirm whether the source file is incomplete or whether the value should be carried forward manually.",
                    location=f"{section.source_sheet} row 8; {_expected_missing_column(month_cols, missing_month)}",
                    suggested_fix="Ask for the missing month, or leave the section partial.",
                    confidence=80,
                )
            )

        # Explicit-section alignment check: these sections are written by fixed
        # source-row -> report-row pairs. If many source labels no longer sit
        # where the map expects, values can land on the wrong report rows.
        explicit_map = ABHI_EXPLICIT_ROW_MAPS.get(section.source_sheet, {})
        if len(explicit_map) >= 10:
            base_start = ABHI_BASE_SECTION_START_ROWS.get(section.source_sheet)
            actual_start = abhi_section_start_row(report_sheet, section.source_sheet) or section.start_row
            shift = actual_start - base_start if base_start is not None else 0
            section_label_counts: dict[str, int] = {}
            for section_row in range(section.start_row, section.end_row + 1):
                section_label = abhi_report_label(report_sheet, section_row)
                if section_label:
                    section_label_counts[section_label] = section_label_counts.get(section_label, 0) + 1

            compared = 0
            shifted_rows: list[int] = []
            for source_row, base_row in explicit_map.items():
                source_label = abhi_label_key(source_sheet.cell(row=source_row, column=3).value) or abhi_label_key(
                    source_sheet.cell(row=source_row, column=2).value
                )
                target_row = base_row + shift
                target_label = abhi_report_label(report_sheet, target_row)
                if not source_label or not target_label:
                    continue
                compared += 1
                if source_label == target_label:
                    continue
                if section_label_counts.get(source_label, 0) != 1:
                    # Duplicate labels within the section make neighbour
                    # matching unreliable; skip them.
                    continue
                # A mismatch that matches a NEIGHBOURING report row is a shift
                # (inserted/removed line); a plain mismatch is usually just a
                # renamed label and is handled elsewhere.
                for offset in (-2, -1, 1, 2):
                    if abhi_report_label(report_sheet, target_row + offset) == source_label:
                        shifted_rows.append(target_row)
                        break
            # Merged cells make a few labels read one row off even in clean
            # periods (~8 in this workbook family), so only a clearly higher
            # count indicates genuine drift.
            if compared >= 10 and len(shifted_rows) >= 10:
                findings.append(
                    IntelligenceFinding(
                        company=profile.display_name,
                        severity="Review",
                        check="Section rows may be misaligned",
                        finding=(
                            f"In '{section.source_sheet}', {len(shifted_rows)} mapped rows look shifted by a row "
                            f"or two (e.g. report rows {', '.join(str(r) for r in shifted_rows[:4])})."
                        ),
                        action="Spot-check this section in the generated file; lines may have been inserted or removed.",
                        location=f"{section.source_sheet} vs report rows {section.start_row}:{section.end_row}",
                        suggested_fix="Compare the section against the KPI file and correct any misplaced rows manually.",
                        confidence=70,
                        approval_required=True,
                    )
                )

        report_labels = _abhi_report_labels(report_sheet, section.start_row, section.end_row)
        source_labels = _abhi_source_labels(source_sheet)
        if report_labels:
            missing_labels = sorted(report_labels - source_labels)
            missing_ratio = len(missing_labels) / len(report_labels)
            if missing_labels and missing_ratio >= 0.25:
                findings.append(
                    IntelligenceFinding(
                        company=profile.display_name,
                        severity="Review",
                        check="Possible format change",
                        finding=(
                            f"{len(missing_labels)} expected labels in '{section.source_sheet}' were not recognized. "
                            f"Examples: {', '.join(missing_labels[:5])}."
                        ),
                        action="Review whether ABHI renamed rows or changed the section format.",
                        location=f"{section.source_sheet} vs report rows {section.start_row}:{section.end_row}",
                        suggested_fix="Compare the renamed rows against the report section and approve mappings if valid.",
                        confidence=70,
                        approval_required=True,
                    )
                )

    return findings


def _check_label_copy_profile(
    template_workbook,
    kpi_workbook,
    profile_name: str,
    profile: LabelCopyProfile,
    months: tuple[tuple[int, int], ...],
) -> list[IntelligenceFinding]:
    findings: list[IntelligenceFinding] = []
    candidates = _profile_sheet_candidates(profile_name, profile)
    matched_sheets = [
        actual_sheet_name
        for sheet_name in candidates
        if (actual_sheet_name := infer_profile_source_sheet_name(kpi_workbook, profile, sheet_name, months)) is not None
    ]

    if profile_name == "roomy":
        for sheet_name in kpi_workbook.sheetnames:
            if any(_sheet_name_mentions_month(sheet_name, month_key) for month_key in months):
                if sheet_name not in matched_sheets:
                    matched_sheets.append(sheet_name)

    # Several expected candidate names can resolve to the same actual sheet;
    # keep each actual sheet once so checks do not repeat.
    matched_sheets = list(dict.fromkeys(matched_sheets))

    if not matched_sheets:
        findings.append(
            IntelligenceFinding(
                company=profile.display_name,
                severity="Blocked",
                check="Missing source sheet",
                finding="No expected KPI sheet was found in the uploaded workbook.",
                action="Upload the right KPI file or check whether the company changed workbook tabs.",
                location=", ".join(kpi_workbook.sheetnames[:8]),
                suggested_fix="Approve a renamed workbook tab in the normalization review below, if one is suggested.",
                confidence=90,
            )
        )
        return findings

    found_months: set[tuple[int, int]] = set()
    for sheet_name in matched_sheets:
        sheet = kpi_workbook[sheet_name]
        found_months.update(find_flexible_month_columns(sheet, months).keys())
        for month_key in months:
            if _sheet_name_mentions_month(sheet_name, month_key):
                found_months.add(month_key)

    missing_months = [month for month in months if month not in found_months]
    if missing_months:
        severity = "Blocked" if profile_name == "roomy" else "Warning"
        findings.append(
            IntelligenceFinding(
                company=profile.display_name,
                severity=severity,
                check="Missing source month",
                finding=f"Could not confirm source data for {_format_months(missing_months)}.",
                action="Upload the missing monthly file(s) or confirm that this company should be left partial.",
                location=", ".join(matched_sheets[:6]),
                suggested_fix="Upload the missing month file(s), or generate a partial report on purpose.",
                confidence=80,
            )
        )
        consolidated = _consolidated_finding(profile.display_name, kpi_workbook, matched_sheets, missing_months)
        if consolidated is not None:
            findings.append(consolidated)

    report_sheet = template_workbook[profile.report_sheet]
    mapped_keys = _report_metric_keys(report_sheet) | _profile_expected_label_keys(profile)
    for sheet_name in matched_sheets:
        sheet = kpi_workbook[sheet_name]
        findings.extend(_duplicate_month_findings(profile.display_name, sheet, sheet_name, months))
        findings.extend(
            _duplicate_label_findings(
                profile.display_name,
                sheet,
                sheet_name,
                mapped_keys,
                months,
                profile.row_rules,
            )
        )
        findings.extend(_magnitude_findings(profile.display_name, sheet, sheet_name, mapped_keys, months))

    expected_labels = _profile_expected_label_keys(profile)
    if expected_labels:
        available_labels = _available_label_keys(kpi_workbook, matched_sheets)
        missing_labels = sorted(expected_labels - available_labels)
        if missing_labels:
            matched_count = len(expected_labels) - len(missing_labels)
            coverage = matched_count / len(expected_labels)
            if coverage < 0.6:
                findings.append(
                    IntelligenceFinding(
                        company=profile.display_name,
                        severity="Warning",
                        check="Possible format change",
                        finding=(
                            f"{len(missing_labels)} expected KPI labels were not recognized. "
                            f"Examples: {', '.join(missing_labels[:5])}."
                        ),
                        action="Review whether the company renamed rows, added lines, or sent a different format.",
                        location=", ".join(matched_sheets[:6]),
                        suggested_fix="Approve renamed-label suggestions in the normalization review below.",
                        confidence=70,
                        approval_required=True,
                    )
                )

    return findings


def run_intelligence_checks(
    template_path: Path,
    kpi_files: Mapping[str, Path],
    months: tuple[tuple[int, int], ...],
) -> dict[str, object]:
    findings: list[IntelligenceFinding] = []

    try:
        template_workbook = load_workbook(template_path, data_only=True, read_only=False)
    except Exception as exc:
        return {
            "findings": [
                IntelligenceFinding(
                    company="Monthly Reporting Template",
                    severity="Blocked",
                    check="Template open check",
                    finding=f"Could not open template: {exc}",
                    action="Upload a valid Monthly Reporting Template workbook.",
                    confidence=100,
                ).row()
            ],
            "status_counts": {"Blocked": 1, "Warning": 0, "Review": 0},
            "blocked": True,
        }

    try:
        for profile_name, kpi_path in kpi_files.items():
            profile = PROFILES.get(profile_name)
            if profile is None:
                continue

            if profile.report_sheet not in template_workbook.sheetnames:
                findings.append(
                    IntelligenceFinding(
                        company=profile.display_name,
                        severity="Blocked",
                        check="Missing report tab",
                        finding=f"Template is missing the '{profile.report_sheet}' tab.",
                        action="Use the correct Monthly Reporting Template.",
                        location=profile.report_sheet,
                        confidence=100,
                    )
                )
                continue

            report_sheet = template_workbook[profile.report_sheet]
            report_month_cols = find_flexible_month_columns(report_sheet, months)
            missing_report_months = [month for month in months if month not in report_month_cols]
            if missing_report_months:
                findings.append(
                    IntelligenceFinding(
                        company=profile.display_name,
                        severity="Blocked",
                        check="Missing report month columns",
                        finding=f"Template is missing {_format_months(missing_report_months)} columns for this company.",
                        action="Use a template with the selected period columns already added.",
                        location=profile.report_sheet,
                        confidence=95,
                    )
                )
                continue

            try:
                kpi_workbook = load_workbook(kpi_path, data_only=True, read_only=False)
            except Exception as exc:
                findings.append(
                    IntelligenceFinding(
                        company=profile.display_name,
                        severity="Blocked",
                        check="KPI open check",
                        finding=f"Could not open KPI workbook: {exc}",
                        action="Upload a valid Excel KPI file.",
                        location=str(kpi_path),
                        confidence=100,
                    )
                )
                continue

            try:
                if isinstance(profile, CompanyProfile):
                    findings.extend(_check_abhi(template_workbook, kpi_workbook, profile, months))
                else:
                    findings.extend(
                        _check_label_copy_profile(template_workbook, kpi_workbook, profile_name, profile, months)
                    )
            finally:
                kpi_workbook.close()
    finally:
        template_workbook.close()

    findings = list(dict.fromkeys(findings))
    status_counts = {
        "Blocked": sum(finding.severity == "Blocked" for finding in findings),
        "Warning": sum(finding.severity == "Warning" for finding in findings),
        "Review": sum(finding.severity == "Review" for finding in findings),
    }
    return {
        "findings": [finding.row() for finding in findings],
        "status_counts": status_counts,
        "blocked": status_counts["Blocked"] > 0,
    }
