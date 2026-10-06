from __future__ import annotations

import io
import json
import os
import re
import tempfile
import ast
import operator
from calendar import month_abbr, month_name
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
import streamlit as st

from monthly_reporting_automation import PROFILES, find_month_columns, workbook_sheet_name
from monthly_reporting_page import QOQ_METHOD_MANUAL, _qoq_rows_for_workbook


PROJECT_ROOT = Path(__file__).resolve().parent
QUARTERLY_OUTPUTS_DIR = PROJECT_ROOT / "Outputs" / "Quarterly Reports"
DEFAULT_OPENAI_MODEL = "gpt-5-mini"
COMMENTARY_SHEET_NAME = "Quarterly Commentary"
SECTION_HIGHLIGHTS = "Highlights"
SECTION_PORTFOLIO_COMPANY = "Portfolio Company"
SECTION_OPERATOR_ANGEL = "Operator Angel"
FINANCIAL_EXHIBIT_SECTIONS = ("Balance Sheet", "Income Statement", "Financial Summary")

OPERATOR_ANGEL_COMPANIES = (
    "Savvy Technologies",
    "Bolandi Technologies",
    "Delsys Technologies",
    "Startup Early",
    "Scholar Den",
    "House Call",
    "Orko",
    "L.L.M.Bots",
    "Aabshar",
    "Raptr Games",
)

HIGHLIGHT_ROWS = (
    ("Investment Activity", "Fund-level investment activity bullets for the first highlights page."),
    ("Portfolio Highlights", "Portfolio company bullets for the first highlights page."),
    ("Other Firm Matters", "Firm-level bullets for the first highlights page."),
)

MAX_BULLETS = 4
MAX_COMMENTARY_BULLETS = 12
MAX_BULLET_CHARS = 190
MAX_DESCRIPTION_CHARS = 520
MAX_UPDATES_CHARS = 820

EXCLUDED_LABEL_TOKENS = (
    "source:",
    "note:",
    "check",
    "commentary",
    "what worked",
    "what didn't work",
    "what didnt work",
    "targets for next month",
    "key updates",
    "company description",
    "key financial",
    "standard reporting template",
    "forecast",
)

METRIC_PRIORITIES = (
    ("total revenue", 120),
    ("net revenue", 115),
    ("revenue", 110),
    ("gmv", 105),
    ("gtv", 105),
    ("ntv", 105),
    ("gross profit", 100),
    ("gross margin", 96),
    ("total expenses", 92),
    ("net burn", 90),
    ("ebitda", 88),
    ("net income", 86),
    ("operating profit", 84),
    ("monthly active users", 80),
    ("mau", 80),
    ("bookings", 78),
    ("transactions", 76),
    ("active corporates", 74),
    ("cash", 72),
    ("runway", 70),
    ("subscriptions", 68),
    ("total students", 66),
    ("teacher", 64),
)

MAX_FIGMA_METRIC_ROWS = 18

COMPANY_METRIC_MAP: dict[str, tuple[tuple[str, str], ...]] = {
    "ABHI": (
        ("NTV", "NTV"),
        ("Net Revenue", "Net Revenue"),
        ("COS", "COS"),
        ("Gross Profit", "Gross Profit"),
        ("Other Income", "Other Income"),
        ("Salaries", "Salaries"),
        ("Tech Development", "Tech Development"),
        ("Marketing", "Marketing"),
        ("SG&A", "SG&A"),
        ("Other Expenses (inc.one time)", "Other Expenses"),
        ("Total Expenses", "Total Expenses"),
        ("Capex", "Capex"),
        ("Net Burn", "Net Burn"),
        ("Transactions (EWA + Payroll)", "Transactions (EWA + Payroll)"),
    ),
    "SimPaisa": (
        ("Total GTV", "GTV"),
        ("Net Revenue", "Net Revenue"),
        ("Salaries", "Salaries"),
        ("SG&A", "SG&A"),
        ("Business Development", "Business Development"),
        ("Professional Fees", "Professional Fees"),
        ("Other expenses", "Other expenses"),
        ("Total Expenses", "Total Expenses"),
        ("Net Income (loss)", "Net Burn / Income"),
        ("Wallet Transactions", "Wallet Transactions"),
        ("Telco Transactions", "Telco Transactions"),
        ("Card Transactions", "Card Transactions"),
        ("Disbursement Transactions", "Disbursement Transactions"),
        ("Total No. of Transactions", "Total No. of Transactions"),
    ),
    "Bykea": (
        ("GTV", "GTV"),
        ("Net Revenue", "Net Revenue"),
        ("Driver Incentives", "Driver Incentives"),
        ("Marketing", "Marketing"),
        ("Tech", "Tech"),
        ("Overheads", "Overheads"),
        ("Net Burn", "Net Burn"),
        ("Bookings / day", "Bookings / day"),
        ("Net Tranactions/day", "Net Transactions / day"),
        ("Fulfilment", "Fulfilment"),
        ("MAU (Monthly Active Users) Net", "MAU"),
        ("MAD (Monthly Active Drivers)", "MAD"),
    ),
    "Dot & Line": (
        ("GMV", "GMV"),
        ("Net revenue", "Net Revenue"),
        ("Other revenue", "Other Revenue"),
        ("Salaries", "Salaries"),
        ("SG&A", "SG&A"),
        ("Marketing", "Marketing"),
        ("Other expenses", "Other Expenses"),
        ("Total Expenses", "Total Expenses"),
        ("Net Income (loss)", "Net Burn"),
        ("Total Students", "Total Students"),
        ("Total Teacher Partners", "Total Teacher Partners"),
    ),
    "Procheck": (
        ("Revenue (New Lines/Codes)", "Revenue (New lines/codes)"),
        ("Revenue (Recurring Lines/Codes)", "Revenue (Recurring lines/codes)"),
        ("Total Revenue", "Total Revenue"),
        ("Cost of Sales", "Deployment cost"),
        ("Gross Margin", "Gross Profit"),
        ("Salaries", "Salaries"),
        ("SG&A", "SG&A"),
        ("Software Development Cost", "Software Development Cost"),
        ("Other Expenses", "Other Expenses"),
        ("Total Expenses", "Total Expenses"),
        ("Total Profit (Loss) / Burn", "Net Burn"),
        ("Codes Generated (TnT)", "Codes Generated (TnT)"),
        ("No. of Lines (OEE)", "No. of Lines (OEE)"),
    ),
    "Roomy": (
        ("Revenue - Total", "Total Revenue"),
        ("Gross Profit", "Gross Profit"),
        ("Salaries & Wages", "Salaries"),
        ("Sales & Marketing", "Sales & Marketing"),
        ("Rent & CAM / Renovation Costs", "Rent & CAM / Renovation Costs"),
        ("Admin & General", "Admin & General"),
        ("Total Operating Expenses", "Total Operating Expenses"),
        ("Net Burn", "Net Burn"),
        ("Operational Rooms", "Operational Rooms"),
        ("Occupancy %", "Occupancy Rate (%)"),
        ("ADR", "Average Daily Rent ($)"),
    ),
    "Oladoc": (
        ("GMV (USD)", "GMV (USD)"),
        ("NMV (USD)", "NMV (USD)"),
        ("Total Revenue", "Total Revenue"),
        ("Gross Profit", "Gross Profit"),
        ("Salaries", "Salaries"),
        ("Marketing", "Marketing"),
        ("Tech - Subscriptions", "Tech - Subscriptions"),
        ("G&A", "G&A"),
        ("Others", "Others"),
        ("Total Expenses", "Total Expenses"),
        ("Net Burn", "Net Burn"),
        ("Total Bookings", "Total Bookings"),
        ("Billed Bookings (0nline)", "Billed Bookings (Online)"),
        ("Total Active Doctors", "Total Active Doctors"),
    ),
    "Revolving Games": (
        ("Revenue", "Revenue"),
        ("Operations and BD", "Operations"),
        ("SG&A", "SG&A"),
        ("Marketing", "Marketing"),
        ("Software and Servers", "Software and Servers"),
        ("Other expenses", "Other expenses"),
        ("Total Expenses", "Total Expenses"),
        ("Net Income (loss)", "Net Burn"),
    ),
    "Tapmad": (
        ("Gross Revenue", "Gross Revenue"),
        ("Net Revenue", "Net Revenue"),
        ("Cost of Sales & Services", "COS"),
        ("Gross Profit", "Gross Profit"),
        ("Salaries & Wages (incl. bonuses/perks)", "Salaries"),
        ("SG&A (Rent, utilities, office expenses)", "SG&A"),
        ("Marketing", "Marketing"),
        ("Professional & consultancy fee", "Professional / Consultancy Fee"),
        ("Other expenses", "Other"),
        ("OPEX", "Total Operating Expenses"),
        ("CAPEX", "CAPEX"),
        ("Net Profit / (Loss)", "Net Profit / (Loss)"),
        ("Subscriptions", "Subscriptions"),
        ("Monthly Active Users (MAUs)", "Monthly Active Users (MAUs)"),
    ),
    "Jiye Technologies": (
        ("NMV (USD)", "NMV (USD)"),
        ("Net Revenue", "Net Revenue"),
        ("Gross Profit", "Gross Profit"),
        ("Gross Margin", "Gross Margin"),
        ("Salaries", "Salaries"),
        ("Tech", "Tech"),
        ("Marketing", "Marketing"),
        ("Others", "Other Expenses"),
        ("Total Expenses", "Total Expenses"),
        ("Net Burn", "Net Burn"),
        ("Net Orders Delivered", "Net Orders Delivered"),
        ("AOV", "AOV"),
        ("Total Customers", "Total Customers"),
        ("New Customers", "New Customers"),
    ),
    "OneLoad": (
        ("GMV - Total", "GMV"),
        ("Net Revenue - Total", "Net Revenue"),
        ("Variable Costs", "Variable costs"),
        ("Gross Profit", "Gross Profit"),
        ("Salaries", "Salaries"),
        ("SG&A", "SG&A"),
        ("Other Expenses", "Other Expenses"),
        ("Total Operating Expenses", "Total Operating Expenses"),
        ("Net Burn", "Net Burn"),
        ("Active Retailers", "Active Retailers"),
        ("Total No. of Transactions", "No. of Transactions"),
    ),
}


@dataclass(frozen=True)
class MetricRow:
    label: str
    values: dict[tuple[int, int], object]


@dataclass(frozen=True)
class CompanyReportData:
    company: str
    profile_name: str
    sheet_name: str | None
    description: str
    key_updates: tuple[str, ...]
    metrics: tuple[MetricRow, ...]
    qoq_changes: tuple[dict[str, object], ...]
    warnings: tuple[str, ...]


@dataclass(frozen=True)
class FinancialExhibitRow:
    section: str
    label: str
    values: dict[str, object]


def quarter_months(year: int, quarter_end_month: int) -> tuple[tuple[int, int], ...]:
    quarter_start_month = ((quarter_end_month - 1) // 3) * 3 + 1
    return tuple((year, quarter_start_month + offset) for offset in range(3))


def quarter_label(months: Iterable[tuple[int, int]]) -> str:
    months = tuple(months)
    if not months:
        return "Quarterly Report"
    year, ending_month = months[-1]
    quarter = ((ending_month - 1) // 3) + 1
    return f"Q{quarter} {year}"


def company_page_months(months: Iterable[tuple[int, int]]) -> tuple[tuple[int, int], ...]:
    months = tuple(months)
    if not months:
        return ()
    start_year, start_month = months[0]
    previous: list[tuple[int, int]] = []
    year = start_year
    month = start_month
    for _ in range(3):
        month -= 1
        if month == 0:
            month = 12
            year -= 1
        previous.append((year, month))
    return tuple(reversed(previous)) + months


def _cell_text(value: object) -> str:
    if not isinstance(value, str):
        return ""
    text = str(value).replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _normalise(value: str) -> str:
    return " ".join(value.lower().replace("&", "and").split())


def _strip_bullet(value: str) -> str:
    text = value.strip()
    text = re.sub(r"^[\u2022*\-]+\s*", "", text)
    return text.strip()


def _first_nonempty_text_in_row(sheet: Worksheet, row: int, max_col: int = 10) -> str:
    for col in range(1, min(sheet.max_column, max_col) + 1):
        text = _cell_text(sheet.cell(row=row, column=col).value)
        if text:
            return text
    return ""


def _find_heading_row(sheet: Worksheet, heading: str, max_row: int = 80) -> tuple[int, int] | None:
    wanted = _normalise(heading)
    for row in range(1, min(sheet.max_row, max_row) + 1):
        for col in range(1, min(sheet.max_column, 10) + 1):
            text = _cell_text(sheet.cell(row=row, column=col).value)
            if wanted in _normalise(text):
                return row, col
    return None


def _extract_description(sheet: Worksheet, company: str) -> str:
    heading = _find_heading_row(sheet, "Company Description", max_row=25)
    if heading is not None:
        row, _col = heading
        candidates = [
            _first_nonempty_text_in_row(sheet, next_row, max_col=8)
            for next_row in range(row + 1, min(row + 6, sheet.max_row) + 1)
        ]
        return max(candidates, key=len, default="").strip()

    company_key = _normalise(company)
    candidates: list[str] = []
    for row in range(1, min(sheet.max_row, 10) + 1):
        text = _first_nonempty_text_in_row(sheet, row, max_col=8)
        if not text:
            continue
        text_key = _normalise(text)
        if text_key == company_key or "key financial" in text_key or "key updates" in text_key:
            continue
        if len(text) >= 24:
            candidates.append(text)
    return max(candidates, key=len, default="").strip()


def _extract_key_updates(sheet: Worksheet) -> tuple[str, ...]:
    heading = _find_heading_row(sheet, "Key Updates", max_row=35)
    if heading is None:
        return ()

    start_row, _col = heading
    updates: list[str] = []
    blank_streak = 0
    for row in range(start_row + 1, min(start_row + 12, sheet.max_row) + 1):
        text = _first_nonempty_text_in_row(sheet, row, max_col=8)
        key = _normalise(text)
        if not text:
            blank_streak += 1
            if blank_streak >= 3 and updates:
                break
            continue
        blank_streak = 0
        if "key financial" in key or "reporting link" in key or "source:" in key:
            break
        updates.append(_strip_bullet(text))
    return tuple(update for update in updates if update)


def _label_for_row(sheet: Worksheet, row: int, first_month_col: int) -> str:
    for col in range(max(1, first_month_col - 6), 0, -1):
        text = _cell_text(sheet.cell(row=row, column=col).value)
        if text and not text.startswith("#"):
            return text
    return ""


def _metric_label_rows(sheet: Worksheet, first_month_col: int, max_row: int = 160) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row in range(1, min(sheet.max_row, max_row) + 1):
        label = _label_for_row(sheet, row, first_month_col)
        key = _normalise(label)
        if label and key not in rows:
            rows[key] = row
    return rows


def _is_metric_value(value: object) -> bool:
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    return False


def _metric_score(label: str) -> int:
    key = _normalise(label)
    if not key or any(token in key for token in EXCLUDED_LABEL_TOKENS):
        return 0
    for token, score in METRIC_PRIORITIES:
        if token in key:
            return score
    return 25 if len(label) <= 60 else 0


_CELL_REF_RE = re.compile(r"(?<![A-Za-z0-9_])\$?([A-Z]{1,3})\$?([0-9]{1,7})(?![A-Za-z0-9_])")
_CELL_REF_ONLY_RE = re.compile(r"^\$?([A-Z]{1,3})\$?([0-9]{1,7})$")
_RANGE_REF_ONLY_RE = re.compile(
    r"^\$?([A-Z]{1,3})\$?([0-9]{1,7}):\$?([A-Z]{1,3})\$?([0-9]{1,7})$"
)


def _safe_number_eval(expression: str) -> float | int | None:
    operators = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.USub: operator.neg,
        ast.UAdd: operator.pos,
    }

    def evaluate(node):
        if isinstance(node, ast.Expression):
            return evaluate(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.BinOp) and type(node.op) in operators:
            return operators[type(node.op)](evaluate(node.left), evaluate(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in operators:
            return operators[type(node.op)](evaluate(node.operand))
        raise ValueError("Unsupported formula expression")

    if not re.fullmatch(r"[0-9eE.+\-*/() ]+", expression):
        return None
    try:
        result = evaluate(ast.parse(expression, mode="eval"))
    except (SyntaxError, ValueError, ZeroDivisionError, TypeError):
        return None
    if isinstance(result, (int, float)) and not isinstance(result, bool):
        return result
    return None


def _resolved_formula_cell_value(
    data_sheet: Worksheet,
    formula_sheet: Worksheet,
    row: int,
    col: int,
    cache: dict[tuple[int, int], object],
    stack: set[tuple[int, int]],
) -> object:
    key = (row, col)
    if key in cache:
        return cache[key]

    data_value = data_sheet.cell(row=row, column=col).value
    if _is_metric_value(data_value):
        cache[key] = data_value
        return data_value

    formula_value = formula_sheet.cell(row=row, column=col).value
    if _is_metric_value(formula_value):
        cache[key] = formula_value
        return formula_value
    if not (isinstance(formula_value, str) and formula_value.startswith("=")):
        cache[key] = data_value
        return data_value
    if key in stack:
        return None

    stack.add(key)
    try:
        resolved = _evaluate_simple_formula(data_sheet, formula_sheet, formula_value, cache, stack)
    finally:
        stack.remove(key)
    cache[key] = resolved
    return resolved


def _sum_formula_argument(
    data_sheet: Worksheet,
    formula_sheet: Worksheet,
    argument: str,
    cache: dict[tuple[int, int], object],
    stack: set[tuple[int, int]],
) -> float | int | None:
    argument = argument.strip()
    range_match = _RANGE_REF_ONLY_RE.fullmatch(argument.replace("$", ""))
    if range_match:
        min_col, min_row, max_col, max_row = range_boundaries(argument.replace("$", ""))
        total = 0
        found = False
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                value = _resolved_formula_cell_value(data_sheet, formula_sheet, row, col, cache, stack)
                if _is_metric_value(value):
                    total += value
                    found = True
        return total if found else None

    cell_match = _CELL_REF_ONLY_RE.fullmatch(argument.replace("$", ""))
    if cell_match:
        col = column_index_from_string(cell_match.group(1))
        row = int(cell_match.group(2))
        value = _resolved_formula_cell_value(data_sheet, formula_sheet, row, col, cache, stack)
        return value if _is_metric_value(value) else 0

    return _evaluate_simple_formula(data_sheet, formula_sheet, f"={argument}", cache, stack)


def _evaluate_simple_formula(
    data_sheet: Worksheet,
    formula_sheet: Worksheet,
    formula: str,
    cache: dict[tuple[int, int], object],
    stack: set[tuple[int, int]],
) -> float | int | None:
    expression = formula.strip()[1:].replace("$", "")
    if not expression or "[" in expression or "!" in expression:
        return None
    if "=" in expression:
        return None

    direct_cell = _CELL_REF_ONLY_RE.fullmatch(expression)
    if direct_cell:
        col = column_index_from_string(direct_cell.group(1))
        row = int(direct_cell.group(2))
        value = _resolved_formula_cell_value(data_sheet, formula_sheet, row, col, cache, stack)
        return value if _is_metric_value(value) else None

    def replace_sum(match: re.Match) -> str:
        pieces = [piece.strip() for piece in match.group(1).split(",")]
        total = 0
        found = False
        for piece in pieces:
            value = _sum_formula_argument(data_sheet, formula_sheet, piece, cache, stack)
            if _is_metric_value(value):
                total += value
                found = True
        return str(total if found else 0)

    previous = None
    while previous != expression:
        previous = expression
        expression = re.sub(r"SUM\(([^()]+)\)", replace_sum, expression, flags=re.IGNORECASE)

    expression = re.sub(r"(\d+(?:\.\d+)?)%", r"(\1/100)", expression)

    def replace_cell(match: re.Match) -> str:
        col = column_index_from_string(match.group(1))
        row = int(match.group(2))
        value = _resolved_formula_cell_value(data_sheet, formula_sheet, row, col, cache, stack)
        return str(value if _is_metric_value(value) else 0)

    expression = _CELL_REF_RE.sub(replace_cell, expression)
    return _safe_number_eval(expression)


def _extract_metrics(
    sheet: Worksheet,
    formula_sheet: Worksheet,
    months: tuple[tuple[int, int], ...],
    company: str | None = None,
) -> tuple[tuple[MetricRow, ...], tuple[str, ...]]:
    warnings: list[str] = []
    month_cols = find_month_columns(sheet, months)
    missing_months = [month for month in months if month not in month_cols]
    if missing_months:
        missing = ", ".join(f"{month_abbr[month]} {str(year)[-2:]}" for year, month in missing_months)
        warnings.append(f"Missing month columns: {missing}.")
    if not month_cols:
        return (), tuple(warnings)

    first_month_col = min(month_cols.values())
    formula_cache: dict[tuple[int, int], object] = {}
    metric_map = COMPANY_METRIC_MAP.get(company or "")
    if metric_map:
        label_rows = _metric_label_rows(sheet, first_month_col)
        selected: list[MetricRow] = []
        missing_labels: list[str] = []
        empty_value_labels: list[str] = []
        for source_label, display_label in metric_map:
            row = label_rows.get(_normalise(source_label))
            if row is None:
                missing_labels.append(source_label)
                continue
            values = {
                month: _resolved_formula_cell_value(sheet, formula_sheet, row, col, formula_cache, set())
                for month, col in month_cols.items()
            }
            if not any(_is_metric_value(value) for value in values.values()):
                empty_value_labels.append(display_label)
            selected.append(MetricRow(label=display_label, values=values))
        if missing_labels:
            missing = ", ".join(missing_labels[:4])
            if len(missing_labels) > 4:
                missing += f", +{len(missing_labels) - 4} more"
            warnings.append(f"Configured metric labels missing from workbook: {missing}.")
        if empty_value_labels:
            empty = ", ".join(empty_value_labels[:4])
            if len(empty_value_labels) > 4:
                empty += f", +{len(empty_value_labels) - 4} more"
            warnings.append(f"Configured metric rows have no usable values for this period: {empty}.")
        if selected:
            for metric in selected:
                missing_values = [month for month in months if not _is_metric_value(metric.values.get(month))]
                if missing_values:
                    warnings.append(f"Metric '{metric.label}' has missing values for part of the quarter.")
                    break
            return tuple(selected), tuple(warnings)

    candidates: list[tuple[int, int, MetricRow]] = []
    seen_labels: set[str] = set()
    for row in range(1, min(sheet.max_row, 130) + 1):
        label = _label_for_row(sheet, row, first_month_col)
        score = _metric_score(label)
        if not score:
            continue
        label_key = _normalise(label)
        if label_key in seen_labels:
            continue
        values = {
            month: _resolved_formula_cell_value(sheet, formula_sheet, row, col, formula_cache, set())
            for month, col in month_cols.items()
        }
        if not any(_is_metric_value(value) for value in values.values()):
            continue
        seen_labels.add(label_key)
        candidates.append((-score, row, MetricRow(label=label, values=values)))

    candidates.sort()
    selected = tuple(metric for _score, _row, metric in candidates[:7])
    if not selected:
        period = ", ".join(f"{month_abbr[month]} {str(year)[-2:]}" for year, month in months)
        warnings.append(f"No usable actual metric values found for {period}. Check that the Monthly Reporting workbook has populated Actual columns for this quarter.")
    else:
        for metric in selected:
            missing_values = [month for month in months if not _is_metric_value(metric.values.get(month))]
            if missing_values:
                warnings.append(f"Metric '{metric.label}' has missing values for part of the quarter.")
                break
    return selected, tuple(warnings)


def _top_qoq_changes(
    qoq_rows: list[dict[str, object]],
    company: str,
    limit: int = 6,
) -> tuple[dict[str, object], ...]:
    company_rows = [row for row in qoq_rows if row.get("Company") == company]

    def score(row: dict[str, object]) -> float:
        qoq_pct = row.get("QoQ %")
        change = row.get("Change")
        if isinstance(qoq_pct, (int, float)):
            return abs(float(qoq_pct))
        if isinstance(change, (int, float)):
            return abs(float(change))
        return 0.0

    return tuple(sorted(company_rows, key=score, reverse=True)[:limit])


def _text_warnings(description: str, updates: tuple[str, ...]) -> list[str]:
    warnings: list[str] = []
    if not description:
        warnings.append("Missing company description.")
    elif len(description) > MAX_DESCRIPTION_CHARS:
        warnings.append(f"Description may overflow fixed text box ({len(description)} chars).")

    if not updates:
        warnings.append("Missing key updates.")
    if len(updates) > MAX_BULLETS:
        warnings.append(f"Too many key update bullets ({len(updates)}; target {MAX_BULLETS}).")
    if sum(len(update) for update in updates) > MAX_UPDATES_CHARS:
        warnings.append("Key updates may overflow fixed text box.")
    if any(len(update) > MAX_BULLET_CHARS for update in updates):
        warnings.append(f"One or more bullets exceed {MAX_BULLET_CHARS} characters.")
    return warnings


def parse_monthly_reporting_workbook(
    workbook_path: Path,
    months: tuple[tuple[int, int], ...],
    metric_months: tuple[tuple[int, int], ...] | None = None,
) -> tuple[CompanyReportData, ...]:
    metric_months = metric_months or months
    workbook = load_workbook(workbook_path, data_only=True)
    formula_workbook = load_workbook(workbook_path, data_only=False)
    qoq_rows, qoq_warnings = _qoq_rows_for_workbook(
        workbook_path=workbook_path,
        months=months,
        selected_profile_names=list(PROFILES.keys()),
        calculation_method=QOQ_METHOD_MANUAL,
    )
    companies: list[CompanyReportData] = []
    try:
        for profile_name, profile in PROFILES.items():
            company_warnings: list[str] = [
                warning
                for warning in qoq_warnings
                if warning.startswith(f"{profile.display_name}:")
            ]
            sheet_name = workbook_sheet_name(workbook, profile.report_sheet)
            if sheet_name is None:
                companies.append(
                    CompanyReportData(
                        company=profile.display_name,
                        profile_name=profile_name,
                        sheet_name=None,
                        description="",
                        key_updates=(),
                        metrics=(),
                        qoq_changes=(),
                        warnings=(f"Missing workbook sheet '{profile.report_sheet}'.",),
                    )
                )
                continue

            sheet = workbook[sheet_name]
            formula_sheet = formula_workbook[sheet_name]
            description = _extract_description(sheet, profile.display_name)
            key_updates = _extract_key_updates(sheet)
            metrics, metric_warnings = _extract_metrics(sheet, formula_sheet, metric_months, profile.display_name)
            company_warnings.extend(_text_warnings(description, key_updates))
            company_warnings.extend(metric_warnings)
            companies.append(
                CompanyReportData(
                    company=profile.display_name,
                    profile_name=profile_name,
                    sheet_name=sheet_name,
                    description=description,
                    key_updates=key_updates,
                    metrics=metrics,
                    qoq_changes=_top_qoq_changes(qoq_rows, profile.display_name),
                    warnings=tuple(company_warnings),
                )
            )
    finally:
        workbook.close()
        formula_workbook.close()
    return tuple(companies)


def _format_metric_value(label: str, value: object) -> str:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return "-"

    label_key = _normalise(label)
    if any(token in label_key for token in ("%", "margin", "rate", "yield")) and abs(value) <= 5:
        return f"{value * 100:.1f}%"
    if abs(value) >= 1_000_000:
        return f"{value / 1_000_000:.1f}M"
    if abs(value) >= 1_000:
        return f"{value / 1_000:.1f}K"
    if float(value).is_integer():
        return f"{int(value):,}"
    return f"{value:.1f}"


def _openai_setting(name: str, default: str | None = None) -> str | None:
    try:
        value = st.secrets.get(name)  # type: ignore[attr-defined]
    except Exception:
        value = None
    if value:
        return str(value)
    return os.environ.get(name, default)


def _openai_api_key_configured() -> bool:
    return bool(_openai_setting("OPENAI_API_KEY"))


def _metric_context(company: CompanyReportData, months: tuple[tuple[int, int], ...]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for metric in company.metrics[:8]:
        values = {
            f"{month_abbr[month]} {str(year)[-2:]}": _format_metric_value(metric.label, metric.values.get((year, month)))
            for year, month in months
        }
        numeric_values = [
            metric.values.get(month)
            for month in months
            if _is_metric_value(metric.values.get(month))
        ]
        change = None
        if len(numeric_values) >= 2 and numeric_values[0] not in (0, None):
            try:
                change = (numeric_values[-1] / numeric_values[0]) - 1
            except ZeroDivisionError:
                change = None
        rows.append(
            {
                "metric": metric.label,
                "values": values,
                "quarter_change": f"{change * 100:.1f}%" if isinstance(change, (int, float)) else None,
            }
        )
    return rows


def _format_qoq_value(value: object, percent: bool = False) -> str | None:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    if percent:
        return f"{value:.1f}%"
    return _format_metric_value("", value)


def _qoq_context(company: CompanyReportData) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for row in company.qoq_changes:
        rows.append(
            {
                "metric": row.get("Metric"),
                "method": row.get("Method"),
                "current_quarter": _format_qoq_value(row.get("Current quarter")),
                "previous_quarter": _format_qoq_value(row.get("Previous quarter")),
                "change": _format_qoq_value(row.get("Change")),
                "qoq_percent": _format_qoq_value(row.get("QoQ %"), percent=True),
            }
        )
    return rows


def _companies_for_ai(
    companies: tuple[CompanyReportData, ...],
    months: tuple[tuple[int, int], ...],
) -> list[dict[str, object]]:
    return [
        {
            "company": company.company,
            "description": company.description,
            "existing_key_updates": list(company.key_updates),
            "metrics": _metric_context(company, months),
            "qoq_changes": _qoq_context(company),
            "warnings": list(company.warnings),
        }
        for company in companies
        if company.sheet_name
    ]


def _normalise_ai_updates(value: object) -> str:
    if isinstance(value, str):
        lines = value.splitlines()
    elif isinstance(value, list):
        lines = [str(item) for item in value]
    else:
        return ""
    return "\n".join(_strip_bullet(line) for line in lines if _strip_bullet(line))[:1200]


def _draft_key_updates_with_ai(
    companies: tuple[CompanyReportData, ...],
    months: tuple[tuple[int, int], ...],
) -> tuple[dict[str, str], str | None]:
    try:
        from openai import OpenAI
    except ImportError:
        return {}, "The OpenAI Python package is not installed. Add openai to requirements.txt and redeploy."

    api_key = _openai_setting("OPENAI_API_KEY")
    if not api_key:
        return {}, "OPENAI_API_KEY is not configured in Streamlit secrets or environment variables."

    model = _openai_setting("OPENAI_MODEL", DEFAULT_OPENAI_MODEL) or DEFAULT_OPENAI_MODEL
    period = quarter_label(months)
    payload = {
        "period": period,
        "companies": _companies_for_ai(companies, months),
    }
    client = OpenAI(api_key=api_key)
    prompt = (
        "You are drafting Sarmayacar quarterly investor report key updates from a Monthly Reporting workbook.\n"
        "Use only the supplied company descriptions, existing updates, warnings, and metric values.\n"
        "Do not invent numbers, customer names, fundraise details, causes, or forward-looking claims.\n"
        "For each company, write 2 to 4 concise bullets. Prefer the supplied QoQ changes when available, then the quarter metrics.\n"
        "If the metrics are missing, write one conservative bullet based on existing updates and one bullet saying data is pending review.\n"
        "Return only valid JSON with this schema: {\"companies\":[{\"company\":\"...\",\"key_updates\":[\"...\",\"...\"]}]}.\n\n"
        f"Workbook context:\n{json.dumps(payload, ensure_ascii=False)}"
    )

    try:
        response = client.responses.create(
            model=model,
            input=prompt,
            store=False,
        )
    except Exception as exc:
        return {}, f"AI drafting failed: {exc}"

    output_text = getattr(response, "output_text", "") or ""
    try:
        parsed = json.loads(output_text)
    except json.JSONDecodeError:
        return {}, "AI drafting returned text that was not valid JSON. Try again or use a different OPENAI_MODEL."

    drafts: dict[str, str] = {}
    for item in parsed.get("companies", []):
        if not isinstance(item, dict):
            continue
        company = str(item.get("company", "")).strip()
        updates = _normalise_ai_updates(item.get("key_updates", []))
        if company and updates:
            drafts[company] = updates
    return drafts, None


def _apply_update_drafts(
    rows: list[dict[str, object]],
    drafts: dict[str, str],
    replace_existing: bool,
) -> list[dict[str, object]]:
    updated: list[dict[str, object]] = []
    for row in rows:
        if row.get("Section") != SECTION_PORTFOLIO_COMPANY:
            updated.append(row)
            continue
        company = str(row.get("Item") or row.get("Company", ""))
        draft = drafts.get(company)
        if not draft:
            updated.append(row)
            continue
        next_row = dict(row)
        existing = str(next_row.get("Key updates", "")).strip()
        next_row["Key updates"] = draft if replace_existing or not existing else f"{existing}\n{draft}"
        updated.append(next_row)
    return updated


def _as_bool(value: object, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"yes", "y", "true", "1", "include", "included"}:
        return True
    if text in {"no", "n", "false", "0", "exclude", "excluded"}:
        return False
    return default


def _split_update_text(value: object) -> list[str]:
    return [
        _strip_bullet(line)
        for line in str(value or "").replace("\r", "\n").splitlines()
        if _strip_bullet(line)
    ]


def _commentary_key(section: str, item: str) -> str:
    return f"{section.strip()}::{item.strip()}"


def _highlight_review_rows() -> list[dict[str, object]]:
    return [
        {
            "Section": SECTION_HIGHLIGHTS,
            "Include": True,
            "Company": item,
            "Item": item,
            "Description": guidance,
            "Key updates": "",
            "Notes": "",
        }
        for item, guidance in HIGHLIGHT_ROWS
    ]


def _operator_angel_review_rows() -> list[dict[str, object]]:
    return [
        {
            "Section": SECTION_OPERATOR_ANGEL,
            "Include": True,
            "Company": company,
            "Item": company,
            "Description": "",
            "Key updates": "",
            "Notes": "",
        }
        for company in OPERATOR_ANGEL_COMPANIES
    ]


def _review_rows(companies: tuple[CompanyReportData, ...]) -> list[dict[str, object]]:
    portfolio_rows = [
        {
            "Section": SECTION_PORTFOLIO_COMPANY,
            "Include": bool(company.sheet_name),
            "Company": company.company,
            "Item": company.company,
            "Description": company.description,
            "Key updates": "\n".join(company.key_updates),
            "Notes": "",
        }
        for company in companies
    ]
    return [*_highlight_review_rows(), *portfolio_rows, *_operator_angel_review_rows()]


def _commentary_rows_from_workbook(workbook_path: Path) -> dict[str, dict[str, object]]:
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheet_name = workbook_sheet_name(workbook, COMMENTARY_SHEET_NAME)
        if sheet_name is None:
            return {}
        sheet = workbook[sheet_name]
        headers = {
            _normalise(sheet.cell(row=1, column=col).value or ""): col
            for col in range(1, sheet.max_column + 1)
        }
        item_col = headers.get("item") or headers.get("company")
        if item_col is None:
            return {}

        def value(row: int, header: str) -> object:
            col = headers.get(_normalise(header))
            return sheet.cell(row=row, column=col).value if col else None

        rows: dict[str, dict[str, object]] = {}
        for row in range(2, sheet.max_row + 1):
            item = _cell_text(value(row, "Item")) or _cell_text(value(row, "Company"))
            if not item:
                continue
            section = _cell_text(value(row, "Section")) or SECTION_PORTFOLIO_COMPANY
            updates = [
                _cell_text(value(row, f"Bullet {index}")) or _cell_text(value(row, f"Key Update {index}"))
                for index in range(1, MAX_COMMENTARY_BULLETS + 1)
            ]
            key = _commentary_key(section, item)
            rows[key] = {
                "Section": section,
                "Include": _as_bool(value(row, "Include"), default=True),
                "Company": item,
                "Item": item,
                "Description": _cell_text(value(row, "Description")),
                "Key updates": "\n".join(update for update in updates if update),
                "Notes": _cell_text(value(row, "Notes")),
                "Reviewer Status": _cell_text(value(row, "Reviewer Status")),
            }
        return rows
    finally:
        workbook.close()


def _apply_commentary_rows(
    base_rows: list[dict[str, object]],
    commentary_rows: dict[str, dict[str, object]],
) -> list[dict[str, object]]:
    if not commentary_rows:
        return base_rows
    by_key = {_normalise(key): row for key, row in commentary_rows.items()}
    merged: list[dict[str, object]] = []
    for row in base_rows:
        key = _commentary_key(str(row.get("Section", SECTION_PORTFOLIO_COMPANY)), str(row.get("Company", "") or row.get("Item", "")))
        legacy_key = str(row.get("Company", "") or row.get("Item", ""))
        override = by_key.get(_normalise(key)) or by_key.get(_normalise(legacy_key))
        if not override:
            merged.append(row)
            continue
        next_row = dict(row)
        for key in ("Include", "Description", "Key updates", "Notes"):
            value = override.get(key)
            if key == "Include" or value not in (None, ""):
                next_row[key] = value
        merged.append(next_row)
    return merged


def build_monthly_workbook_with_commentary_sheet(
    workbook_path: Path,
    review_rows: list[dict[str, object]],
    months: tuple[tuple[int, int], ...],
) -> bytes:
    workbook = load_workbook(workbook_path)
    try:
        existing = workbook_sheet_name(workbook, COMMENTARY_SHEET_NAME)
        if existing:
            del workbook[existing]
        sheet = workbook.create_sheet(COMMENTARY_SHEET_NAME, 0)
        headers = [
            "Section",
            "Include",
            "Item",
            "Description",
            *[f"Bullet {index}" for index in range(1, MAX_COMMENTARY_BULLETS + 1)],
            "Notes",
            "Reviewer Status",
        ]
        green_fill = PatternFill("solid", fgColor="008C78")
        gray_fill = PatternFill("solid", fgColor="F3F5F7")
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.fill = green_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_index, row in enumerate(review_rows, start=2):
            updates = _split_update_text(row.get("Key updates"))
            values = [
                row.get("Section", SECTION_PORTFOLIO_COMPANY),
                "Yes" if row.get("Include", True) else "No",
                row.get("Item") or row.get("Company", ""),
                row.get("Description", ""),
                *updates[:MAX_COMMENTARY_BULLETS],
            ]
            while len(values) < 4 + MAX_COMMENTARY_BULLETS:
                values.append("")
            values.extend([row.get("Notes", ""), "Draft"])
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=col, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = gray_fill

        widths = [24, 12, 28, 72, *([58] * MAX_COMMENTARY_BULLETS), 48, 20]
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"
        sheet["A" + str(len(review_rows) + 4)] = (
            f"Generated by FundOps Quarterly Report for {quarter_label(months)}. "
            "Edit this sheet, then re-upload the workbook in the Quarterly Report tab."
        )
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")
    return slug or "company"


def _review_updates(row: dict[str, object], limit: int | None = MAX_BULLETS) -> list[str]:
    updates = [
        _strip_bullet(line)
        for line in str(row.get("Key updates", "")).replace("\r", "\n").splitlines()
        if _strip_bullet(line)
    ]
    return updates if limit is None else updates[:limit]


def _figma_metric_rows(
    company: CompanyReportData,
    months: tuple[tuple[int, int], ...],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for index, metric in enumerate(company.metrics[:MAX_FIGMA_METRIC_ROWS], start=1):
        rows.append(
            {
                "index": index,
                "label": metric.label,
                "values": [
                    {
                        "month": f"{month_abbr[month]} {str(year)[-2:]}",
                        "value": _format_metric_value(metric.label, metric.values.get((year, month))),
                    }
                    for year, month in months
                ],
            }
        )
    return rows


def _figma_qoq_rows(company: CompanyReportData) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for index, row in enumerate(company.qoq_changes[:6], start=1):
        rows.append(
            {
                "index": index,
                "metric": row.get("Metric") or "",
                "method": row.get("Method") or "",
                "current_quarter": _format_qoq_value(row.get("Current quarter")) or "",
                "previous_quarter": _format_qoq_value(row.get("Previous quarter")) or "",
                "change": _format_qoq_value(row.get("Change")) or "",
                "qoq_percent": _format_qoq_value(row.get("QoQ %"), percent=True) or "",
            }
        )
    return rows


MONTH_YEAR_RE = re.compile(r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[ -]?\d{2,4}$", re.IGNORECASE)


def _format_exhibit_header(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%b-%y")
    if value is None:
        return ""
    text = str(value).strip()
    if MONTH_YEAR_RE.match(text):
        month_part, year_part = re.split(r"[ -]", text, maxsplit=1)
        return f"{month_part[:3].title()}-{year_part[-2:]}"
    return ""


def _format_exhibit_value(value: object, *, ratio_style: str | None = None) -> str:
    if value is None or value == "":
        return "-"
    if isinstance(value, (datetime, date)):
        return value.strftime("%b-%y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if ratio_style == "percent":
            return f"{value * 100:.1f}%"
        if ratio_style == "multiple":
            return f"{value:.2f}x"
        return f"{int(round(value)):,}"
    return str(value).strip()


def _date_columns_in_range(sheet: Worksheet, row: int, start_col: int, end_col: int) -> list[tuple[int, str]]:
    date_cols: list[tuple[int, str]] = []
    for col in range(start_col, min(sheet.max_column, end_col) + 1):
        label = _format_exhibit_header(sheet.cell(row=row, column=col).value)
        if label:
            date_cols.append((col, label))
    return date_cols


def _financial_label(sheet: Worksheet, row: int, label_cols: Iterable[int]) -> str:
    for col in label_cols:
        text = _cell_text(sheet.cell(row=row, column=col).value)
        if text:
            return text
    return ""


def _find_financial_heading(sheet: Worksheet, heading: str) -> tuple[int, int] | None:
    target = _normalise(heading)
    for row in range(1, min(sheet.max_row, 180) + 1):
        for col in range(1, min(sheet.max_column, 60) + 1):
            if _normalise(_cell_text(sheet.cell(row=row, column=col).value)) == target:
                return row, col
    return None


def _sheet_financial_score(sheet: Worksheet, target_sheet_tokens: set[str]) -> int:
    score = 0
    headings = {
        "balance sheet": _find_financial_heading(sheet, "Balance Sheet"),
        "income statement": _find_financial_heading(sheet, "Income Statement"),
        "financial summary": _find_financial_heading(sheet, "Financial Summary"),
    }
    score += sum(100 for location in headings.values() if location is not None)
    title = " ".join(sheet.title.lower().strip().split())
    if title in target_sheet_tokens:
        score += 80
    elif any(token in title for token in target_sheet_tokens):
        score += 40
    if "irr" in title:
        score -= 200
    balance = headings["balance sheet"]
    summary = headings["financial summary"]
    if balance and summary and balance[0] == summary[0] and balance[1] < summary[1]:
        header_row = balance[0] + 4
        left_headers = _date_columns_in_range(sheet, header_row, balance[1], summary[1] - 1)
        right_headers = _date_columns_in_range(sheet, header_row, summary[1] + 1, sheet.max_column)
        score += min(len(left_headers), 10) * 4
        score += min(len(right_headers), 10) * 4
    return score


def _financial_sheet_for_months(workbook, months: tuple[tuple[int, int], ...]) -> str | None:
    if not months:
        return None
    year, month = months[-1]
    wanted_tokens = {
        f"{month_name[month].lower()} {year}",
        f"{month_name[month].lower()} {str(year)[-2:]}",
        f"{month_abbr[month].lower()} {year}",
        f"{month_abbr[month].lower()} {str(year)[-2:]}",
    }
    scored = [
        (_sheet_financial_score(workbook[sheet_name], wanted_tokens), index, sheet_name)
        for index, sheet_name in enumerate(workbook.sheetnames)
    ]
    viable = [candidate for candidate in scored if candidate[0] >= 220]
    if not viable:
        return None
    return max(viable, key=lambda candidate: (candidate[0], -candidate[1]))[2]


def parse_financial_exhibits_workbook(
    workbook_path: Path,
    months: tuple[tuple[int, int], ...],
    max_rows_per_section: int = 34,
) -> tuple[dict[str, object] | None, list[str]]:
    warnings: list[str] = []
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheet_name = _financial_sheet_for_months(workbook, months)
        if sheet_name is None:
            return None, ["No financial exhibit sheets found."]
        sheet = workbook[sheet_name]
        sections: list[dict[str, object]] = []
        balance_heading = _find_financial_heading(sheet, "Balance Sheet")
        income_heading = _find_financial_heading(sheet, "Income Statement")
        summary_heading = _find_financial_heading(sheet, "Financial Summary")
        if balance_heading is None or income_heading is None or summary_heading is None:
            return None, [f"`{sheet_name}` does not look like a financial exhibits sheet."]

        balance_row, balance_col = balance_heading
        income_row, income_col = income_heading
        summary_row, summary_col = summary_heading
        header_row = balance_row + 4
        left_date_cols = _date_columns_in_range(sheet, header_row, balance_col, summary_col - 1)[-7:]
        right_date_cols = _date_columns_in_range(sheet, header_row, summary_col + 1, sheet.max_column)[-7:]
        if not left_date_cols:
            warnings.append(f"No financial statement date columns found on row {header_row}.")
        if not right_date_cols:
            warnings.append(f"No financial summary date columns found on row {header_row}.")
        expected_period = f"{month_abbr[months[-1][1]]}-{str(months[-1][0])[-2:]}" if months else ""
        periods = [period for _col, period in (right_date_cols or left_date_cols)]
        if expected_period and periods and expected_period not in periods:
            warnings.append(f"Financial exhibits do not include the selected quarter period `{expected_period}`.")

        def append_statement_section(section_name: str, start_row: int, end_row: int) -> None:
            rows: list[dict[str, object]] = []
            for row in range(start_row, min(end_row, sheet.max_row) + 1):
                label = _financial_label(sheet, row, (balance_col, balance_col + 1, balance_col + 2))
                if not label:
                    continue
                label_key = _normalise(label)
                if label_key in {_normalise(section_name), "in usd $ 000's", "investments at fair value"}:
                    continue
                values = [
                    {"period": period, "value": _format_exhibit_value(sheet.cell(row=row, column=col).value)}
                    for col, period in left_date_cols
                ]
                if not any(value["value"] != "-" for value in values):
                    continue
                rows.append({"label": label, "values": values})
                if len(rows) >= max_rows_per_section:
                    break
            if rows:
                sections.append({"section": section_name, "rows": rows})

        append_statement_section("Balance Sheet", balance_row + 5, income_row - 2)
        append_statement_section("Income Statement", income_row + 4, min(sheet.max_row, income_row + 85))

        summary_rows: list[dict[str, object]] = []
        summary_headings = {
            "commitments, drawdowns and distributions",
            "% of committed capital",
            "portfolio investments and nav",
            "performance metrics",
            "carried interest",
        }
        summary_section = ""
        for row in range(summary_row + 5, min(sheet.max_row, summary_row + 45) + 1):
            label = _financial_label(sheet, row, (summary_col,))
            if not label:
                continue
            label_key = _normalise(label)
            if label_key in summary_headings:
                summary_section = label_key
                summary_rows.append({"label": label, "subsection": True, "values": []})
                continue
            ratio_style = None
            if summary_section == "% of committed capital":
                ratio_style = "percent"
            elif summary_section == "performance metrics":
                ratio_style = "percent" if "irr" in label_key else "multiple"
            values = [
                {
                    "period": period,
                    "value": _format_exhibit_value(sheet.cell(row=row, column=col).value, ratio_style=ratio_style),
                }
                for col, period in right_date_cols
            ]
            if not any(value["value"] != "-" for value in values):
                continue
            summary_rows.append({"label": label, "values": values})
            if len(summary_rows) >= max_rows_per_section:
                break
        if summary_rows:
            sections.append({"section": "Financial Summary", "rows": summary_rows})

        sections = [
            section
            for section in sections
            if isinstance(section.get("rows"), list) and section["rows"]
        ]
        if not sections:
            warnings.append("No usable financial exhibit rows found.")
        return {
            "source_sheet": sheet_name,
            "periods": periods,
            "sections": sections,
        }, warnings
    finally:
        workbook.close()


def build_figma_data_pack(
    companies: tuple[CompanyReportData, ...],
    review_rows: list[dict[str, object]],
    months: tuple[tuple[int, int], ...],
    warnings: list[dict[str, str]],
    financial_exhibits: dict[str, object] | None = None,
) -> dict[str, object]:
    companies_by_name = {company.company: company for company in companies}
    report_companies: list[dict[str, object]] = []
    highlight_sections: list[dict[str, object]] = []
    operator_angels: list[dict[str, object]] = []
    for row in review_rows:
        if not row.get("Include", True):
            continue
        section = str(row.get("Section", SECTION_PORTFOLIO_COMPANY)).strip() or SECTION_PORTFOLIO_COMPANY
        item_name = str(row.get("Item") or row.get("Company", "")).strip()
        if section == SECTION_HIGHLIGHTS:
            highlight_sections.append(
                {
                    "name": item_name,
                    "description": str(row.get("Description", "")).strip(),
                    "bullets": _review_updates(row, limit=None),
                    "notes": str(row.get("Notes", "")).strip(),
                }
            )
            continue
        if section == SECTION_OPERATOR_ANGEL:
            operator_angels.append(
                {
                    "name": item_name,
                    "slug": _slug(item_name),
                    "description": str(row.get("Description", "")).strip(),
                    "key_updates": _review_updates(row),
                    "notes": str(row.get("Notes", "")).strip(),
                }
            )
            continue
        company_name = item_name
        company = companies_by_name.get(company_name)
        if company is None:
            continue
        report_companies.append(
            {
                "name": company.company,
                "slug": _slug(company.company),
                "description": str(row.get("Description", "")).strip(),
                "key_updates": _review_updates(row),
                "notes": str(row.get("Notes", "")).strip(),
                "metrics": _figma_metric_rows(company, months),
                "qoq_changes": _figma_qoq_rows(company),
                "warnings": [
                    warning["Warning"]
                    for warning in warnings
                    if warning.get("Company") == company.company
                ],
            }
        )

    payload = {
        "schema_version": "sarmayacar.quarterly_report.v1",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "report": {
            "quarter": quarter_label(months),
            "months": [f"{month_abbr[month]} {str(year)[-2:]}" for year, month in months],
            "source": "Sarmayacar FundOps Streamlit Quarterly Report tab",
        },
        "figma": {
            "template_frame": "Company Page Template",
            "layer_naming": {
                "company.name": "Company name text layer",
                "company.description": "Company description text layer",
                "company.update_1": "First update bullet text layer",
                "company.update_2": "Second update bullet text layer",
                "company.update_3": "Third update bullet text layer",
                "company.update_4": "Fourth update bullet text layer",
                "metric_1.label": "Metric label text layer",
                "metric_1.month_1": "Metric first month value",
                "metric_1.month_2": "Metric second month value",
                "metric_1.month_3": "Metric third month value",
                "metric_1.month_4": "Metric fourth month value",
                "metric_1.month_5": "Metric fifth month value",
                "metric_1.month_6": "Metric sixth month value",
                "qoq_1.metric": "QoQ metric name",
                "qoq_1.percent": "QoQ percentage change",
                "notes.disclaimer": "Notes or disclaimer text layer",
            },
        },
        "highlights": highlight_sections,
        "companies": report_companies,
        "operator_angels": operator_angels,
        "financial_exhibits": financial_exhibits,
    }
    return payload


def _json_download_bytes(payload: dict[str, object]) -> bytes:
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def _add_textbox(slide, x, y, w, h, text, font_size=12, bold=False, color="1F2937"):
    from pptx.util import Inches, Pt
    from pptx.enum.text import MSO_AUTO_SIZE

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor

    value = hex_color.strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _add_rect(slide, x, y, w, h, fill, line="FFFFFF", transparency=0):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.fill.transparency = transparency
    shape.line.color.rgb = _rgb(line)
    return shape


def _add_metric_table(slide, company: CompanyReportData, months: tuple[tuple[int, int], ...]) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    rows = max(2, len(company.metrics) + 1)
    cols = len(months) + 1
    table_shape = slide.shapes.add_table(rows, cols, Inches(6.35), Inches(2.0), Inches(6.35), Inches(3.85))
    table = table_shape.table
    widths = [2.75, 1.2, 1.2, 1.2]
    for idx, width in enumerate(widths[:cols]):
        table.columns[idx].width = Inches(width)

    headers = ["Metric", *[f"{month_abbr[month]} {str(year)[-2:]}" for year, month in months]]
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb("123D40")
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER if col else PP_ALIGN.LEFT
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = _rgb("FFFFFF")

    for row_idx, metric in enumerate(company.metrics, start=1):
        values = [metric.label, *[_format_metric_value(metric.label, metric.values.get(month)) for month in months]]
        for col, value in enumerate(values):
            cell = table.cell(row_idx, col)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb("F7F9F8") if row_idx % 2 else _rgb("FFFFFF")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.RIGHT if col else PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.2)
                    run.font.color.rgb = _rgb("1F2937")


def _add_bullets(slide, updates: Iterable[str]) -> None:
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(0.72), Inches(3.7), Inches(5.25), Inches(2.35))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for idx, update in enumerate(list(updates)[:MAX_BULLETS]):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = _strip_bullet(update)
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(12)
        paragraph.space_after = Pt(7)
        paragraph._p.get_or_add_pPr().set("marL", "171450")
        paragraph._p.get_or_add_pPr().set("indent", "-171450")


def build_company_pages_pptx(
    companies: tuple[CompanyReportData, ...],
    review_rows: list[dict[str, object]],
    months: tuple[tuple[int, int], ...],
    output_path: Path,
) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("Quarterly Report generation needs python-pptx. Add python-pptx to requirements.txt and install dependencies.") from exc

    data_by_company = {company.company: company for company in companies}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    report_label = quarter_label(months)

    for review in review_rows:
        if not review.get("Include", True):
            continue
        if review.get("Section", SECTION_PORTFOLIO_COMPANY) != SECTION_PORTFOLIO_COMPANY:
            continue
        company_name = str(review.get("Item") or review.get("Company", ""))
        company = data_by_company.get(company_name)
        if company is None:
            continue

        slide = prs.slides.add_slide(blank_layout)
        _add_rect(slide, 0, 0, 13.333, 7.5, "FFFFFF")
        _add_rect(slide, 0, 0, 13.333, 0.24, "123D40", line="123D40")
        _add_rect(slide, 0, 7.2, 13.333, 0.3, "123D40", line="123D40")
        _add_rect(slide, 0.55, 0.62, 0.08, 5.55, "C49A5A", line="C49A5A")

        _add_textbox(slide, 0.72, 0.48, 7.6, 0.45, company.company, font_size=26, bold=True, color="123D40")
        _add_textbox(slide, 10.85, 0.58, 1.75, 0.28, report_label, font_size=11, bold=True, color="5B6770")
        _add_textbox(slide, 0.72, 1.16, 5.2, 0.24, "Company Description", font_size=10.5, bold=True, color="C49A5A")
        _add_textbox(slide, 0.72, 1.48, 5.25, 1.55, str(review.get("Description", "")).strip(), font_size=11.2, color="1F2937")
        _add_textbox(slide, 0.72, 3.36, 5.2, 0.24, "Key Updates", font_size=10.5, bold=True, color="C49A5A")
        updates = str(review.get("Key updates", "")).replace("\r", "\n").splitlines()
        _add_bullets(slide, [update for update in updates if update.strip()])

        _add_textbox(slide, 6.35, 1.16, 4.4, 0.26, "Key Financial & Operating Metrics", font_size=10.5, bold=True, color="C49A5A")
        _add_metric_table(slide, company, months)
        notes = str(review.get("Notes", "")).strip()
        if notes:
            _add_textbox(slide, 6.35, 6.05, 5.95, 0.38, f"Notes: {notes}", font_size=8.8, color="5B6770")
        _add_textbox(slide, 0.72, 6.86, 8.6, 0.18, "Source: Monthly Reporting workbook. Draft for management review.", font_size=7.5, color="FFFFFF")

        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Aptos"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def _warning_rows(
    companies: tuple[CompanyReportData, ...],
    edited_rows: list[dict[str, object]] | None = None,
) -> list[dict[str, str]]:
    by_company = {
        str(row.get("Item") or row.get("Company", "")): row
        for row in edited_rows or []
        if row.get("Section", SECTION_PORTFOLIO_COMPANY) == SECTION_PORTFOLIO_COMPANY
    }
    rows: list[dict[str, str]] = []
    for company in companies:
        warnings = list(company.warnings)
        if company.company in by_company:
            row = by_company[company.company]
            description = str(row.get("Description", "")).strip()
            updates = tuple(
                _strip_bullet(line)
                for line in str(row.get("Key updates", "")).splitlines()
                if line.strip()
            )
            warnings = [warning for warning in warnings if not warning.startswith(("Description", "Missing company", "Missing key", "Too many", "Key updates", "One or more bullets"))]
            warnings.extend(_text_warnings(description, updates))
        for warning in warnings:
            rows.append({"Company": company.company, "Warning": warning})
    return rows


def _latest_quarterly_outputs(limit: int = 8) -> list[dict[str, str]]:
    if not QUARTERLY_OUTPUTS_DIR.exists():
        return []
    files = sorted(QUARTERLY_OUTPUTS_DIR.glob("*.pptx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return [
        {
            "File": path.name,
            "Modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%b %d, %Y %H:%M"),
            "Folder": str(path.parent),
        }
        for path in files[:limit]
    ]


def render_quarterly_report_page() -> None:
    st.header("Quarterly Report")
    st.caption("Generate editable company-page slides from the completed Monthly Reporting workbook. The Adobe XD PDF is used only as a format reference.")

    first_historical_year = 2021
    year_choices = list(range(first_historical_year, date.today().year + 16))
    quarter_end_options = [("Mar", 3), ("Jun", 6), ("Sep", 9), ("Dec", 12)]

    with st.container(border=True):
        st.subheader("Run setup")
        setup_cols = st.columns([1, 1, 2])
        with setup_cols[0]:
            selected_year = st.selectbox(
                "Year",
                options=year_choices,
                index=year_choices.index(2026) if 2026 in year_choices else len(year_choices) - 1,
                key="quarterly_report_year",
            )
        with setup_cols[1]:
            selected_quarter_label = st.selectbox(
                "Quarter ending",
                options=[label for label, _month in quarter_end_options],
                index=0,
                key="quarterly_report_quarter_end",
            )
        months = quarter_months(selected_year, dict(quarter_end_options)[selected_quarter_label])
        table_months = company_page_months(months)
        with setup_cols[2]:
            st.metric("Report quarter", quarter_label(months))

        monthly_report_file = st.file_uploader(
            "Completed Monthly Reporting workbook",
            type=["xlsx"],
            key="quarterly_monthly_report_file",
        )
        financial_exhibits_file = st.file_uploader(
            "Financial exhibits workbook (optional)",
            type=["xlsx"],
            key="quarterly_financial_exhibits_file",
        )

    if monthly_report_file is None:
        st.info("Upload the completed Monthly Reporting workbook to build the review screen.")
        latest_rows = _latest_quarterly_outputs()
        if latest_rows:
            with st.expander("Latest generated quarterly files", expanded=False):
                st.dataframe(latest_rows, hide_index=True, width="stretch")
        return

    financial_exhibits = None
    financial_exhibit_warnings: list[str] = []

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        workbook_path = tmp_dir / "monthly_reporting.xlsx"
        workbook_path.write_bytes(monthly_report_file.getbuffer())
        with st.spinner("Reading company pages from the monthly workbook"):
            companies = parse_monthly_reporting_workbook(workbook_path, months, metric_months=table_months)
            commentary_rows = _commentary_rows_from_workbook(workbook_path)
            base_review_rows = _apply_commentary_rows(_review_rows(companies), commentary_rows)
            commentary_workbook_bytes = build_monthly_workbook_with_commentary_sheet(
                workbook_path,
                base_review_rows,
                months,
            )
        if financial_exhibits_file is not None:
            exhibits_path = tmp_dir / "financial_exhibits.xlsx"
            exhibits_path.write_bytes(financial_exhibits_file.getbuffer())
            with st.spinner("Reading financial exhibits workbook"):
                financial_exhibits, financial_exhibit_warnings = parse_financial_exhibits_workbook(
                    exhibits_path,
                    months,
                )

    parsed_count = sum(1 for company in companies if company.sheet_name)
    qoq_signal_count = sum(len(company.qoq_changes) for company in companies)
    warning_rows = _warning_rows(companies)

    with st.container(border=True):
        st.subheader("Workbook scan")
        metric_cols = st.columns(4)
        metric_cols[0].metric("Company sheets found", parsed_count)
        metric_cols[1].metric("Companies in scope", len(companies))
        metric_cols[2].metric("QoQ signals", qoq_signal_count)
        metric_cols[3].metric("Warnings", len(warning_rows))
        if financial_exhibits_file is not None:
            if financial_exhibits:
                sections = financial_exhibits.get("sections", [])
                st.success(f"Financial exhibits loaded from `{financial_exhibits.get('source_sheet')}` with {len(sections)} sections.")
            if financial_exhibit_warnings:
                st.warning(" ".join(financial_exhibit_warnings))
        if warning_rows:
            st.caption("Warnings are review flags and do not block PPTX generation.")
            st.dataframe(warning_rows, hide_index=True, width="stretch")
        else:
            st.success("No missing-data or overflow warnings detected.")

    with st.container(border=True):
        st.subheader("Human review")
        st.caption("Use the workbook commentary sheet as the source of truth, or generate AI draft updates from workbook metrics and review them here.")
        review_state_key = f"quarterly_report_rows_{quarter_label(months)}_{monthly_report_file.name}"
        if review_state_key not in st.session_state:
            st.session_state[review_state_key] = base_review_rows

        commentary_cols = st.columns([1, 2])
        with commentary_cols[0]:
            st.download_button(
                "Download workbook with commentary sheet",
                data=commentary_workbook_bytes,
                file_name=f"{Path(monthly_report_file.name).stem} - with Quarterly Commentary.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
                key=f"{review_state_key}_download_commentary_workbook",
            )
        with commentary_cols[1]:
            if commentary_rows:
                st.success(f"Loaded {len(commentary_rows)} rows from `{COMMENTARY_SHEET_NAME}`.")
            else:
                st.caption(f"No `{COMMENTARY_SHEET_NAME}` sheet found yet. Download the workbook copy, edit the sheet, then re-upload it here.")

        ai_cols = st.columns([1, 1, 2])
        with ai_cols[0]:
            replace_ai_updates = st.checkbox(
                "Replace existing updates",
                value=True,
                key=f"{review_state_key}_replace_ai_updates",
            )
        with ai_cols[1]:
            draft_with_ai = st.button(
                "Draft key updates with AI",
                disabled=not _openai_api_key_configured(),
                width="stretch",
                key=f"{review_state_key}_draft_ai",
            )
        with ai_cols[2]:
            if _openai_api_key_configured():
                st.caption(f"AI model: `{_openai_setting('OPENAI_MODEL', DEFAULT_OPENAI_MODEL)}`")
            else:
                st.caption("Add `OPENAI_API_KEY` in Streamlit secrets to enable AI drafting.")

        if draft_with_ai:
            with st.spinner("Drafting company key updates from workbook metrics"):
                drafts, error = _draft_key_updates_with_ai(companies, months)
            if error:
                st.error(error)
            elif not drafts:
                st.warning("AI did not return any usable company updates.")
            else:
                st.session_state[review_state_key] = _apply_update_drafts(
                    st.session_state[review_state_key],
                    drafts,
                    replace_existing=replace_ai_updates,
                )
                st.success(f"Drafted key updates for {len(drafts)} companies. Review and edit before generating PPTX.")
                st.rerun()

        edited_rows = st.data_editor(
            st.session_state[review_state_key],
            hide_index=True,
            width="stretch",
            disabled=["Section", "Item", "Company"],
            column_order=["Section", "Include", "Item", "Description", "Key updates", "Notes"],
            key=f"quarterly_report_review_{quarter_label(months)}_{monthly_report_file.name}",
            column_config={
                "Section": st.column_config.TextColumn("Section", width="medium"),
                "Include": st.column_config.CheckboxColumn("Include"),
                "Item": st.column_config.TextColumn("Item", width="medium"),
                "Description": st.column_config.TextColumn("Description", width="large"),
                "Key updates": st.column_config.TextColumn("Key updates", width="large"),
                "Notes": st.column_config.TextColumn("Notes", width="medium"),
            },
        )

    edited_warnings = _warning_rows(companies, edited_rows)
    if edited_warnings:
        with st.expander("Warnings after edits", expanded=True):
            st.caption("Warnings are review flags and do not block PPTX generation.")
            st.dataframe(edited_warnings, hide_index=True, width="stretch")

    included_company_count = sum(
        1
        for row in edited_rows
        if row.get("Include") and row.get("Section", SECTION_PORTFOLIO_COMPANY) == SECTION_PORTFOLIO_COMPANY
    )
    included_operator_count = sum(
        1
        for row in edited_rows
        if row.get("Include") and row.get("Section") == SECTION_OPERATOR_ANGEL
    )
    figma_data_pack = build_figma_data_pack(
        companies=companies,
        review_rows=edited_rows,
        months=table_months,
        warnings=edited_warnings,
        financial_exhibits=financial_exhibits,
    )

    with st.container(border=True):
        st.subheader("Figma handoff")
        st.caption("Download this JSON and import it with the private Figma plugin inside the quarterly report template.")
        handoff_cols = st.columns(3)
        handoff_cols[0].metric("Portfolio companies", included_company_count)
        handoff_cols[1].metric("Operator angels", included_operator_count)
        handoff_cols[2].metric("Financial exhibit sections", len((financial_exhibits or {}).get("sections", [])))
        st.download_button(
            "Download Figma data pack",
            data=_json_download_bytes(figma_data_pack),
            file_name=f"Sarmayacar {quarter_label(months)} Figma Data Pack.json",
            mime="application/json",
            width="stretch",
        )

    generate = st.button(
        "Generate company pages PPTX",
        type="primary",
        disabled=included_company_count == 0,
        width="stretch",
    )
    if not generate:
        return

    filename = f"Sarmayacar {quarter_label(months)} Quarterly Company Pages.pptx"
    with tempfile.TemporaryDirectory() as tmp:
        output_path = Path(tmp) / filename
        try:
            build_company_pages_pptx(companies, edited_rows, table_months, output_path)
        except RuntimeError as exc:
            st.error(str(exc))
            return
        pptx_bytes = output_path.read_bytes()

    QUARTERLY_OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    permanent_output_path = QUARTERLY_OUTPUTS_DIR / filename
    permanent_output_path.write_bytes(pptx_bytes)

    with st.container(border=True):
        st.success("Generated! Your editable quarterly company-page deck is ready.")
        st.caption(f"Saved final file: {permanent_output_path}")
        st.download_button(
            "Download Quarterly Report PPTX",
            data=io.BytesIO(pptx_bytes),
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            width="stretch",
        )
