import unittest
from datetime import datetime

from openpyxl import Workbook

from quarterly_report import (
    COMPANY_METRIC_SECTION_MAP, CompanyReportData, _extract_metric_sections, _format_table_metric_value,
    _warning_rows, build_figma_data_pack,
)


MONTHS = tuple((2026, month) for month in range(1, 7))


def sheet_for(company):
    sheet = Workbook().active
    sheet.title = company
    for month in range(1, 7):
        col = month * 2 + 8
        sheet.cell(1, col, datetime(2026, month, 1))
        sheet.cell(2, col, "Actual")
        sheet.cell(2, col + 1, "Forecast")
    return sheet


def put(sheet, row, label, actual, forecast=None):
    sheet.cell(row, {"Oladoc": 1, "Tapmad": 3}.get(sheet.title, 2), label)
    for month in range(1, 7):
        sheet.cell(row, month * 2 + 8, actual)
        sheet.cell(row, month * 2 + 9, forecast)


def row_named(sections, label):
    return next(row for section in sections for row in section["rows"] if row["label"] == label)


class CompanyTableTests(unittest.TestCase):
    def test_simpaisa_reads_usd_block_and_total_income_includes_interest(self):
        sheet = sheet_for("SimPaisa")
        put(sheet, 19, "Total GTV", 14_004_968_831)
        put(sheet, 52, "Total GTV", 50_030_253.389, 54_728_000)
        put(sheet, 55, "Net Revenue", 817_217.58)
        put(sheet, 64, "Net Income (loss)", 251_407.86)
        put(sheet, 65, "Interest Income", 4_495.57)
        put(sheet, 66, "Total Income", 255_903.43)
        sections = _extract_metric_sections(sheet, sheet, MONTHS, "SimPaisa")
        self.assertEqual(len(sections), 1)
        self.assertEqual(len(sections[0]["rows"]), 9)
        gtv = row_named(sections, "GTV")
        self.assertEqual(gtv["currency"], "$")
        self.assertEqual(gtv["values"][0]["actual"], "50,030")
        self.assertEqual(gtv["values"][0]["forecast"], "54,728")
        self.assertEqual(gtv["values"][0]["source"]["actual"], "J52")
        self.assertEqual(row_named(sections, "Net Income")["values"][0]["actual"], "256")

    def test_operating_counts_are_not_divided_and_percentages_not_zeroed(self):
        sheet = sheet_for("Roomy")
        put(sheet, 13, "Revenue - Total", 317_419)
        put(sheet, 26, "Operational Rooms", 236, 224)
        put(sheet, 27, "Occupancy %", 0.48, 0.40)
        put(sheet, 28, "ADR", 77.802, 74)
        sections = _extract_metric_sections(sheet, sheet, MONTHS, "Roomy")
        self.assertEqual(row_named(sections, "Total Revenue")["values"][0]["actual"], "317")
        self.assertEqual(row_named(sections, "Operational Rooms")["values"][0]["actual"], "236")
        occupancy = row_named(sections, "Occupancy Rate (%)")
        self.assertEqual(occupancy["values"][0]["actual"], "48")
        self.assertEqual(occupancy["values"][0]["forecast"], "40")
        self.assertEqual(occupancy["currency"], "")

    def test_already_scaled_reporting_blocks_are_not_scaled_twice(self):
        for company, row, source, label, value, expected in (
            ("Tapmad", 46, "Gross Revenue", "Gross Revenue", 2173.314, "2,173"),
            ("OneLoad", 61, "GMV - Total", "GMV", 3519.673, "3,520"),
            ("Procheck", 18, "Total Revenue", "Total Revenue", 2942, "2,942"),
        ):
            with self.subTest(company=company):
                sheet = sheet_for(company)
                put(sheet, row, source, value)
                result = row_named(_extract_metric_sections(sheet, sheet, MONTHS, company), label)
                self.assertEqual(result["values"][0]["actual"], expected)

    def test_pkr_accounting_strings_and_detail_net_income_are_preserved(self):
        sheet = sheet_for("Dot & Line")
        put(sheet, 15, "GMV", "PKR 7,550,817")
        put(sheet, 17, "Other revenue", "-PKR 112,449")
        put(sheet, 24, "Net Income (loss)", "PKR 1,782,870")
        put(sheet, 89, "Net Income", -1_782_870)
        sections = _extract_metric_sections(sheet, sheet, MONTHS, "Dot & Line")
        gmv = row_named(sections, "GMV")
        self.assertEqual(gmv["currency"], "PKR")
        self.assertEqual(gmv["values"][0]["actual"], "7,551")
        self.assertEqual(row_named(sections, "Other Revenue")["values"][0]["actual"], "(112)")
        self.assertEqual(row_named(sections, "Net Burn")["values"][0]["actual"], "(1,783)")

    def test_full_company_row_contract_including_missing_rows(self):
        expected = {"SimPaisa": 9, "ABHI": 13, "Bykea": 13, "Oladoc": 14, "Tapmad": 13,
                    "Revolving Games": 8, "OneLoad": 11, "Jiye Technologies": 9,
                    "Roomy": 11, "Procheck": 13, "Dot & Line": 11}
        self.assertEqual(set(COMPANY_METRIC_SECTION_MAP), set(expected))
        for company, count in expected.items():
            with self.subTest(company=company):
                sheet = sheet_for(company)
                sections = _extract_metric_sections(sheet, sheet, MONTHS, company)
                self.assertEqual(sum(len(s["rows"]) for s in sections), count)
                self.assertTrue(all(s["show_forecast"] for s in sections))
                for section in sections:
                    for row in section["rows"]:
                        self.assertTrue(row["source_missing"])
                        self.assertTrue(all(v["actual"] == v["forecast"] == "-" for v in row["values"]))

    def test_zero_missing_decimal_and_accounting_rounding_are_distinct(self):
        for value in (None, True, "-", "#VALUE!", "#DIV/0!", "PKR -"):
            self.assertEqual(_format_table_metric_value(value, scale="whole"), "-")
        self.assertEqual(_format_table_metric_value(0, scale="whole"), "0")
        self.assertEqual(_format_table_metric_value(-1500, scale="thousands"), "(2)")
        self.assertEqual(_format_table_metric_value(1.818528745, scale="decimal"), "1.82")
        self.assertEqual(_format_table_metric_value(48, scale="percent"), "-")
        self.assertEqual(_format_table_metric_value(-10, scale="thousands"), "(0)")

    def test_review_row_order_does_not_change_approved_company_order(self):
        names = ("Dot & Line", "ABHI", "SimPaisa", "Bykea")
        companies = tuple(CompanyReportData(name, "test", "sheet", "description", (), (), (), (), ()) for name in names)
        review = [{"Item": name, "Description": "Reviewed", "Key updates": "Reviewed update"} for name in names]
        payload = build_figma_data_pack(companies, review, MONTHS, [])
        self.assertEqual([company["name"] for company in payload["companies"]], ["SimPaisa", "ABHI", "Bykea", "Dot & Line"])

    def test_old_quarter_commentary_is_flagged_without_rewriting_it(self):
        company = CompanyReportData("SimPaisa", "test", "sheet", "description", (), (), (), (), ())
        review = [{"Item": "SimPaisa", "Description": "Reviewed", "Key updates": "Q1 2026 revenue declined."}]
        warnings = _warning_rows((company,), review, ((2026, 4), (2026, 5), (2026, 6)))
        self.assertTrue(any("Commentary references Q1 2026, not Q2 2026" in row["Warning"] for row in warnings))
        self.assertEqual(review[0]["Key updates"], "Q1 2026 revenue declined.")
        review[0]["Key updates"] = "Q2 2026 grew compared with Q1 2026."
        self.assertFalse(any("Commentary references" in row["Warning"] for row in _warning_rows((company,), review, ((2026, 6),))))


if __name__ == "__main__":
    unittest.main()
