import io
import json
import unittest
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook

from quarterly_report import (
    PORTFOLIO_SUMMARY_ROWS,
    SECTION_PORTFOLIO_COMPANY,
    _format_exhibit_header,
    _format_portfolio_money,
    _format_portfolio_ownership,
    _money_to_number,
    _portfolio_key,
    _portfolio_summary_data,
    _read_portfolio_source,
    build_figma_data_pack,
    parse_financial_exhibits_workbook,
)


MONTHS = ((2026, 4), (2026, 5), (2026, 6))


class PortfolioSummaryTests(unittest.TestCase):
    def setUp(self):
        self.workbook = Workbook()
        self.addCleanup(self.workbook.close)
        self.sheet = sheet = self.workbook.active
        sheet.title = "June 2026"
        for cell, value in {
            "D8": "Balance Sheet", "D59": "Income Statement", "Y8": "Financial Summary",
            "R12": datetime(2025, 12, 31), "S12": datetime(2026, 6, 30),
            "AH12": datetime(2026, 6, 30), "E65": "Realized gain (loss)",
            "Y25": "Total invested in portfolio companies", "Y26": "Fair value of portfolio",
            "D89": "As at Mar 2026", "E93": "Segment", "F93": "Date",
            "K93": "SV Total Investment", "L93": "SV Total Ownership",
            "N93": "SV Total Value", "O93": "Unrealized Gains", "D126": "Total",
        }.items():
            sheet[cell] = value
        self.costs, self.carrying = [], []
        for index, row in enumerate(PORTFOLIO_SUMMARY_ROWS):
            balance_row, schedule_row = index + 14, index + 94
            cost = 100000 + index * 123
            carrying = 75000 + index * 31
            self.costs.append(cost)
            self.carrying.append(carrying)
            sheet.cell(balance_row, 6, row["source_label"])
            sheet.cell(balance_row, 18, 99999)
            sheet.cell(balance_row, 19, carrying / 1000)
            sheet.cell(schedule_row, 4, row["source_label"])
            sheet.cell(schedule_row, 6, datetime.strptime(row["first_cash_injection"], "%b-%y"))
            sheet.cell(schedule_row, 11, cost)
            sheet.cell(schedule_row, 12, 0.064)
            sheet.cell(schedule_row, 14, 999999)
            sheet.cell(schedule_row, 15, 888888)
        sheet["AH25"] = sum(self.costs) / 1000
        sheet["AH26"] = sum(self.carrying) / 1000

    def source(self, period="Jun-26"):
        return _read_portfolio_source(self.sheet, period)

    def summary(self, period="Jun-26"):
        return _portfolio_summary_data({"portfolio_source": self.source(period)}, period)

    def parse(self, max_rows=34):
        data = io.BytesIO()
        self.workbook.save(data)
        data.seek(0)
        return parse_financial_exhibits_workbook(data, MONTHS, max_rows_per_section=max_rows)

    def test_all_companies_map_current_cost_carrying_and_ownership(self):
        summary = self.summary()
        self.assertEqual(len(summary["rows"]), 26)
        for index, row in enumerate(summary["rows"]):
            self.assertEqual(row["name"], PORTFOLIO_SUMMARY_ROWS[index]["name"])
            self.assertEqual(row["amounts"]["cost"], self.costs[index])
            self.assertEqual(row["amounts"]["carrying_value"], self.carrying[index])
            self.assertEqual(row["amounts"]["unrealized_gain_loss"], self.carrying[index] - self.costs[index])
            self.assertEqual(row["ownership"], "6.4%")
            self.assertEqual(row["sources"]["cost"], f"June 2026!K{index + 94} (USD)")
            self.assertEqual(row["sources"]["carrying_value"], f"June 2026!S{index + 14} (USD thousands)")
            self.assertEqual(row["sources"]["ownership"], f"June 2026!L{index + 94}")

    def test_full_precision_conversion_and_accounting_format(self):
        for value, multiplier, expected in (
            (753.872, 1000, 753872), (115.955, 1000, 115955), (34.6, 1000, 34600),
            (0, 1000, 0), ("$(25,000)", 1, -25000), ("$1,234", 1, 1234),
            ("1.5", 1, 2), ("-1.5", 1, -2), (Decimal("1.2345"), 1000, 1235),
        ):
            with self.subTest(value=value):
                self.assertEqual(_money_to_number(value, multiplier), expected)
        self.assertEqual(_format_portfolio_money(-25000), "$(25,000)")
        self.assertEqual(_format_portfolio_money(0), "$0")
        self.assertEqual(_format_portfolio_money(None), "-")

    def test_invalid_values_do_not_become_zero(self):
        for value in (None, "", "-", "#DIV/0!", "#VALUE!", "unknown 100", True, datetime.now(), float("nan"), float("inf")):
            with self.subTest(value=value):
                self.assertIsNone(_money_to_number(value))

    def test_ownership_uses_fraction_or_explicit_percent(self):
        for value, expected in ((0, "0.0%"), (1, "100.0%"), (.116, "11.6%"),
                                (.0214285714, "2.1%"), ("6.45%", "6.5%")):
            self.assertEqual(_format_portfolio_ownership(value), expected)
        for value in (None, "#DIV/0!", True, -0.1, 1.1, "6.4", "NaN", "Infinity"):
            self.assertIsNone(_format_portfolio_ownership(value))

    def test_missing_or_invalid_ownership_has_no_approved_fallback(self):
        self.sheet["L94"] = "#DIV/0!"
        self.sheet["L95"] = None
        self.sheet["L96"] = 0
        summary = self.summary()
        self.assertEqual([row["ownership"] for row in summary["rows"][:3]], ["-", "-", "0.0%"])
        self.assertTrue(any("June 2026!L94 (#DIV/0!)" in warning for warning in summary["warnings"]))

    def test_approved_descriptions_dates_board_seats_and_order_remain(self):
        self.sheet["F94"] = datetime(2024, 1, 1)
        summary = self.summary()
        for expected, row in zip(PORTFOLIO_SUMMARY_ROWS, summary["rows"]):
            for field in ("name", "description", "first_cash_injection", "sv_board_seats", "section"):
                self.assertEqual(row[field], expected[field])
        self.assertTrue(any("kept approved first cash injection" in warning for warning in summary["warnings"]))
        self.assertIn("Approved report", summary["rows"][0]["sources"]["first_cash_injection"])

    def test_aliases_join_balance_and_investment_rows(self):
        aliases = {"AbhiFinance": "Abhi", "Procheck 1": "Procheck", "Buyzilla 2": "Buyzilla",
                   "Jiye Tech": "Jive Technologies", "Truckistan": "L.L.M.Bots",
                   "Raptr": "Raptr Games", "Bolandi": "Taraki", "Savvy": "Savvy Technologies"}
        for index, row in enumerate(PORTFOLIO_SUMMARY_ROWS):
            for original, alias in aliases.items():
                if _portfolio_key(row["source_label"]) == _portfolio_key(original):
                    self.sheet.cell(index + 94, 4, alias)
        summary = self.summary()
        self.assertTrue(all(row["amounts"]["cost"] is not None for row in summary["rows"]))
        self.assertTrue(all(check["matches"] for check in summary["checks"]))

    def test_portfolio_and_grand_totals_reconcile(self):
        summary = self.summary()
        indices = [i for i, row in enumerate(PORTFOLIO_SUMMARY_ROWS) if row["section"] == SECTION_PORTFOLIO_COMPANY]
        self.assertEqual(summary["totals"]["cost"], _format_portfolio_money(sum(self.costs)))
        self.assertEqual(summary["totals"]["carrying_value"], _format_portfolio_money(sum(self.carrying)))
        self.assertEqual(summary["totals"]["unrealized_gain_loss"], _format_portfolio_money(sum(self.carrying) - sum(self.costs)))
        self.assertEqual(summary["portfolio_totals"]["cost"], _format_portfolio_money(sum(self.costs[i] for i in indices)))
        self.assertEqual(summary["portfolio_totals"]["carrying_value"], _format_portfolio_money(sum(self.carrying[i] for i in indices)))
        self.assertEqual(len(summary["checks"]), 2)
        self.assertTrue(all(check["matches"] for check in summary["checks"]))

    def test_missing_amount_withholds_affected_total_not_partial_sum(self):
        self.sheet["S14"] = None
        summary = self.summary()
        self.assertEqual(summary["rows"][0]["carrying_value"], "-")
        self.assertEqual(summary["totals"]["carrying_value"], "-")
        self.assertEqual(summary["portfolio_totals"]["unrealized_gain_loss"], "-")
        self.assertNotEqual(summary["totals"]["cost"], "-")
        self.assertFalse(summary["checks"][1]["matches"])

    def test_duplicate_rows_withhold_ambiguous_fields(self):
        self.sheet["F40"] = PORTFOLIO_SUMMARY_ROWS[0]["source_label"]
        self.sheet["S40"] = 1
        self.sheet["D120"] = PORTFOLIO_SUMMARY_ROWS[0]["source_label"]
        self.sheet["K120"] = 1
        self.sheet["L120"] = .1
        summary = self.summary()
        row = summary["rows"][0]
        for field in ("cost", "carrying_value", "ownership", "unrealized_gain_loss"):
            self.assertEqual(row[field], "-")
        self.assertTrue(any("Duplicate investment row" in warning for warning in summary["warnings"]))

    def test_realized_gains_missing_are_not_assumed_zero(self):
        self.sheet["R65"] = 0
        summary = self.summary()
        self.assertTrue(all(row["realized_gain_loss"] == "-" for row in summary["rows"]))
        self.assertEqual(summary["totals"]["realized_gain_loss"], "-")

    def test_explicit_per_company_realized_zero_is_preserved(self):
        self.sheet["T93"] = "Realized Gain/Loss ($)"
        for index in range(26):
            self.sheet.cell(index + 94, 20, 0)
        summary = self.summary()
        self.assertTrue(all(row["realized_gain_loss"] == "$0" for row in summary["rows"]))
        self.assertEqual(summary["totals"]["realized_gain_loss"], "$0")

    def test_wrong_period_and_ambiguous_period_do_not_fallback(self):
        for period in ("Mar-26", "Dec-25"):
            with self.subTest(period=period):
                summary = self.summary(period)
                self.assertTrue(all(row["cost"] == row["carrying_value"] == row["ownership"] == "-" for row in summary["rows"]))
        self.sheet["R12"] = datetime(2026, 6, 30)
        self.assertFalse(self.source()["rows"])

    def test_pack_period_mismatch_discards_values_and_controls(self):
        summary = _portfolio_summary_data({"portfolio_source": self.source()}, "Sep-26")
        self.assertEqual(summary["rows"][0]["cost"], "-")
        self.assertEqual(summary["checks"], [])

    def test_reconciliation_discrepancy_is_not_silently_corrected(self):
        self.sheet["AH25"] = 1
        summary = self.summary()
        self.assertFalse(summary["checks"][0]["matches"])
        self.assertEqual(summary["totals"]["cost"], _format_portfolio_money(sum(self.costs)))
        self.assertTrue(any("does not reconcile" in warning for warning in summary["warnings"]))

    def test_missing_control_is_flagged(self):
        self.sheet["Y26"] = None
        summary = self.summary()
        self.assertEqual(len(summary["checks"]), 1)
        self.assertTrue(any("No carrying value control total" in warning for warning in summary["warnings"]))

    def test_missing_schedule_does_not_use_static_costs(self):
        self.sheet["K93"] = None
        summary = self.summary()
        self.assertEqual(summary["totals"]["cost"], "-")
        self.assertNotEqual(summary["totals"]["carrying_value"], "-")

    def test_all_rows_survive_exhibit_display_cap_and_json_handoff(self):
        exhibits, warnings = self.parse(max_rows=3)
        self.assertFalse(warnings)
        self.assertEqual(len(exhibits["sections"][0]["rows"]), 3)
        payload = build_figma_data_pack((), [], MONTHS, [], exhibits)
        payload = json.loads(json.dumps(payload, allow_nan=False))
        self.assertEqual(payload["report"]["quarter"], "Q2 2026")
        self.assertEqual(len(payload["portfolio_summary"]["rows"]), 26)
        self.assertEqual(payload["portfolio_summary"]["rows"][-1]["amounts"]["cost"], self.costs[-1])
        self.assertTrue(all(check["matches"] for check in payload["portfolio_summary"]["checks"]))

    def test_current_quarter_sheet_selected_over_historical_sheet(self):
        old = self.workbook.copy_worksheet(self.sheet)
        old.title = "Mar 2026"
        old["S12"] = old["AH12"] = datetime(2026, 3, 31)
        old["K94"] = 999
        self.workbook.move_sheet(old, offset=-1)
        exhibits, _ = self.parse()
        self.assertEqual(exhibits["source_sheet"], "June 2026")
        self.assertEqual(_portfolio_summary_data(exhibits, "Jun-26")["rows"][0]["amounts"]["cost"], self.costs[0])

    def test_missing_financial_file_keeps_order_but_no_fabricated_numbers(self):
        summary = _portfolio_summary_data(None, "Jun-26")
        self.assertEqual(len(summary["rows"]), 26)
        self.assertEqual(summary["totals"]["cost"], "-")
        self.assertTrue(all(row["ownership"] == "-" for row in summary["rows"]))

    def test_date_headers_accept_compact_and_separated_month_year(self):
        for value in ("Jun26", "Jun-26", "Jun 2026", datetime(2026, 6, 30)):
            self.assertEqual(_format_exhibit_header(value), "Jun-26")


if __name__ == "__main__":
    unittest.main()
