from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Callable, Mapping

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


PROJECT_ROOT = Path(__file__).resolve().parent


def _default_formula_memory_path() -> Path:
    mapped_path = PROJECT_ROOT / "Mappings" / "formula_memory.json"
    flat_path = PROJECT_ROOT / "formula_memory.json"
    if mapped_path.exists() or not flat_path.exists():
        return mapped_path
    return flat_path


DEFAULT_FORMULA_MEMORY_PATH = _default_formula_memory_path()
FORMULA_ERROR_TOKENS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A")


MonthKey = tuple[int, int]
FindMonthColumns = Callable[[Worksheet, tuple[MonthKey, ...]], dict[MonthKey, int]]
TranslateFormula = Callable[[str, str, str], str]


def _label_key(value: object) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    text = text.replace("\xa0", " ")
    return text or None


def blank_formula_memory() -> dict[str, object]:
    return {
        "version": 1,
        "source_workbook": "",
        "generated_at": "",
        "companies": {},
    }


def load_formula_memory(path: Path = DEFAULT_FORMULA_MEMORY_PATH) -> dict[str, object]:
    if not path.exists():
        return blank_formula_memory()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return blank_formula_memory()
    if not isinstance(data, dict):
        return blank_formula_memory()
    data.setdefault("version", 1)
    data.setdefault("source_workbook", "")
    data.setdefault("generated_at", "")
    data.setdefault("companies", {})
    if not isinstance(data["companies"], dict):
        data["companies"] = {}
    return data


def save_formula_memory(data: dict[str, object], path: Path = DEFAULT_FORMULA_MEMORY_PATH) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, sort_keys=True), encoding="utf-8")


def _formula_like(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _has_formula_error(value: object) -> bool:
    if not isinstance(value, str):
        return False
    upper = value.upper()
    return any(token in upper for token in FORMULA_ERROR_TOKENS)


def _nearest_formula_source(sheet: Worksheet, row: int, target_col: int) -> tuple[str, str] | None:
    # Actual and forecast columns are paired every two columns. Search the same
    # actual/forecast lane first, moving left because future blank columns should
    # inherit from completed history.
    for distance in range(2, sheet.max_column + 1, 2):
        for source_col in (target_col - distance, target_col + distance):
            if source_col < 1 or source_col > sheet.max_column:
                continue
            source_cell = sheet.cell(row=row, column=source_col)
            if _formula_like(source_cell.value):
                return source_cell.coordinate, str(source_cell.value)
    return None


def _stored_formula_source(row_memory: Mapping[str, object], kind: str, row: int) -> tuple[str, str] | None:
    examples = row_memory.get(f"{kind}_examples")
    if not isinstance(examples, list):
        return None
    for example in examples:
        if not _formula_like(example):
            continue
        match = re.search(r"\b([A-Z]+)\$?(\d+)\b", str(example))
        if not match:
            continue
        ref_col = column_index_from_string(match.group(1))
        ref_row = int(match.group(2))
        if ref_row == row:
            origin_col = ref_col + 1
        else:
            origin_col = ref_col
        return f"{get_column_letter(origin_col)}{row}", str(example)
    return None


def _formula_references_cell(formula: object, coordinate: str) -> bool:
    if not _formula_like(formula):
        return False
    column_match = re.match(r"([A-Z]+)(\d+)", coordinate)
    if not column_match:
        return False
    column, row = column_match.groups()
    pattern = rf"(?<![A-Z0-9_])\$?{re.escape(column)}\$?{re.escape(row)}(?![A-Z0-9_])"
    return re.search(pattern, str(formula), flags=re.IGNORECASE) is not None


def apply_formula_memory(
    workbook,
    months: tuple[MonthKey, ...],
    sheet_names: Mapping[str, str],
    find_month_columns: FindMonthColumns,
    translate_formula: TranslateFormula,
    path: Path = DEFAULT_FORMULA_MEMORY_PATH,
) -> dict[str, object]:
    memory = load_formula_memory(path)
    companies = memory.get("companies", {})
    if not isinstance(companies, dict) or not companies:
        return {
            "formulas_written": 0,
            "missing_formula_sources": 0,
            "companies": [],
        }

    formulas_written = 0
    static_values_overwritten = 0
    missing_formula_sources = 0
    company_summaries: list[dict[str, object]] = []

    for profile_name, company_memory in companies.items():
        if not isinstance(company_memory, dict):
            continue
        if profile_name not in sheet_names:
            continue
        sheet_name = sheet_names.get(profile_name) or str(company_memory.get("sheet_name", ""))
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        target_cols = find_month_columns(sheet, months)
        rows = company_memory.get("rows", {})
        if not isinstance(rows, dict):
            continue

        company_written = 0
        company_missing = 0
        for row_text, row_memory in rows.items():
            if not isinstance(row_memory, dict):
                continue
            try:
                row = int(row_text)
            except ValueError:
                continue
            if profile_name == "simpaisa":
                stored_label_key = _label_key(row_memory.get("label"))
                current_label_key = (
                    _label_key(sheet.cell(row=row, column=2).value)
                    or _label_key(sheet.cell(row=row, column=1).value)
                )
                if stored_label_key and stored_label_key != current_label_key:
                    continue
            for month_key, actual_col in target_cols.items():
                for kind, target_col in (("actual", actual_col), ("forecast", actual_col + 1)):
                    if not row_memory.get(kind):
                        continue
                    if target_col > sheet.max_column:
                        continue
                    target_cell = sheet.cell(row=row, column=target_col)
                    if isinstance(target_cell, MergedCell):
                        continue
                    if _formula_like(target_cell.value):
                        continue
                    source = _nearest_formula_source(sheet, row, target_col)
                    if source is None:
                        source = _stored_formula_source(row_memory, kind, row)
                    if source is None:
                        missing_formula_sources += 1
                        company_missing += 1
                        continue
                    source_ref, source_formula = source
                    if target_cell.value not in (None, ""):
                        static_values_overwritten += 1
                    translated = translate_formula(source_formula, source_ref, target_cell.coordinate)
                    if _formula_references_cell(translated, target_cell.coordinate):
                        missing_formula_sources += 1
                        company_missing += 1
                        continue
                    target_cell.value = translated
                    formulas_written += 1
                    company_written += 1

        if company_written or company_missing:
            company_summaries.append(
                {
                    "profile": profile_name,
                    "company": company_memory.get("display_name", profile_name),
                    "formulas_written": company_written,
                    "missing_formula_sources": company_missing,
                }
            )

    return {
        "formulas_written": formulas_written,
        "static_values_overwritten": static_values_overwritten,
        "missing_formula_sources": missing_formula_sources,
        "companies": company_summaries,
    }


def audit_formula_memory(
    workbook,
    months: tuple[MonthKey, ...],
    sheet_names: Mapping[str, str],
    find_month_columns: FindMonthColumns,
    path: Path = DEFAULT_FORMULA_MEMORY_PATH,
    sample_limit: int = 100,
) -> dict[str, object]:
    memory = load_formula_memory(path)
    companies = memory.get("companies", {})
    if not isinstance(companies, dict) or not companies:
        return {
            "issue_count": 0,
            "issues": [],
            "companies": [],
        }

    issues: list[dict[str, object]] = []
    company_summaries: list[dict[str, object]] = []

    for profile_name, company_memory in companies.items():
        if not isinstance(company_memory, dict):
            continue
        if profile_name not in sheet_names:
            continue
        sheet_name = sheet_names.get(profile_name) or str(company_memory.get("sheet_name", ""))
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        target_cols = find_month_columns(sheet, months)
        rows = company_memory.get("rows", {})
        if not isinstance(rows, dict):
            continue

        company_issue_count = 0
        for row_text, row_memory in rows.items():
            if not isinstance(row_memory, dict):
                continue
            try:
                row = int(row_text)
            except ValueError:
                continue
            if profile_name == "simpaisa":
                stored_label_key = _label_key(row_memory.get("label"))
                current_label_key = (
                    _label_key(sheet.cell(row=row, column=2).value)
                    or _label_key(sheet.cell(row=row, column=1).value)
                )
                if stored_label_key and stored_label_key != current_label_key:
                    continue
            label = row_memory.get("label", "")
            for month_key, actual_col in target_cols.items():
                for kind, target_col in (("actual", actual_col), ("forecast", actual_col + 1)):
                    if not row_memory.get(kind):
                        continue
                    if target_col > sheet.max_column:
                        continue
                    cell = sheet.cell(row=row, column=target_col)
                    if isinstance(cell, MergedCell):
                        continue
                    value = cell.value
                    issue_type = ""
                    if value in (None, ""):
                        issue_type = "blank_formula_cell"
                    elif not _formula_like(value):
                        issue_type = "static_value_in_formula_row"
                    elif _has_formula_error(value):
                        issue_type = "formula_error_token"
                    elif _formula_references_cell(value, cell.coordinate):
                        issue_type = "circular_self_reference"
                    if not issue_type:
                        continue
                    company_issue_count += 1
                    if len(issues) < sample_limit:
                        issues.append(
                            {
                                "profile": profile_name,
                                "company": company_memory.get("display_name", profile_name),
                                "sheet": sheet_name,
                                "cell": cell.coordinate,
                                "month": f"{month_key[0]}-{month_key[1]:02d}",
                                "kind": kind,
                                "row": row,
                                "label": label,
                                "issue": issue_type,
                                "value": value,
                            }
                        )

        if company_issue_count:
            company_summaries.append(
                {
                    "profile": profile_name,
                    "company": company_memory.get("display_name", profile_name),
                    "formula_issue_count": company_issue_count,
                }
            )

    return {
        "issue_count": sum(item["formula_issue_count"] for item in company_summaries),
        "issues": issues,
        "companies": company_summaries,
    }


def formula_memory_timestamp() -> str:
    return datetime.now().isoformat(timespec="seconds")
