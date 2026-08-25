from __future__ import annotations

import csv
import io
import json
import re
import tempfile
from calendar import month_abbr
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
import streamlit as st

from google_drive_connector import (
    build_drive_service,
    drive_api_available,
    list_folder_items,
    parse_drive_folder_id,
    service_account_email,
)
from intelligence_detector import run_intelligence_checks
from monthly_reporting_automation import (
    INACTIVE_COMPANIES,
    PENDING_COMPANIES,
    PROFILES,
    find_month_columns,
    run_batch_update,
    validate_kpi_file,
    validate_template_file,
    workbook_sheet_name,
)
from normalization_memory import (
    list_memory_entries,
    remove_memory_alias,
    remove_memory_sheet_alias,
    save_memory_alias,
    save_memory_sheet_alias,
    source_label_from_note,
    source_sheet_from_note,
)
from normalization_layer import build_normalization_review


PROJECT_ROOT = Path(__file__).resolve().parent
FINAL_OUTPUTS_DIR = PROJECT_ROOT / "Outputs" / "Final Generated Files"


def _month_options() -> list[tuple[str, tuple[int, int]]]:
    current_year = date.today().year
    first_historical_year = 2021
    return [
        (f"{month_abbr[month]} {str(year)[-2:]}", (year, month))
        for year in range(first_historical_year, current_year + 16)
        for month in (3, 6, 9, 12)
    ]


def _default_month_index(
    month_choices: list[tuple[str, tuple[int, int]]],
    default_month: tuple[int, int],
    fallback: int,
) -> int:
    labels = [label for label, month in month_choices if month == default_month]
    if not labels:
        return fallback
    return [label for label, _ in month_choices].index(labels[0])


def _save_upload(uploaded_file, path: Path) -> None:
    path.write_bytes(uploaded_file.getbuffer())


def _uploaded_suffix(uploaded_file) -> str:
    return Path(getattr(uploaded_file, "name", "")).suffix.lower()


def _month_from_filename(filename: str, months: tuple[tuple[int, int], ...]) -> tuple[int, int] | None:
    text = filename.lower().replace("'", "")
    text = re.sub(r"[_\-.]+", " ", text)
    tokens = set(re.findall(r"[a-z]+|\d+", text))
    for year, month in months:
        year_tokens = {str(year), str(year)[-2:]}
        month_tokens = {
            month_abbr[month].lower(),
            datetime(year, month, 1).strftime("%B").lower(),
            f"{month:02d}",
            str(month),
        }
        if tokens.intersection(year_tokens) and tokens.intersection(month_tokens):
            return year, month
    return None


def _month_sheet_title(month_key: tuple[int, int]) -> str:
    year, month = month_key
    return f"{month_abbr[month]}-{str(year)[-2:]}"


def _unique_sheet_title(workbook: Workbook, title: str) -> str:
    base = re.sub(r"[:\\/?*\[\]]", " ", title).strip() or "Sheet"
    base = base[:31]
    if base not in workbook.sheetnames:
        return base

    for index in range(2, 100):
        suffix = f" {index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
    return base[:27] + " 99"


def _copy_sheet_values(source_sheet, target_sheet) -> None:
    for row in source_sheet.iter_rows():
        for source_cell in row:
            target_sheet.cell(row=source_cell.row, column=source_cell.column).value = source_cell.value


def _parse_pdf_number(value: str) -> float | None:
    if not re.fullmatch(r"-?\$?[0-9,]+(?:\.\d+)?", value):
        return None
    return float(value.replace("$", "").replace(",", ""))


def _extract_revolving_games_pdf_rows(pdf_path: Path) -> dict[str, dict[tuple[int, int], float]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("PDF uploads need pdfplumber. Add pdfplumber to requirements.txt and redeploy.") from exc

    month_centers: dict[tuple[int, int], float] = {}
    rows: dict[str, dict[tuple[int, int], float]] = {}

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False)

            for index, word in enumerate(words[:-1]):
                if not re.fullmatch(r"[A-Za-z]{3}", word["text"]):
                    continue
                next_word = words[index + 1]
                if abs(next_word["top"] - word["top"]) > 2:
                    continue
                if not re.fullmatch(r"20\d{2}", next_word["text"]):
                    continue
                try:
                    month = datetime.strptime(word["text"], "%b").month
                    year = int(next_word["text"])
                except ValueError:
                    continue
                month_centers[(year, month)] = (word["x0"] + next_word["x1"]) / 2

            if not month_centers:
                continue

            line_tops = sorted({
                round(word["top"] / 3) * 3
                for word in words
                if 120 <= word["top"] <= 760
            })
            for line_top in line_tops:
                line_words = [
                    word
                    for word in words
                    if round(word["top"] / 3) * 3 == line_top
                ]
                label = " ".join(
                    word["text"]
                    for word in line_words
                    if word["x0"] < 300
                ).strip()
                if not label or label.startswith("Accrual Basis"):
                    continue

                for word in line_words:
                    number = _parse_pdf_number(word["text"])
                    if number is None or word["x0"] < 300:
                        continue
                    x_center = (word["x0"] + word["x1"]) / 2
                    month_key = min(month_centers, key=lambda key: abs(x_center - month_centers[key]))
                    if abs(x_center - month_centers[month_key]) > 45:
                        continue
                    rows.setdefault(label_key(label), {})[month_key] = number

    return rows


def label_key(value: object) -> str:
    return " ".join(str(value).lower().replace("&", "and").split())


def _sum_pdf_rows(
    source_rows: dict[str, dict[tuple[int, int], float]],
    labels: tuple[str, ...],
    month_key: tuple[int, int],
) -> float:
    total = 0.0
    for label in labels:
        total += source_rows.get(label_key(label), {}).get(month_key, 0.0) or 0.0
    return total


def _convert_revolving_games_pdf_to_workbook(
    pdf_path: Path,
    output_path: Path,
    months: tuple[tuple[int, int], ...],
) -> None:
    source_rows = _extract_revolving_games_pdf_rows(pdf_path)

    operations_labels = (
        "Total for 6200 Payroll expenses",
        "6270 Contract labor",
        "6275 International Entity Labor",
        "6420 Business licenses",
        "6475 Dues, Membership and Subscription",
        "6540 Payment Processing Fee",
        "Total for 6790 Utilities",
        "6920 State Tax Expense",
    )
    sga_labels = (
        "6340 Entertainment",
        "6370 Meals",
        "6415 Gifts & Donations",
        "Total for 6390 Travel",
        "Total for 6520 Professional Services",
        "Total for 6740 General business expenses",
    )
    other_labels = (
        "Total for 6400 Office expenses",
        "6810 Other Personal Expenses",
        "Total for 6820 Interest paid",
        "6910 Uncategorized Expense",
        "6940 Depreciation",
        "Total for Other Expenses",
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "P_L_for_Q1_2025"
    sheet["A1"] = "Revolving Games, Inc."
    if months:
        sheet["B1"] = f"Q{((months[-1][1] - 1) // 3) + 1} {months[-1][0]}"

    month_columns = {month_key: index + 2 for index, month_key in enumerate(months)}
    for month_key, col in month_columns.items():
        year, month = month_key
        sheet.cell(row=2, column=col).value = datetime(year, month, 1)

    metric_rows = {
        "Revenue": 3,
        "Operations": 4,
        "SG&A": 5,
        "Marketing": 6,
        "Software and Servers": 7,
        "Other expenses": 8,
        "Total Expenses": 9,
        "Net Burn": 10,
    }
    for label, row in metric_rows.items():
        sheet.cell(row=row, column=1).value = label

    for month_key, col in month_columns.items():
        revenue = source_rows.get(label_key("Total for Income"), {}).get(month_key, 0.0) or 0.0
        operations = _sum_pdf_rows(source_rows, operations_labels, month_key)
        sga = _sum_pdf_rows(source_rows, sga_labels, month_key)
        marketing = source_rows.get(label_key("Total for 6510 Advertising & marketing"), {}).get(month_key, 0.0) or 0.0
        software = source_rows.get(label_key("Computer Equipment"), {}).get(month_key, 0.0) or 0.0
        other = _sum_pdf_rows(source_rows, other_labels, month_key)
        total_expenses = -(operations + sga + marketing + software + other)

        values = {
            "Revenue": revenue,
            "Operations": operations,
            "SG&A": sga,
            "Marketing": marketing,
            "Software and Servers": software,
            "Other expenses": other,
            "Total Expenses": total_expenses,
            "Net Burn": revenue + total_expenses,
        }
        for label, row in metric_rows.items():
            sheet.cell(row=row, column=col).value = round(values[label], 2)

    sheet["A12"] = "*Converted from Revolving Games PDF P&L"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def _save_kpi_uploads(
    profile_name: str,
    uploaded_files: list,
    path: Path,
    months: tuple[tuple[int, int], ...],
) -> None:
    if len(uploaded_files) == 1:
        if profile_name == "revolving_games" and _uploaded_suffix(uploaded_files[0]) == ".pdf":
            with tempfile.TemporaryDirectory() as pdf_tmp:
                pdf_path = Path(pdf_tmp) / "revolving_games.pdf"
                _save_upload(uploaded_files[0], pdf_path)
                _convert_revolving_games_pdf_to_workbook(pdf_path, path, months)
            return
        _save_upload(uploaded_files[0], path)
        return

    combined_workbook = Workbook()
    default_sheet = combined_workbook.active
    combined_workbook.remove(default_sheet)

    with tempfile.TemporaryDirectory() as combine_tmp:
        combine_dir = Path(combine_tmp)
        for file_index, uploaded_file in enumerate(uploaded_files, start=1):
            source_suffix = _uploaded_suffix(uploaded_file)
            source_path = combine_dir / f"source_{file_index}{source_suffix or '.xlsx'}"
            _save_upload(uploaded_file, source_path)
            if profile_name == "revolving_games" and source_suffix == ".pdf":
                converted_path = combine_dir / f"source_{file_index}_converted.xlsx"
                _convert_revolving_games_pdf_to_workbook(source_path, converted_path, months)
                source_path = converted_path
            source_workbook = load_workbook(source_path, data_only=True)
            try:
                file_month = _month_from_filename(uploaded_file.name, months)
                for source_sheet in source_workbook.worksheets:
                    if profile_name == "roomy" and file_month is not None:
                        title = _month_sheet_title(file_month)
                    else:
                        title = source_sheet.title
                    target_sheet = combined_workbook.create_sheet(
                        _unique_sheet_title(combined_workbook, title)
                    )
                    _copy_sheet_values(source_sheet, target_sheet)
            finally:
                source_workbook.close()

    path.parent.mkdir(parents=True, exist_ok=True)
    combined_workbook.save(path)
    combined_workbook.close()


def _quarter_months(ending_month: tuple[int, int]) -> tuple[tuple[int, int], ...]:
    year, month = ending_month
    quarter_start_month = ((month - 1) // 3) * 3 + 1
    return tuple((year, quarter_start_month + offset) for offset in range(3))


def _output_name(months: tuple[tuple[int, int], ...]) -> str:
    if not months:
        return "Monthly Reporting.xlsx"

    year, month = months[-1]
    return f"Monthly Reporting {month_abbr[month]} {str(year)[-2:]}.xlsx"


def _summary_name(months: tuple[tuple[int, int], ...]) -> str:
    return f"{Path(_output_name(months)).stem} - run summary.json"


def _source_trace_name(months: tuple[tuple[int, int], ...]) -> str:
    return f"{Path(_output_name(months)).stem} - source trace.csv"


def _latest_output_rows(limit: int = 8) -> list[dict[str, str]]:
    if not FINAL_OUTPUTS_DIR.exists():
        return []

    files = sorted(
        FINAL_OUTPUTS_DIR.glob("Monthly Reporting *.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    rows = []
    for path in files[:limit]:
        rows.append(
            {
                "File": path.name,
                "Modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%b %d, %Y %H:%M"),
                "Folder": str(path.parent),
            }
        )
    return rows


def _status_icon(status: str) -> str:
    if status == "Ready":
        return "Ready"
    if status == "Warning":
        return "Warning"
    if status == "Missing":
        return "Missing"
    if status == "Not uploaded":
        return "Not uploaded"
    return "Blocked"


def _validation_details(issues: tuple[str, ...], warnings: tuple[str, ...]) -> str:
    messages = [*issues, *warnings]
    return " | ".join(messages) if messages else ""


def _rows_to_csv(rows: list[dict[str, object]]) -> bytes:
    if not rows:
        return b""

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()))
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


QOQ_METHOD_MANUAL = "Manual workbook rules"
QOQ_RULE_QUARTER_TOTAL = "quarter_total"
QOQ_RULE_ENDING_VALUE = "ending_value"


def _manual_qoq_rule(
    method: str,
    source_rows: tuple[int, ...] | None = None,
    note: str | None = None,
) -> dict[str, object]:
    return {
        "method": method,
        "source_rows": source_rows,
        "note": note,
    }


QOQ_MANUAL_RULES: dict[str, dict[int, dict[str, object]]] = {
    "abhi": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(14, 29)},
        30: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE),
    },
    "simpaisa": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(45, 69)},
    },
    "bykea": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(16, 24)},
        **{row: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE) for row in range(24, 29)},
    },
    "dot_and_line": {
        15: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL),
        16: _manual_qoq_rule(
            QOQ_RULE_QUARTER_TOTAL,
            source_rows=(16, 17),
            note="Dec 25 manual formula combines Net revenue and Other revenue.",
        ),
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(17, 26)},
        27: _manual_qoq_rule(
            QOQ_RULE_ENDING_VALUE,
            source_rows=(26,),
            note="Dec 25 manual formula displays row 27 but compares row 26.",
        ),
    },
    "procheck": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(16, 29)},
    },
    "roomy": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(13, 25)},
        **{row: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE) for row in range(25, 29)},
    },
    "oladoc": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(16, 38)},
        38: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE),
        39: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE),
    },
    "tapmad": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(46, 57)},
        57: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE),
        58: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE),
    },
    "jiye_technologies": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(13, 30)},
    },
    "oneload": {
        **{row: _manual_qoq_rule(QOQ_RULE_QUARTER_TOTAL) for row in range(61, 78)},
        76: _manual_qoq_rule(QOQ_RULE_ENDING_VALUE),
    },
}


def _previous_quarter_months(months: tuple[tuple[int, int], ...]) -> tuple[tuple[int, int], ...]:
    if not months:
        return ()

    first_year, first_month = months[0]
    previous_end_month = first_month - 1
    previous_end_year = first_year
    if previous_end_month == 0:
        previous_end_month = 12
        previous_end_year -= 1
    return _quarter_months((previous_end_year, previous_end_month))


def _quarter_label(months: tuple[tuple[int, int], ...]) -> str:
    return ", ".join(f"{month_abbr[month]} {str(year)[-2:]}" for year, month in months)


def _number_value(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None

    text = value.strip()
    if not text or text.lower() in {"nan", "-", "n/a", "na"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    is_percent = "%" in text
    text = (
        text.replace("$", "")
        .replace("PKR", "")
        .replace("USD", "")
        .replace(",", "")
        .replace("%", "")
        .strip()
    )
    try:
        number = float(text)
    except ValueError:
        return None
    if negative:
        number *= -1
    if is_percent:
        number /= 100
    return number


def _row_label(sheet, row: int) -> str:
    labels = []
    for col in range(1, min(8, sheet.max_column) + 1):
        value = sheet.cell(row=row, column=col).value
        if isinstance(value, str):
            text = " ".join(value.split())
            if text:
                labels.append(text)
    if not labels:
        return ""
    return labels[-1]


def _auto_qoq_method(metric_label: str) -> str:
    label = metric_label.lower()
    if any(token in label for token in ("%", "margin", "rate", "ratio", "yield", "mdr", "take rate")):
        return "Quarter average"
    if any(token in label for token in ("/ day", "per day", "/day", "daily", "average", "avg")):
        return "Quarter average"
    if any(token in label for token in ("monthly active", "mau", "mad", "active users", "active drivers")):
        return "Quarter ending value"
    if "transaction" in label or "revenue" in label or "gmv" in label or "gtv" in label or "nmv" in label:
        return "Quarter total"
    ending_tokens = (
        "cash",
        "balance",
        "runway",
        "assets",
        "liabilities",
        "equity",
        "portfolio",
        "loan book",
        "deposits",
        "headcount",
        "users",
        "clients",
        "merchants",
        "captains",
        "corporates",
        "stores",
    )
    if any(token in label for token in ending_tokens):
        return "Quarter ending value"
    return "Quarter total"


def _aggregate_qoq_values(values: list[float], method: str) -> float | None:
    clean_values = [value for value in values if value is not None]
    if not clean_values:
        return None
    if method == "Quarter average":
        return sum(clean_values) / len(clean_values)
    if method == "Quarter ending value":
        return clean_values[-1]
    return sum(clean_values)


def _qoq_percent(current_total: float, previous_total: float) -> float | None:
    if previous_total == 0:
        return None
    return (current_total / previous_total - 1) * 100


def _qoq_values_for_rows(
    sheet,
    source_rows: tuple[int, ...],
    month_columns: dict[tuple[int, int], int],
    months: tuple[tuple[int, int], ...],
) -> list[float]:
    values: list[float] = []
    for source_row in source_rows:
        for month in months:
            column = month_columns.get(month)
            if column is None:
                continue
            value = _number_value(sheet.cell(row=source_row, column=column).value)
            if value is not None:
                values.append(value)
    return values


def _qoq_formula_description(rule_method: str) -> str:
    if rule_method == QOQ_RULE_ENDING_VALUE:
        return "Current quarter-end / previous quarter-end - 1"
    return "Current quarter total / previous quarter total - 1"


def _manual_qoq_rows_for_profile(
    sheet,
    profile_name: str,
    current_cols: dict[tuple[int, int], int],
    previous_cols: dict[tuple[int, int], int],
    months: tuple[tuple[int, int], ...],
    previous_months: tuple[tuple[int, int], ...],
) -> list[dict[str, object]]:
    profile = PROFILES[profile_name]
    profile_rules = QOQ_MANUAL_RULES.get(profile_name, {})
    rows: list[dict[str, object]] = []

    for row_index, rule in profile_rules.items():
        metric = _row_label(sheet, row_index)
        if not metric or metric.lower() in {"actual", "forecast", "source: company information", "kpis"}:
            continue

        source_rows = rule.get("source_rows") or (row_index,)
        source_rows = tuple(int(row) for row in source_rows)
        rule_method = str(rule["method"])
        formula = _qoq_formula_description(rule_method)

        if rule_method == QOQ_RULE_ENDING_VALUE:
            current_months = months[-1:]
            previous_months_for_calc = previous_months[-1:]
            method_label = "Manual: quarter-end month"
        else:
            current_months = months
            previous_months_for_calc = previous_months
            method_label = "Manual: quarter total"

        current_values = _qoq_values_for_rows(sheet, source_rows, current_cols, current_months)
        previous_values = _qoq_values_for_rows(sheet, source_rows, previous_cols, previous_months_for_calc)
        if not current_values and not previous_values:
            continue

        current_total = sum(current_values) if current_values else None
        previous_total = sum(previous_values) if previous_values else None
        if current_total is None or previous_total is None:
            qoq_pct = None
            change = None
        else:
            change = current_total - previous_total
            qoq_pct = _qoq_percent(current_total, previous_total)

        rows.append(
            {
                "Company": profile.display_name,
                "Row": row_index,
                "Source rows": ", ".join(str(row) for row in source_rows),
                "Metric": metric,
                "Method": method_label,
                "Formula": formula,
                "Current quarter": current_total,
                "Previous quarter": previous_total,
                "Change": change,
                "QoQ %": qoq_pct,
            }
        )

    return rows


def _qoq_rows_for_workbook(
    workbook_path: Path,
    months: tuple[tuple[int, int], ...],
    selected_profile_names: list[str],
    calculation_method: str,
) -> tuple[list[dict[str, object]], list[str]]:
    previous_months = _previous_quarter_months(months)
    warnings: list[str] = []
    rows: list[dict[str, object]] = []

    workbook = load_workbook(workbook_path, data_only=True)
    try:
        for profile_name in selected_profile_names:
            profile = PROFILES[profile_name]
            actual_sheet_name = workbook_sheet_name(workbook, profile.report_sheet)
            if actual_sheet_name is None:
                warnings.append(f"{profile.display_name}: sheet not found.")
                continue

            sheet = workbook[actual_sheet_name]
            current_cols = find_month_columns(sheet, months, header_search_rows=30)
            previous_cols = find_month_columns(sheet, previous_months, header_search_rows=30)
            missing_current = [month for month in months if month not in current_cols]
            missing_previous = [month for month in previous_months if month not in previous_cols]
            if missing_current:
                warnings.append(f"{profile.display_name}: current quarter columns missing for {_quarter_label(tuple(missing_current))}.")
            if missing_previous:
                warnings.append(f"{profile.display_name}: previous quarter columns missing for {_quarter_label(tuple(missing_previous))}.")
            if len(current_cols) == 0 or len(previous_cols) == 0:
                continue

            if calculation_method == QOQ_METHOD_MANUAL:
                if profile_name not in QOQ_MANUAL_RULES:
                    warnings.append(
                        f"{profile.display_name}: no manual QoQ formula rules were found in the Dec 25 sample workbook."
                    )
                    continue

                rows.extend(
                    _manual_qoq_rows_for_profile(
                        sheet=sheet,
                        profile_name=profile_name,
                        current_cols=current_cols,
                        previous_cols=previous_cols,
                        months=months,
                        previous_months=previous_months,
                    )
                )
                continue

            for row_index in range(1, sheet.max_row + 1):
                metric = _row_label(sheet, row_index)
                if not metric or metric.lower() in {"actual", "forecast", "source: company information"}:
                    continue

                current_values = [
                    _number_value(sheet.cell(row=row_index, column=current_cols[month]).value)
                    for month in months
                    if month in current_cols
                ]
                previous_values = [
                    _number_value(sheet.cell(row=row_index, column=previous_cols[month]).value)
                    for month in previous_months
                    if month in previous_cols
                ]
                if not any(value is not None for value in current_values):
                    continue
                if not any(value is not None for value in previous_values):
                    continue

                row_method = _auto_qoq_method(metric) if calculation_method == "Auto" else calculation_method
                current_total = _aggregate_qoq_values(current_values, row_method)
                previous_total = _aggregate_qoq_values(previous_values, row_method)
                if current_total is None or previous_total is None:
                    continue

                change = current_total - previous_total
                qoq_pct = _qoq_percent(current_total, previous_total)

                rows.append(
                    {
                        "Company": profile.display_name,
                        "Row": row_index,
                        "Source rows": str(row_index),
                        "Metric": metric,
                        "Method": row_method,
                        "Formula": "Selected method current / previous - 1",
                        "Current quarter": current_total,
                        "Previous quarter": previous_total,
                        "Change": change,
                        "QoQ %": qoq_pct,
                    }
                )
    finally:
        workbook.close()

    return rows, warnings


def _render_qoq_calculator(selected_months: tuple[tuple[int, int], ...]) -> None:
    with st.container(border=True):
        st.subheader("QoQ changes")
        st.caption("Upload a completed Monthly Reporting workbook to calculate company-by-company quarter-over-quarter changes.")

        qoq_file = st.file_uploader(
            "Completed Monthly Reporting workbook",
            type=["xlsx"],
            key="qoq_completed_workbook",
        )

        qoq_cols = st.columns([1, 2])
        with qoq_cols[0]:
            calculation_method = st.selectbox(
                "Calculation method",
                options=[
                    QOQ_METHOD_MANUAL,
                    "Auto",
                    "Quarter total",
                    "Quarter average",
                    "Quarter ending value",
                ],
                index=0,
                help="Manual workbook rules uses the row-by-row QoQ formulas from the Dec 25 manually prepared report. Auto remains available for exploratory checks.",
            )
        with qoq_cols[1]:
            profile_options = list(PROFILES.keys())
            selected_profile_names = st.multiselect(
                "Companies",
                options=profile_options,
                default=profile_options,
                format_func=lambda profile_name: PROFILES[profile_name].display_name,
            )

        current_label = _quarter_label(selected_months)
        previous_label = _quarter_label(_previous_quarter_months(selected_months))
        st.caption(f"Comparing {current_label} against {previous_label}.")

        calculate_qoq = st.button(
            "Calculate QoQ changes",
            disabled=qoq_file is None or not selected_profile_names,
            width="stretch",
        )
        if not calculate_qoq:
            return

        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "completed_monthly_reporting.xlsx"
            _save_upload(qoq_file, workbook_path)
            with st.spinner("Calculating QoQ changes"):
                qoq_rows, warnings = _qoq_rows_for_workbook(
                    workbook_path=workbook_path,
                    months=selected_months,
                    selected_profile_names=selected_profile_names,
                    calculation_method=calculation_method,
                )

        if warnings:
            with st.expander("QoQ warnings", expanded=False):
                for warning in warnings:
                    st.warning(warning)

        if not qoq_rows:
            st.info("No comparable QoQ rows were found for the selected companies and quarter.")
            return

        st.success(f"Calculated {len(qoq_rows):,} QoQ rows.")
        st.dataframe(
            qoq_rows,
            hide_index=True,
            width="stretch",
            column_config={
                "Current quarter": st.column_config.NumberColumn(format="%.2f"),
                "Previous quarter": st.column_config.NumberColumn(format="%.2f"),
                "Change": st.column_config.NumberColumn(format="%.2f"),
                "QoQ %": st.column_config.NumberColumn(format="%.1f%%"),
            },
        )
        st.download_button(
            "Download QoQ changes CSV",
            data=_rows_to_csv(qoq_rows),
            file_name=f"QoQ Changes {month_abbr[selected_months[-1][1]]} {str(selected_months[-1][0])[-2:]}.csv",
            mime="text/csv",
            width="stretch",
        )


def _render_latest_outputs() -> None:
    rows = _latest_output_rows()
    with st.expander("Latest generated files", expanded=False):
        if rows:
            st.dataframe(rows, hide_index=True, width="stretch")
        else:
            st.caption("No generated files yet.")


def render_monthly_reporting_page() -> None:
    st.header("Monthly Reporting")
    st.caption("Upload the reporting template and whichever company KPI files have arrived. Missing companies or missing months are left blank.")

    first_historical_year = 2021
    year_choices = list(range(first_historical_year, date.today().year + 16))
    quarter_end_options = [("Mar", 3), ("Jun", 6), ("Sep", 9), ("Dec", 12)]

    with st.container(border=True):
        st.subheader("Run setup")
        source_mode = "Upload files"
        st.caption("Input source: file upload. Google Drive integration is parked for now.")

        setup_cols = st.columns([1, 1, 2, 2])
        with setup_cols[0]:
            selected_year = st.selectbox(
                "Year",
                options=year_choices,
                index=year_choices.index(2026) if 2026 in year_choices else len(year_choices) - 1,
            )
        with setup_cols[1]:
            selected_quarter_label = st.selectbox(
                "Quarter ending",
                options=[label for label, _month in quarter_end_options],
                index=0,
            )
        quarter_month_lookup = dict(quarter_end_options)
        selected_months = _quarter_months((selected_year, quarter_month_lookup[selected_quarter_label]))
        selected_month_labels = ", ".join(
            f"{month_abbr[month]} {str(year)[-2:]}" for year, month in selected_months
        )
        with setup_cols[2]:
            st.metric("Months populated", selected_month_labels)
        with setup_cols[3]:
            st.metric("Output file", _output_name(selected_months))

        if source_mode == "Google Drive folder":
            st.caption("Use this when the shared Drive is online-only and not synced to Finder.")
            folder_link = st.text_input(
                "Shared Drive folder link or ID",
                placeholder="https://drive.google.com/drive/folders/...",
            )
            credentials_file = st.file_uploader(
                "Google service account credentials",
                type=["json"],
                key="google_drive_service_account",
            )

            available, missing_package = drive_api_available()
            if not available:
                st.warning(
                    "Google Drive connection packages still need to be installed "
                    f"before live Drive access can run. Missing package: {missing_package}."
                )

            if credentials_file is not None:
                try:
                    email = service_account_email(credentials_file.getvalue())
                    st.info(f"Share the online Drive folder with this service account: {email}")
                except ValueError as exc:
                    st.error(str(exc))

            check_drive = st.button(
                "Check Drive access",
                disabled=not folder_link or credentials_file is None or not available,
                width="stretch",
            )
            if check_drive:
                try:
                    folder_id = parse_drive_folder_id(folder_link)
                    service = build_drive_service(credentials_file.getvalue())
                    items = list_folder_items(service, folder_id)
                except Exception as exc:
                    st.error(f"Could not access that Drive folder yet: {exc}")
                else:
                    st.success(f"Connected. Found {len(items)} items in the folder.")
                    st.dataframe(
                        [
                            {
                                "Name": item.name,
                                "Type": item.mime_type,
                                "Modified": item.modified_time or "",
                            }
                            for item in items
                        ],
                        hide_index=True,
                        width="stretch",
                    )

            st.info("Once the Drive folder is connected, the next step is auto-detecting the template and company KPI files from this folder, then saving the generated report back online.")
            _render_latest_outputs()
            return

    _render_qoq_calculator(selected_months)

    with st.container(border=True):
        st.subheader("Template")
        template_file = st.file_uploader(
            "Monthly Reporting Template",
            type=["xlsx"],
            key="monthly_reporting_template",
        )

    with st.container(border=True):
        st.subheader("KPI files")
        st.caption("Upload only the files received so far. You can attach multiple files for a company when monthly KPI packs are separate.")
        uploaded_kpis = {}
        upload_cols = st.columns(2)
        for index, (profile_name, profile) in enumerate(PROFILES.items()):
            with upload_cols[index % 2]:
                accepted_types = ["xlsx", "pdf"] if profile_name == "revolving_games" else ["xlsx"]
                uploaded = st.file_uploader(
                    profile.display_name,
                    type=accepted_types,
                    key=f"kpi_{profile_name}",
                    accept_multiple_files=True,
                )
            if uploaded:
                uploaded_kpis[profile_name] = uploaded

    with st.expander("Mapping status", expanded=False):
        mapped_companies = {profile.display_name for profile in PROFILES.values()}
        pending_companies = [
            company for company in PENDING_COMPANIES
            if company not in mapped_companies
        ]
        status_rows = [
            {"Company": profile.display_name, "Status": "Mapped"}
            for profile in PROFILES.values()
        ]
        if pending_companies:
            status_rows.extend(
                {"Company": company, "Status": "Mapping pending"}
                for company in pending_companies
            )
        status_rows.extend(
            {"Company": company, "Status": "Inactive - left untouched"}
            for company in INACTIVE_COMPANIES
        )
        st.dataframe(status_rows, hide_index=True, width="stretch")

    validation_rows = []
    validation_blocked = False
    omitted_companies = []
    intelligence_summary = {
        "findings": [],
        "status_counts": {"Blocked": 0, "Warning": 0, "Review": 0},
        "blocked": False,
    }
    normalization_summary = {
        "rows": [],
        "status_counts": {"Auto-ready": 0, "Review": 0, "Missing": 0},
        "blocked": False,
        "model_status": "Upload a template and at least one KPI file to run normalization checks.",
    }
    with tempfile.TemporaryDirectory() as validation_tmp:
        validation_dir = Path(validation_tmp)
        template_validation_path = None
        kpi_validation_paths = {}
        if template_file is not None:
            template_validation_path = validation_dir / "monthly_reporting_template.xlsx"
            _save_upload(template_file, template_validation_path)
            template_validation = validate_template_file(template_validation_path, selected_months)
            validation_blocked = validation_blocked or template_validation.status == "Blocked"
            validation_rows.append(
                {
                    "Company": template_validation.company,
                    "Status": _status_icon(template_validation.status),
                    "Confidence": f"{template_validation.confidence}%",
                    "Details": _validation_details(template_validation.issues, template_validation.warnings),
                }
            )
        else:
            validation_rows.append(
                {
                    "Company": "Monthly Reporting Template",
                    "Status": "Missing",
                    "Confidence": "0%",
                    "Details": "Upload the monthly reporting template.",
                }
            )

        for profile_name, profile in PROFILES.items():
            uploaded = uploaded_kpis.get(profile_name)
            if uploaded is None:
                omitted_companies.append(profile.display_name)
                validation_rows.append(
                    {
                        "Company": profile.display_name,
                        "Status": "Not uploaded",
                        "Confidence": "",
                        "Details": "Will be left blank/untouched.",
                    }
                )
                continue

            kpi_validation_path = validation_dir / f"{profile_name}_kpi.xlsx"
            _save_kpi_uploads(profile_name, uploaded, kpi_validation_path, selected_months)
            kpi_validation_paths[profile_name] = kpi_validation_path
            validation = validate_kpi_file(profile_name, kpi_validation_path, selected_months)
            validation_blocked = validation_blocked or validation.status == "Blocked"
            details = _validation_details(validation.issues, validation.warnings)
            if len(uploaded) > 1 and not details:
                details = f"{len(uploaded)} files combined for this company."
            validation_rows.append(
                {
                    "Company": validation.company,
                    "Status": _status_icon(validation.status),
                    "Confidence": f"{validation.confidence}%",
                    "Details": details,
                }
            )

        if template_validation_path is not None and kpi_validation_paths:
            intelligence_summary = run_intelligence_checks(
                template_path=template_validation_path,
                kpi_files=kpi_validation_paths,
                months=selected_months,
            )
            normalization_summary = build_normalization_review(
                template_path=template_validation_path,
                kpi_files=kpi_validation_paths,
                months=selected_months,
            )
            validation_blocked = validation_blocked or bool(intelligence_summary.get("blocked"))

    with st.container(border=True):
        st.subheader("Readiness checks")
        status_counts = {
            "Ready": sum(row["Status"] == "Ready" for row in validation_rows),
            "Warning": sum(row["Status"] == "Warning" for row in validation_rows),
            "Blocked": sum(row["Status"] == "Blocked" for row in validation_rows),
            "Not uploaded": sum(row["Status"] == "Not uploaded" for row in validation_rows),
        }
        metric_cols = st.columns(4)
        metric_cols[0].metric("Ready", status_counts["Ready"])
        metric_cols[1].metric("Warnings", status_counts["Warning"])
        metric_cols[2].metric("Blocked", status_counts["Blocked"])
        metric_cols[3].metric("Left blank", status_counts["Not uploaded"])

        st.dataframe(validation_rows, hide_index=True, width="stretch")
        if status_counts["Not uploaded"]:
            st.caption("Not uploaded companies will not be changed in this run.")

    with st.container(border=True):
        st.subheader("Intelligence detector")
        counts = intelligence_summary["status_counts"]
        intelligence_cols = st.columns(3)
        intelligence_cols[0].metric("Blocked", counts["Blocked"])
        intelligence_cols[1].metric("Warnings", counts["Warning"])
        intelligence_cols[2].metric("Review", counts["Review"])

        findings = intelligence_summary["findings"]
        if findings:
            st.dataframe(
                [
                    {
                        "Company": finding["company"],
                        "Severity": finding["severity"],
                        "Check": finding["check"],
                        "Finding": finding["finding"],
                        "Suggested fix": finding.get("suggested_fix", ""),
                        "Confidence": f"{finding.get('confidence', 0)}%",
                        "Approval needed": "Yes" if finding.get("approval_required") else "No",
                        "Action": finding["action"],
                        "Location": finding["location"],
                    }
                    for finding in findings
                ],
                hide_index=True,
                width="stretch",
            )
        elif template_file is not None and uploaded_kpis:
            st.success("No intelligence warnings detected for the uploaded files.")
        else:
            st.caption("Upload a template and at least one KPI file to run intelligence checks.")

    with st.container(border=True):
        st.subheader("Normalization review")
        st.caption("This converts uploaded KPI formats into the standard reporting structure before the template is populated. The workbook layout is not changed.")
        normalization_counts = normalization_summary["status_counts"]
        normalization_cols = st.columns(3)
        normalization_cols[0].metric("Auto-ready", normalization_counts["Auto-ready"])
        normalization_cols[1].metric("Review", normalization_counts["Review"])
        normalization_cols[2].metric("Missing", normalization_counts["Missing"])
        st.caption(normalization_summary["model_status"])

        normalization_rows = normalization_summary["rows"]
        if normalization_rows:
            show_all_normalization = st.checkbox(
                "Show auto-ready sample rows",
                value=False,
                key="show_all_normalization_rows",
            )
            display_rows = normalization_rows if show_all_normalization else [
                row for row in normalization_rows if row["status"] != "Auto-ready"
            ]
            if display_rows:
                st.dataframe(display_rows, hide_index=True, width="stretch")
            else:
                st.success("No normalization review items. Source formats match expected mappings.")

            review_rows = [
                row for row in normalization_rows
                if row["status"] == "Review"
                and (
                    source_label_from_note(str(row.get("note", "")))
                    or source_sheet_from_note(str(row.get("note", "")))
                )
            ]
            if review_rows:
                st.caption("Approve a suggested mapping only after checking the source cell or workbook tab. Approved mappings are remembered for future runs.")
                profile_by_company = {
                    profile.display_name: profile_name
                    for profile_name, profile in PROFILES.items()
                }
                approval_candidates = {}
                for row in review_rows:
                    source_label = source_label_from_note(str(row["note"]))
                    source_sheet = source_sheet_from_note(str(row["note"]))
                    if source_label:
                        approval_type = "Source row"
                        expected = row["Metric"] if "Metric" in row else row["metric"]
                        suggested = source_label
                    elif source_sheet:
                        approval_type = "Workbook tab"
                        expected, suggested = source_sheet
                    else:
                        continue
                    key = (approval_type, row["company"], expected, suggested)
                    approval_candidates.setdefault(
                        key,
                        {
                            "Approve": False,
                            "Type": approval_type,
                            "Company": row["company"],
                            "Expected": expected,
                            "Suggested": suggested,
                            "Source": row["source"],
                            "Confidence": row["confidence"],
                        },
                    )

                approval_rows = list(approval_candidates.values())
                edited_approvals = st.data_editor(
                    approval_rows,
                    hide_index=True,
                    width="stretch",
                    disabled=["Type", "Company", "Expected", "Suggested", "Source", "Confidence"],
                    key="normalization_mapping_approvals",
                )
                if st.button("Save approved mappings", width="stretch"):
                    saved = 0
                    for row in edited_approvals:
                        if not row.get("Approve"):
                            continue
                        profile_name = profile_by_company.get(str(row["Company"]))
                        if profile_name is None:
                            continue
                        if row["Type"] == "Workbook tab":
                            if save_memory_sheet_alias(
                                profile_name=profile_name,
                                expected_sheet_name=str(row["Expected"]),
                                actual_sheet_name=str(row["Suggested"]),
                                confidence=int(row["Confidence"]),
                            ):
                                saved += 1
                        else:
                            if save_memory_alias(
                                profile_name=profile_name,
                                report_label=str(row["Expected"]),
                                source_label=str(row["Suggested"]),
                                source_ref=str(row["Source"]),
                                confidence=int(row["Confidence"]),
                            ):
                                saved += 1
                    if saved:
                        st.success(f"Saved {saved} approved mapping{'s' if saved != 1 else ''}. Rechecking with the new memory.")
                        st.rerun()
                    else:
                        st.warning("No mappings were selected for approval.")

            st.download_button(
                "Download normalization review",
                data=_rows_to_csv(normalization_rows),
                file_name=f"{Path(_output_name(selected_months)).stem} - normalization review.csv",
                mime="text/csv",
                width="stretch",
            )
        else:
            st.caption("Upload a template and at least one KPI file to run normalization checks.")

    with st.expander("Approved mapping memory", expanded=False):
        st.caption("Mappings approved in earlier runs. Remove one if it was approved by mistake; the app will go back to asking for review.")
        memory_entries = list_memory_entries()
        if memory_entries:
            memory_rows = [{"Remove": False, **entry} for entry in memory_entries]
            edited_memory = st.data_editor(
                memory_rows,
                hide_index=True,
                width="stretch",
                disabled=["Type", "Profile", "Expected", "Mapped to", "Approved"],
                key="approved_memory_editor",
            )
            if st.button("Remove selected mappings", width="stretch"):
                removed = 0
                for row in edited_memory:
                    if not row.get("Remove"):
                        continue
                    if row["Type"] == "Workbook tab":
                        if remove_memory_sheet_alias(str(row["Profile"]), str(row["Expected"])):
                            removed += 1
                    else:
                        if remove_memory_alias(str(row["Profile"]), str(row["Expected"])):
                            removed += 1
                if removed:
                    st.success(f"Removed {removed} mapping{'s' if removed != 1 else ''}. Rechecking without them.")
                    st.rerun()
                else:
                    st.warning("No mappings were selected for removal.")
        else:
            st.caption("No approved mappings saved yet.")

    if validation_blocked:
        st.error("One or more uploaded files are blocked. Fix the issue before generating the report.")
    elif any(row["Status"] == "Warning" for row in validation_rows):
        st.warning("Some files have warnings. You can generate, but review the readiness details first.")
    if uploaded_kpis:
        st.info(f"This run will update {len(uploaded_kpis)} companies and leave {len(omitted_companies)} companies blank/untouched.")

    generate = st.button(
        "Generate monthly report",
        type="primary",
        disabled=template_file is None or not uploaded_kpis or not selected_months or validation_blocked,
        width="stretch",
    )

    if not generate:
        return

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        template_path = tmp_dir / "monthly_reporting_template.xlsx"
        output_path = tmp_dir / _output_name(selected_months)
        summary_path = tmp_dir / "summary.json"

        _save_upload(template_file, template_path)

        kpi_paths = {}
        for profile_name, uploaded in uploaded_kpis.items():
            kpi_path = tmp_dir / f"{profile_name}_kpi.xlsx"
            _save_kpi_uploads(profile_name, uploaded, kpi_path, selected_months)
            kpi_paths[profile_name] = kpi_path

        with st.spinner("Generating monthly report"):
            intelligence_summary = run_intelligence_checks(
                template_path=template_path,
                kpi_files=kpi_paths,
                months=selected_months,
            )
            source_trace_summary = build_normalization_review(
                template_path=template_path,
                kpi_files=kpi_paths,
                months=selected_months,
            )
            summary = run_batch_update(
                template_path=template_path,
                kpi_files=kpi_paths,
                output_path=output_path,
                months=selected_months,
            )
            summary["omitted_companies"] = [
                profile.display_name
                for profile_name, profile in PROFILES.items()
                if profile_name not in kpi_paths
            ]
            summary["intelligence"] = intelligence_summary
            summary["source_trace"] = source_trace_summary
            summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
            source_trace_bytes = _rows_to_csv(source_trace_summary.get("rows", []))

        output_bytes = output_path.read_bytes()
        summary_bytes = summary_path.read_bytes()
        permanent_output_path = FINAL_OUTPUTS_DIR / _output_name(selected_months)
        permanent_summary_path = FINAL_OUTPUTS_DIR / _summary_name(selected_months)
        permanent_source_trace_path = FINAL_OUTPUTS_DIR / _source_trace_name(selected_months)
        FINAL_OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
        permanent_output_path.write_bytes(output_bytes)
        permanent_summary_path.write_bytes(summary_bytes)
        permanent_source_trace_path.write_bytes(source_trace_bytes)

    with st.container(border=True):
        st.success("Generated! Your Monthly Reporting file is ready.")
        st.caption(f"Saved final file: {permanent_output_path}")
        st.caption(f"Saved run summary: {permanent_summary_path}")
        st.caption(f"Saved source trace: {permanent_source_trace_path}")

        col_a, col_b, col_c = st.columns(3)
        col_a.metric("Values populated", f"{summary['total_values_written']:,}")
        col_b.metric("Rows matched", f"{summary['total_rows_matched']:,}")
        col_c.metric("Companies updated", len(summary["companies"]))

        formula_memory = summary.get("formula_memory", {})
        formula_audit = summary.get("formula_audit", {})
        formula_cols = st.columns(4)
        formula_cols[0].metric("Memory formulas filled", f"{formula_memory.get('formulas_written', 0):,}")
        formula_cols[1].metric("Static values restored", f"{formula_memory.get('static_values_overwritten', 0):,}")
        formula_cols[2].metric("Missing formula sources", f"{formula_memory.get('missing_formula_sources', 0):,}")
        formula_cols[3].metric("Formula audit issues", f"{formula_audit.get('issue_count', 0):,}")
        if formula_audit.get("issues"):
            with st.expander("Formula audit sample", expanded=False):
                st.dataframe(formula_audit["issues"], hide_index=True, width="stretch")

        download_cols = st.columns(3)
        download_cols[0].download_button(
            "Download Monthly Reporting file",
            data=output_bytes,
            file_name=_output_name(selected_months),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
        )
        download_cols[1].download_button(
            "Download run summary",
            data=summary_bytes,
            file_name=_summary_name(selected_months),
            mime="application/json",
            width="stretch",
        )
        download_cols[2].download_button(
            "Download source trace",
            data=source_trace_bytes,
            file_name=_source_trace_name(selected_months),
            mime="text/csv",
            width="stretch",
        )

    company_rows = []
    completed_companies = set()
    for company in summary["companies"]:
        company_name = company.get("company", company.get("profile"))
        completed_companies.add(company_name)
        company_rows.append(
            {
                "Company": company_name,
                "Values": company.get("total_values_written", 0),
                "Rows": company.get("total_rows_matched", 0),
                "Status": "Skipped" if company.get("skipped") else "Complete",
            }
        )
    for company in summary.get("omitted_companies", []):
        if company not in completed_companies:
            company_rows.append(
                {
                    "Company": company,
                    "Values": 0,
                    "Rows": 0,
                    "Status": "Not uploaded - left blank",
                }
            )
    st.dataframe(company_rows, hide_index=True, width="stretch")
    _render_latest_outputs()
