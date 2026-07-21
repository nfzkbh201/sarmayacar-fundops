from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
import re
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from monthly_reporting_automation import (
    CompanyProfile,
    LabelCopyProfile,
    PROFILES,
    SourceSpec,
    abhi_label_key,
    abhi_report_label,
    find_flexible_month_columns,
    find_workbook_month_sheet,
    first_report_label,
    get_bykea_visible_summary_source,
    get_bykea_visible_summary_sources,
    infer_profile_source_sheet_name,
    label_key,
    profile_workbook_sheet,
    profile_workbook_sheet_name,
    resolve_abhi_report_sections,
    workbook_sheet_name,
)
from normalization_memory import (
    memory_aliases_for_profile,
    memory_sheet_aliases_for_profile,
    normalise_memory_key,
)


@dataclass(frozen=True)
class NormalizationRow:
    company: str
    month: str
    template_cell: str
    metric: str
    status: str
    confidence: int
    source: str
    raw_value: str
    normalized_value: str
    conversion: str
    note: str

    def row(self) -> dict[str, object]:
        return asdict(self)


def _month_label(month_key: tuple[int, int]) -> str:
    year, month = month_key
    return datetime(year, month, 1).strftime("%b %y")


def _display(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    return str(value)


def _status_counts(rows: list[NormalizationRow]) -> dict[str, int]:
    return {
        "Auto-ready": sum(row.status == "Auto-ready" for row in rows),
        "Review": sum(row.status == "Review" for row in rows),
        "Missing": sum(row.status == "Missing" for row in rows),
    }


def _conversion_note(
    sign: float = 1,
    rounds: bool = False,
    zero_as_dash: bool = False,
) -> str:
    notes: list[str] = []
    if sign == -1:
        notes.append("flip sign")
    elif sign not in (0, 1):
        notes.append(f"multiply by {sign:g}")
    if rounds:
        notes.append("round")
    if zero_as_dash:
        notes.append("zero to dash")
    return ", ".join(notes) or "direct"


def _should_skip_report_label(label: str | None) -> bool:
    if not label:
        return True
    key = label_key(label) or ""
    if len(label) > 90:
        return True
    skipped_exact = {
        "actual",
        "forecast",
        "company description",
        "key updates",
        "key financial & operating metrics",
        "source company information",
        "kpis",
        "commentary",
        "what worked",
        "what didnt work",
        "targets for next month",
    }
    if key in skipped_exact:
        return True
    skipped_fragments = (
        "source:",
        "check",
        "reporting link",
        "financial statements",
        "company information",
    )
    return any(fragment in key for fragment in skipped_fragments)


def _source_cell(sheet_name: str, row: int, col: int) -> str:
    return f"{sheet_name}!{get_column_letter(col)}{row}"


def _label_tokens(value: str | None) -> set[str]:
    if not value:
        return set()
    ignored = {"and", "of", "the", "to", "in", "for", "from", "with", "usd", "pkr"}
    return {
        token
        for token in re.findall(r"[a-z0-9]+", value.lower())
        if token not in ignored
    }


def _compact_label(value: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", value.lower()))


def _similarity_score(expected: str, candidate: str) -> int:
    expected_tokens = _label_tokens(expected)
    candidate_tokens = _label_tokens(candidate)
    token_score = 0
    if expected_tokens and candidate_tokens:
        overlap = len(expected_tokens & candidate_tokens)
        token_score = int(100 * overlap / max(len(expected_tokens), len(candidate_tokens)))

    compact_expected = _compact_label(expected)
    compact_candidate = _compact_label(candidate)
    sequence_score = int(100 * SequenceMatcher(None, compact_expected, compact_candidate).ratio())
    substring_bonus = 10 if compact_expected in compact_candidate or compact_candidate in compact_expected else 0
    return min(100, max(sequence_score, token_score) + substring_bonus)


def _sheet_label_rows(sheet, label_columns: tuple[int, ...]) -> dict[str, list[int]]:
    labels: dict[str, list[int]] = {}
    for row in range(1, sheet.max_row + 1):
        for col in label_columns:
            key = label_key(sheet.cell(row=row, column=col).value)
            if key:
                labels.setdefault(key, []).append(row)
    return labels


def _matching_source_specs(
    kpi_workbook,
    profile: LabelCopyProfile,
    months: tuple[tuple[int, int], ...],
) -> list[tuple[SourceSpec, str, object, dict[tuple[int, int], int], dict[str, list[int]]]]:
    matches = []
    for spec in profile.source_specs:
        sheet_name = infer_profile_source_sheet_name(kpi_workbook, profile, spec.sheet_name, months)
        if sheet_name is None:
            continue
        sheet = kpi_workbook[sheet_name]
        if spec.month_value_columns:
            month_cols = {
                (year, month): spec.month_value_columns[month]
                for year, month in months
                if month in spec.month_value_columns
            }
        elif spec.sheet_month and spec.value_column:
            month_cols = {
                (year, month): spec.value_column
                for year, month in months
                if month == spec.sheet_month
            }
        else:
            month_cols = find_flexible_month_columns(sheet, months)
        matches.append((spec, sheet_name, sheet, month_cols, _sheet_label_rows(sheet, spec.label_columns)))
    return matches


def _sheet_keys_match(left: str, right: str) -> bool:
    return _compact_label(left) == _compact_label(right)


def _available_sheet_label_keys(sheet, max_rows: int = 350) -> set[str]:
    labels: set[str] = set()
    for row in range(1, min(sheet.max_row, max_rows) + 1):
        for col in range(1, min(sheet.max_column, 6) + 1):
            key = label_key(sheet.cell(row=row, column=col).value)
            if key:
                labels.add(key)
    return labels


def _expected_source_label_keys(profile: LabelCopyProfile) -> set[str]:
    expected: set[str] = set()
    expected.update(label_key(value) for value in (profile.aliases or {}).values())
    expected.update(label_key(value) for value in memory_aliases_for_profile(profile.name).values())
    expected.update(label_key(value) for value in (profile.row_rules or {}).keys())
    expected.update(label_key(value) for value in (profile.fixed_month_sheet_rules or {}).keys())
    expected.discard(None)
    return {value for value in expected if value}


def _source_sheet_review_rows(
    kpi_workbook,
    profile: LabelCopyProfile,
    months: tuple[tuple[int, int], ...],
    source_matches: list[tuple[SourceSpec, str, object, dict[tuple[int, int], int], dict[str, list[int]]]],
) -> list[NormalizationRow]:
    if profile.name == "roomy":
        return []

    rows: list[NormalizationRow] = []
    matched_expected = {(spec.sheet_name, sheet_name) for spec, sheet_name, *_ in source_matches}
    approved_sheet_aliases = memory_sheet_aliases_for_profile(profile.name)

    for expected_sheet, actual_sheet in matched_expected:
        if _sheet_keys_match(expected_sheet, actual_sheet):
            continue
        expected_key = normalise_memory_key(expected_sheet)
        approved_actual_sheet = approved_sheet_aliases.get(expected_key or "")
        if approved_actual_sheet and _sheet_keys_match(approved_actual_sheet, actual_sheet):
            continue
        rows.append(
            NormalizationRow(
                company=profile.display_name,
                month=", ".join(_month_label(month_key) for month_key in months),
                template_cell=profile.report_sheet,
                metric="Workbook tab",
                status="Review",
                confidence=85,
                source=actual_sheet,
                raw_value="",
                normalized_value="",
                conversion="source sheet alias",
                note=f"Possible renamed source sheet: expected '{expected_sheet}' -> actual '{actual_sheet}'.",
            )
        )

    if source_matches:
        return rows

    expected_labels = _expected_source_label_keys(profile)
    for spec in profile.source_specs:
        best_sheet = ""
        best_score = 0
        for sheet_name in kpi_workbook.sheetnames:
            sheet = kpi_workbook[sheet_name]
            month_score = len(find_flexible_month_columns(sheet, months))
            label_score = len(expected_labels & _available_sheet_label_keys(sheet))
            name_score = _similarity_score(spec.sheet_name, sheet_name) // 20
            score = month_score * 5 + label_score * 3 + name_score
            if score > best_score:
                best_score = score
                best_sheet = sheet_name

        if best_sheet and best_score >= 5:
            confidence = min(90, max(55, best_score * 5))
            rows.append(
                NormalizationRow(
                    company=profile.display_name,
                    month=", ".join(_month_label(month_key) for month_key in months),
                    template_cell=profile.report_sheet,
                    metric="Workbook tab",
                    status="Review",
                    confidence=confidence,
                    source=best_sheet,
                    raw_value="",
                    normalized_value="",
                    conversion="source sheet alias",
                    note=f"Possible renamed source sheet: expected '{spec.sheet_name}' -> actual '{best_sheet}'.",
                )
            )
    return rows


def _find_label_source(
    source_matches: list[tuple[SourceSpec, str, object, dict[tuple[int, int], int], dict[str, list[int]]]],
    metric_key: str,
    month_key: tuple[int, int],
) -> tuple[str, object] | None:
    for _spec, sheet_name, sheet, month_cols, label_rows in source_matches:
        source_col = month_cols.get(month_key)
        if source_col is None:
            continue
        source_rows = label_rows.get(metric_key)
        if not source_rows:
            continue
        source_row = source_rows[0]
        return _source_cell(sheet_name, source_row, source_col), sheet.cell(row=source_row, column=source_col).value
    return None


def _find_suggested_label_source(
    source_matches: list[tuple[SourceSpec, str, object, dict[tuple[int, int], int], dict[str, list[int]]]],
    metric_key: str,
    month_key: tuple[int, int],
    min_score: int = 72,
    excluded_keys: set[str] | None = None,
) -> tuple[str, object, str, int] | None:
    best: tuple[str, object, str, int] | None = None
    excluded_keys = excluded_keys or set()
    for _spec, sheet_name, sheet, month_cols, label_rows in source_matches:
        source_col = month_cols.get(month_key)
        if source_col is None:
            continue
        for candidate_key, rows in label_rows.items():
            if not candidate_key or _should_skip_report_label(candidate_key):
                continue
            if candidate_key in excluded_keys:
                continue
            score = _similarity_score(metric_key, candidate_key)
            if score < min_score:
                continue
            source_row = rows[0]
            value = sheet.cell(row=source_row, column=source_col).value
            candidate = (
                _source_cell(sheet_name, source_row, source_col),
                value,
                candidate_key,
                score,
            )
            if best is None or candidate[3] > best[3]:
                best = candidate
    return best


def _find_row_rule_source(
    kpi_workbook,
    profile_name: str,
    rule: tuple[tuple[str, int, float], ...],
    month_key: tuple[int, int],
) -> tuple[str, object, str] | None:
    sources: list[str] = []
    values: list[object] = []
    signs: list[float] = []
    for sheet_name, source_row, sign in rule:
        sheet = profile_workbook_sheet(kpi_workbook, profile_name, sheet_name)
        if sheet is None:
            return None
        source_col = find_flexible_month_columns(sheet, (month_key,)).get(month_key)
        if source_col is None:
            return None
        sources.append(_source_cell(sheet.title, source_row, source_col))
        values.append(sheet.cell(row=source_row, column=source_col).value)
        signs.append(sign)
    conversion = " + ".join(_conversion_note(sign=sign) for sign in signs)
    return "; ".join(sources), " + ".join(_display(value) for value in values), conversion


def _find_fixed_rule_source(
    kpi_workbook,
    profile: LabelCopyProfile,
    rule: tuple[int, int, float],
    month_key: tuple[int, int],
) -> tuple[str, object, str] | None:
    source_row, source_col, sign = rule
    sheet = find_workbook_month_sheet(kpi_workbook, month_key)
    if sheet is None:
        for spec in profile.source_specs:
            if spec.sheet_month == month_key[1]:
                sheet_name = workbook_sheet_name(kpi_workbook, spec.sheet_name)
                if sheet_name is not None:
                    sheet = kpi_workbook[sheet_name]
                    break
    if sheet is None:
        return None
    return (
        _source_cell(sheet.title, source_row, source_col),
        sheet.cell(row=source_row, column=source_col).value,
        _conversion_note(sign=sign),
    )


def _label_copy_normalization_rows(
    template_workbook,
    kpi_workbook,
    profile: LabelCopyProfile,
    months: tuple[tuple[int, int], ...],
) -> list[NormalizationRow]:
    rows: list[NormalizationRow] = []
    report_sheet = template_workbook[profile.report_sheet]
    target_cols = find_flexible_month_columns(report_sheet, months)
    memory_aliases = memory_aliases_for_profile(profile.name)
    aliases = {label_key(k): label_key(v) for k, v in (profile.aliases or {}).items()}
    aliases.update(memory_aliases)
    row_rules = {label_key(k): v for k, v in (profile.row_rules or {}).items()}
    fixed_rules = {label_key(k): v for k, v in (profile.fixed_month_sheet_rules or {}).items()}
    zero_as_dash = {label_key(label) for label in profile.zero_as_dash_labels}
    ignored_rows = set(profile.ignored_rows)
    source_matches = _matching_source_specs(kpi_workbook, profile, months)
    rows.extend(_source_sheet_review_rows(kpi_workbook, profile, months, source_matches))
    alias_values = {label_key(value) for value in (profile.aliases or {}).values()}
    alias_values.discard(None)
    expected_keys = {
        key for key in [
            *(label_key(key) for key in (profile.aliases or {}).keys()),
            *(label_key(value) for value in (profile.aliases or {}).values()),
            *(label_key(key) for key in row_rules.keys()),
            *(label_key(key) for key in fixed_rules.keys()),
            *memory_aliases.keys(),
            *memory_aliases.values(),
        ] if key
    }

    # All template metric keys on this report sheet. Used to stop the
    # suggestion engine from proposing a source label that rightfully belongs
    # to a different template row.
    template_metric_keys: set[str] = set()
    for row in range(1, report_sheet.max_row + 1):
        template_metric = first_report_label(report_sheet, row)
        if _should_skip_report_label(template_metric):
            continue
        template_metric_key = label_key(template_metric)
        if template_metric_key:
            template_metric_keys.add(template_metric_key)

    # (row index, month, suggested candidate key, metric was pre-mapped)
    suggestion_meta: list[tuple[int, tuple[int, int], str, bool]] = []

    if not source_matches:
        return [
            NormalizationRow(
                company=profile.display_name,
                month=_month_label(month_key),
                template_cell=profile.report_sheet,
                metric="Workbook",
                status="Missing",
                confidence=0,
                source="",
                raw_value="",
                normalized_value="",
                conversion="",
                note="No recognizable source sheet was found.",
            )
            for month_key in months
        ]

    for month_key in months:
        target_col = target_cols.get(month_key)
        if target_col is None:
            rows.append(
                NormalizationRow(
                    company=profile.display_name,
                    month=_month_label(month_key),
                    template_cell=profile.report_sheet,
                    metric="Month column",
                    status="Missing",
                    confidence=0,
                    source="",
                    raw_value="",
                    normalized_value="",
                    conversion="",
                    note="Template month column was not found.",
                )
            )
            continue

        auto_ready_samples = 0
        for row in range(1, report_sheet.max_row + 1):
            if row in ignored_rows:
                continue
            if row in profile.formula_rows or row in profile.date_rows:
                continue
            metric = first_report_label(report_sheet, row)
            if _should_skip_report_label(metric):
                continue
            metric_key = label_key(metric)
            if not metric_key:
                continue

            template_cell = report_sheet.cell(row=row, column=target_col).coordinate
            lookup_key = aliases.get(metric_key, metric_key)
            source: str = ""
            raw_value: object = ""
            conversion = "direct"
            status = "Missing"
            confidence = 0
            note = "No matching source metric was found."

            if profile.name == "bykea" and metric_key in {"driver incentives", "marketing"}:
                bykea_sources = get_bykea_visible_summary_sources(kpi_workbook, metric_key, month_key)
                result = get_bykea_visible_summary_source(kpi_workbook, metric_key, month_key)
                if result is not None and bykea_sources:
                    _source_sheet, _source_row, _source_col, raw_value = result
                    source = "; ".join(
                        _source_cell(source_sheet.title, source_row, source_col)
                        for source_sheet, source_row, source_col, _value in bykea_sources
                    )
                    status = "Auto-ready"
                    confidence = 100
                    note = "Visible collapsed Bykea summary row."
                    if len(bykea_sources) > 1:
                        note = "Visible collapsed Bykea summary rows combined for pre-2026 convention."
            elif metric_key in row_rules:
                result = _find_row_rule_source(kpi_workbook, profile.name, row_rules[metric_key], month_key)
                if result is not None:
                    source, raw_value, conversion = result
                    status = "Auto-ready"
                    confidence = 100
                    note = "Explicit mapped rule."
            elif metric_key in fixed_rules:
                result = _find_fixed_rule_source(kpi_workbook, profile, fixed_rules[metric_key], month_key)
                if result is not None:
                    source, raw_value, conversion = result
                    status = "Auto-ready"
                    confidence = 100
                    note = "Fixed monthly source position."
            elif lookup_key in alias_values:
                result = _find_label_source(source_matches, lookup_key, month_key)
                if result is not None:
                    source, raw_value = result
                    status = "Auto-ready"
                    confidence = 100 if lookup_key != metric_key else 95
                    if metric_key in memory_aliases:
                        note = "Approved memory match."
                    else:
                        note = "Approved alias match." if lookup_key != metric_key else "Exact label match."
            else:
                result = _find_label_source(source_matches, lookup_key, month_key)
                if result is not None:
                    source, raw_value = result
                    status = "Auto-ready"
                    confidence = 95
                    note = "Exact label match."

            if status == "Missing":
                # Pre-mapped metrics get a lower threshold; every other
                # template metric is still eligible for a rename suggestion,
                # but only at high similarity to keep the review table quiet.
                min_score = 72 if metric_key in expected_keys else 82
                suggestion = _find_suggested_label_source(
                    source_matches,
                    lookup_key,
                    month_key,
                    min_score=min_score,
                    excluded_keys=(template_metric_keys - {metric_key}) | alias_values,
                )
                if suggestion is not None:
                    source, raw_value, suggested_label, confidence = suggestion
                    status = "Review"
                    conversion = "suggested label match"
                    note = f"Possible renamed source label: {suggested_label}."
                    suggested_candidate_key = suggested_label
                else:
                    suggested_candidate_key = None
            else:
                suggested_candidate_key = None

            if status == "Missing" and metric_key not in expected_keys:
                continue

            if status == "Auto-ready" and profile.round_values:
                conversion = f"{conversion}, round" if conversion != "direct" else "round"
            if metric_key in zero_as_dash:
                conversion = f"{conversion}, zero to dash" if conversion != "direct" else "zero to dash"

            # Keep the UI reviewable: show all issues, all review rows, and a
            # small sample of automatic rows per company/month.
            if status == "Auto-ready":
                auto_ready_samples += 1
                if auto_ready_samples > 8:
                    continue

            rows.append(
                NormalizationRow(
                    company=profile.display_name,
                    month=_month_label(month_key),
                    template_cell=f"{profile.report_sheet}!{template_cell}",
                    metric=metric,
                    status=status,
                    confidence=confidence,
                    source=source,
                    raw_value=_display(raw_value),
                    normalized_value=_display(raw_value),
                    conversion=conversion,
                    note=note,
                )
            )
            if suggested_candidate_key is not None:
                suggestion_meta.append(
                    (len(rows) - 1, month_key, suggested_candidate_key, metric_key in expected_keys)
                )

    # Ambiguity guard: if one source label is the "best rename suggestion" for
    # several different template metrics in the same month, the suggestions are
    # guesses, not renames. Drop them for metrics that were never pre-mapped.
    candidate_counts = Counter((month_key, candidate) for _idx, month_key, candidate, _pre in suggestion_meta)
    drop_indices = {
        idx
        for idx, month_key, candidate, pre_mapped in suggestion_meta
        if not pre_mapped and candidate_counts[(month_key, candidate)] > 1
    }
    if drop_indices:
        rows = [row for idx, row in enumerate(rows) if idx not in drop_indices]

    return rows


def _abhi_normalization_rows(
    template_workbook,
    kpi_workbook,
    profile: CompanyProfile,
    months: tuple[tuple[int, int], ...],
) -> list[NormalizationRow]:
    rows: list[NormalizationRow] = []
    report_sheet = template_workbook[profile.report_sheet]
    sections = resolve_abhi_report_sections(report_sheet, profile.sections)

    for section in sections:
        sheet_name = profile_workbook_sheet_name(kpi_workbook, profile.name, section.source_sheet)
        if sheet_name is None:
            rows.append(
                NormalizationRow(
                    company=profile.display_name,
                    month="All selected",
                    template_cell=f"{profile.report_sheet}!{section.start_row}:{section.end_row}",
                    metric=section.source_sheet,
                    status="Missing",
                    confidence=0,
                    source="",
                    raw_value="",
                    normalized_value="",
                    conversion="",
                    note="ABHI source section was not found.",
                )
            )
            continue

        source_sheet = kpi_workbook[sheet_name]
        source_cols = find_flexible_month_columns(source_sheet, months)
        source_labels: set[str] = set()
        for row in range(1, source_sheet.max_row + 1):
            label = abhi_label_key(source_sheet.cell(row=row, column=3).value) or abhi_label_key(
                source_sheet.cell(row=row, column=2).value
            )
            if label:
                source_labels.add(label)

        report_labels: set[str] = set()
        for row in range(section.start_row, section.end_row + 1):
            label = abhi_report_label(report_sheet, row)
            if label and "financial statements" not in label:
                report_labels.add(label)

        missing_labels = sorted(report_labels - source_labels)
        missing_months = [month_key for month_key in months if month_key not in source_cols]
        if missing_months:
            status = "Missing"
            confidence = 0
            note = "Source section exists, but selected month columns were not found."
        elif missing_labels and len(missing_labels) / max(1, len(report_labels)) >= 0.2:
            status = "Review"
            confidence = 70
            note = f"{len(missing_labels)} labels need review. Examples: {', '.join(missing_labels[:5])}."
        else:
            status = "Auto-ready"
            confidence = 95
            note = "Section and month columns recognized."

        rows.append(
            NormalizationRow(
                company=profile.display_name,
                month=", ".join(_month_label(month_key) for month_key in months),
                template_cell=f"{profile.report_sheet}!{section.start_row}:{section.end_row}",
                metric=section.source_sheet,
                status=status,
                confidence=confidence,
                source=sheet_name,
                raw_value="",
                normalized_value="",
                conversion="section label matching",
                note=note,
            )
        )

    return rows


def build_normalization_review(
    template_path: Path,
    kpi_files: Mapping[str, Path],
    months: tuple[tuple[int, int], ...],
) -> dict[str, object]:
    rows: list[NormalizationRow] = []
    try:
        template_workbook = load_workbook(template_path, data_only=True, read_only=False)
    except Exception as exc:
        row = NormalizationRow(
            company="Monthly Reporting Template",
            month=", ".join(_month_label(month_key) for month_key in months),
            template_cell="",
            metric="Template",
            status="Missing",
            confidence=0,
            source="",
            raw_value="",
            normalized_value="",
            conversion="",
            note=f"Could not open template: {exc}",
        )
        return {
            "rows": [row.row()],
            "status_counts": _status_counts([row]),
            "blocked": True,
            "model_status": "Not connected",
        }

    try:
        for profile_name, kpi_path in kpi_files.items():
            profile = PROFILES.get(profile_name)
            if profile is None or profile.report_sheet not in template_workbook.sheetnames:
                continue
            try:
                kpi_workbook = load_workbook(kpi_path, data_only=True, read_only=False)
            except Exception as exc:
                rows.append(
                    NormalizationRow(
                        company=getattr(profile, "display_name", profile_name),
                        month=", ".join(_month_label(month_key) for month_key in months),
                        template_cell="",
                        metric="KPI workbook",
                        status="Missing",
                        confidence=0,
                        source=str(kpi_path),
                        raw_value="",
                        normalized_value="",
                        conversion="",
                        note=f"Could not open KPI workbook: {exc}",
                    )
                )
                continue

            try:
                if isinstance(profile, CompanyProfile):
                    rows.extend(_abhi_normalization_rows(template_workbook, kpi_workbook, profile, months))
                else:
                    rows.extend(_label_copy_normalization_rows(template_workbook, kpi_workbook, profile, months))
            finally:
                kpi_workbook.close()
    finally:
        template_workbook.close()

    counts = _status_counts(rows)
    return {
        "rows": [row.row() for row in rows],
        "status_counts": counts,
        "blocked": counts["Missing"] > 0,
        "model_status": "Normalization checks and approval memory are active. External AI suggestions are not connected yet.",
    }
